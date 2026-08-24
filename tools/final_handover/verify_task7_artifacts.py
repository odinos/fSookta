#!/usr/bin/env python3
"""Independent, source-derived, fail-closed Task 7 verifier."""
from __future__ import annotations
import hashlib,json,os,re,sys,tarfile,zipfile
from pathlib import Path,PurePosixPath
from xml.etree import ElementTree as ET
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

COMMIT='bf8867a2083357cb9d60915bf6c2233801f923d8';VERSION='1.3.11+28';A='{http://schemas.openxmlformats.org/drawingml/2006/main}';P='{http://schemas.openxmlformats.org/presentationml/2006/main}'
TREE='b4ed5fd0c061492c74dba356ed8a114b5f6621ba'
EDITABLE=['09_Security_Privacy_and_Data_Protection_Report.docx','09_Security_and_Access_Control_Matrices.xlsx','10_End_User_Manual.docx','10_Research_Admin_Manual.docx','10_Developer_Handover_Manual.docx','10_Knowledge_Transfer_Deck.pptx','10_Knowledge_Transfer_Minutes.docx','11_Research_Publication_Package.docx','11_Publication_Tables.xlsx']
PDF_PAGES={'09_Security_Privacy_and_Data_Protection_Report.pdf':4,'09_Security_and_Access_Control_Matrices.pdf':15,'10_End_User_Manual.pdf':3,'10_Research_Admin_Manual.pdf':3,'10_Developer_Handover_Manual.pdf':4,'10_Knowledge_Transfer_Deck.pdf':14,'10_Knowledge_Transfer_Minutes.pdf':3,'11_Research_Publication_Package.pdf':3,'11_Publication_Tables.pdf':9}
def sha(p):return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def git_object(kind,body):return hashlib.sha1(f'{kind} {len(body)}\0'.encode()+body).digest()
def tar_identity(path):
 root={};count=0
 with tarfile.open(path) as archive:
  commit=archive.pax_headers.get('comment')
  for member in archive.getmembers():
   if member.isdir():continue
   if member.isfile():data=archive.extractfile(member).read();mode='100755' if member.mode&0o111 else '100644'
   elif member.issym():data=member.linkname.encode();mode='120000'
   else:continue
   node=root;parts=PurePosixPath(member.name).parts
   for part in parts[:-1]:node=node.setdefault(part,{})
   node[parts[-1]]=(mode,git_object('blob',data));count+=1
 def tree(node):
  entries=[]
  for name,value in sorted(node.items(),key=lambda item:(item[0]+('/' if isinstance(item[1],dict) else '')).encode()):
   mode,digest=('40000',tree(value)) if isinstance(value,dict) else value
   entries.append(f'{mode} {name}'.encode()+b'\0'+digest)
  return git_object('tree',b''.join(entries))
 return {'commit':commit,'git_tree_id':tree(root).hex(),'file_count':count}
def materialization_matches(source,archive_path):
 with tarfile.open(archive_path) as archive:
  for member in archive.getmembers():
   target=source/member.name
   if member.isfile():
    if not target.is_file() or sha(target)!=hashlib.sha256(archive.extractfile(member).read()).hexdigest():return False
   elif member.issym():
    if not target.is_symlink() or os.readlink(target)!=member.linkname:return False
 return True
def source_root(root):
 stable=json.loads((root/'manifests/source_snapshot_manifest.json').read_text());assert stable['authoritative_commit']==COMMIT and stable['git_tree_id']==TREE
 identity=json.loads((root/'manifests/task7_source_identity.json').read_text());assert identity['commit']==COMMIT and identity['git_tree_id']==TREE and identity['stable_identity_record']=='manifests/source_snapshot_manifest.json' and identity['stable_identity_record_sha256']==sha(root/'manifests/source_snapshot_manifest.json')
 valid=[];tar_hashes=set()
 for p in sorted((root/'authoritative-materializations').glob('source-*/source')):
  archive=p.parent/'authoritative-source.tar'
  if (p/'pubspec.yaml').is_file() and archive.is_file() and f'version: {VERSION}' in (p/'pubspec.yaml').read_text() and tar_identity(archive)=={'commit':COMMIT,'git_tree_id':TREE,'file_count':602} and materialization_matches(p,archive):valid.append(p);tar_hashes.add(sha(archive))
 assert valid and len(tar_hashes)==1 and next(iter(tar_hashes))==identity['authoritative_tar_sha256']
 for row in identity['stable_aliases']:
  alias=root/row['alias_path'];assert row['alias_path'].startswith('evidence/task7-source/') and alias.is_file() and sha(alias)==row['sha256']==sha(valid[0]/row['source_relative_path'])
 return valid[0]
def doc_text(p):
 d=Document(p);return '\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for r in t.rows for c in r.cells)
def exact_source_facts(root):
 src=source_root(root);profile=(src/'lib/screens/main/profile_tab.dart').read_text();history=(src/'lib/screens/main/history_tab.dart').read_text();training=(src/'lib/screens/main/training_data_export_screen.dart').read_text();state=(src/'lib/app/app_state.dart').read_text();farmer=(src/'lib/screens/main/farmer_manager_screen.dart').read_text()
 assert "'Manage Farmers'" in profile and 'FarmerManagerScreen.routeName' in profile and "'Export Model Training Data'" in profile and 'TrainingDataExportScreen.routeName' in profile
 assert 'exportHistoryRecordCsv' in history and 'exportAllHistoryCsv' in history
 assert 'files: [' in training and 'bundle.dailyLogisticFile.path' in training and 'bundle.xgBoostFile.path' in training
 production_dart='\n'.join(p.read_text(errors='ignore') for p in (src/'lib').rglob('*.dart'))
 assert 'void deleteFarmer' in state and 'Existing history remains' in farmer and 'File.delete' not in production_dart
 env=json.loads((root/'evidence/verification_environment.json').read_text())['environment'];assert 'Flutter 3.41.9' in env['flutter_version_output'] and 'Dart 3.11.5' in env['flutter_version_output'] and 'Xcode 26.6' in env['xcodebuild_output'] and '17F113' in env['xcodebuild_output'] and env['operating_system']=='macOS-26.6.2-arm64-arm-64bit'
 assert sha(src/'assets/models/xgboost_model.onnx')=='dbedb2ab5ce57f3af0cd620e956f30ef34a64beaea493385afd3d27994002efc'
 return src
def verify_documents(root):
 Aroot=root/'artifacts';src=exact_source_facts(root);texts={n:doc_text(Aroot/n) for n in EDITABLE if n.endswith('.docx')};joined='\n'.join(texts.values())
 assert 'Linked path deletion' not in joined and 'Linked local paths may be deleted' not in joined
 assert 'Farmer removal leaves existing history' in joined and 'No File.delete path' in joined
 for n in ('10_End_User_Manual.docx','10_Research_Admin_Manual.docx'):
  t=texts[n]
  for x in ('Profile > Manage Farmers','History > per-record CSV','History > all-visible CSV','Profile > Export Model Training Data','two CSV files','lib/screens/main/profile_tab.dart','lib/screens/main/history_tab.dart','lib/screens/main/training_data_export_screen.dart'):assert x in t,(n,x)
 dev=texts['10_Developer_Handover_Manual.docx']
 for x in ('Flutter 3.41.9 stable','Dart 3.11.5','Xcode 26.6','17F113','macOS 26.6.2 arm64','assets/models/xgboost_model.onnx','dbedb2ab5ce57f3af0cd620e956f30ef34a64beaea493385afd3d27994002efc'):assert x in dev,x
 pub=texts['11_Research_Publication_Package.docx'];assert 'Exact publication evidence contract' in pub and 'Every publication figure/table row records exact source path' in pub
 return {'documents':6,'docx_pages':20}
def verify_workbooks(root):
 Aroot=root/'artifacts';sec_path=Aroot/'09_Security_and_Access_Control_Matrices.xlsx';pub_path=Aroot/'11_Publication_Tables.xlsx';verify_metadata(sec_path);verify_metadata(pub_path);sec=load_workbook(sec_path,data_only=False);pub=load_workbook(pub_path,data_only=False);assert len(sec.sheetnames)==15 and len(pub.sheetnames)==9
 formula=sec['Control Summary']['B9'].value;assert formula=='=IF(COUNTIF(\'Risk Register\'!E5:E9,"Open*")>0,"OPEN ACTIONS","NO OPEN ROWS")'
 risk_statuses=[sec['Risk Register'].cell(r,5).value for r in range(5,10)];derived='OPEN ACTIONS' if any(str(x).startswith('Open') for x in risk_statuses) else 'NO OPEN ROWS';cached=load_workbook(Aroot/'09_Security_and_Access_Control_Matrices.xlsx',data_only=True)['Control Summary']['B9'].value;assert derived==cached=='OPEN ACTIONS'
 allsec='\n'.join(str(c.value or '') for ws in sec for row in ws.iter_rows() for c in row);assert 'Linked path deletion' not in allsec and 'No File.delete path' in allsec and 'farmer removal leaves existing history' in allsec
 controlled={'Implemented - Evidence Available','Partially Implemented','N/A with Rationale','Open - Pending Owner','Open - Pending Researcher','Open - Pending Owner/Researcher','Pending Signature'}
 for wb in (sec,pub):
  for ws in wb:
   headers={ws.cell(4,c).value:c for c in range(1,ws.max_column+1)}
   if 'Status' in headers:assert {ws.cell(r,headers['Status']).value for r in range(5,ws.max_row+1)}<=controlled,(ws.title,headers)
 assert 'Open - Pending Owner/Researcher' in {sec['Risk Register'].cell(r,5).value for r in range(5,10)}
 sheets=('Chapter 4 Tables','Paper 2 Methods','Paper 2 Results','Variable Definitions','Citation Mapping','Figure Registry','Evidence Index');rows=0
 for sn in sheets:
  ws=pub[sn];headers=[ws.cell(4,c).value for c in range(1,ws.max_column+1)];idx={v:i+1 for i,v in enumerate(headers)}
  for key in ('Exact source path','SHA-256','Evidence version / provenance'):assert key in idx,sn
  for r in range(5,ws.max_row+1):
   rel=ws.cell(r,idx['Exact source path']).value;checksum=ws.cell(r,idx['SHA-256']).value;status=ws.cell(r,idx['Evidence version / provenance']).value
   assert rel and '*' not in rel and ';' not in rel and not rel.startswith('authoritative-materializations/') and (root/rel).is_file(),(sn,r,rel);assert checksum==sha(root/rel) and re.fullmatch(r'[0-9a-f]{64}',checksum);assert status;rows+=1
 methods=pub['Paper 2 Methods'];method_paths={methods.cell(r,1).value:methods.cell(r,5).value for r in range(5,methods.max_row+1)}
 expected={'Static analysis':r'^evidence/flutter_analyze_','Automated tests':r'^evidence/flutter_test_','Android build':r'^evidence/build_android_','iOS build':r'^evidence/build_ios_','Case matrices':r'^artifacts/07_Master_Test_and_Verification_Package\.xlsx$'}
 for key,pattern in expected.items():assert re.match(pattern,method_paths.get(key,'')),(key,method_paths.get(key))
 citations=pub['Citation Mapping'];citation_paths={citations.cell(r,1).value:citations.cell(r,5).value for r in range(5,citations.max_row+1)};assert citation_paths['CIT-04'].startswith('evidence/task7-source/') and citation_paths['CIT-05']=='artifacts/08_UAT_Field_Test_and_Usability_Package.xlsx'
 for wb in (sec,pub):
  for ws in wb:
   assert ws.page_setup.orientation=='landscape' and ws.page_setup.fitToWidth==1 and ws.print_area
 return {'workbooks':2,'sheets':24,'publication_bound_rows':rows,'formulas':1,'formula_errors':0}
def verify_metadata(path):
 with zipfile.ZipFile(path) as z:props='\n'.join(z.read(n).decode('utf-8','ignore') for n in z.namelist() if n in ('docProps/core.xml','docProps/app.xml'))
 assert not re.search(r'python|openpyxl|Walnut Exporter|artifact-tool|LibreOffice|Microsoft Excel',props,re.I) and 'SookTa Project' in props
def verify_deck(root):
 path=root/'artifacts/10_Knowledge_Transfer_Deck.pptx';verify_metadata(path)
 with zipfile.ZipFile(path) as z:
  notes=[]
  for i in range(1,15):
   tree=ET.fromstring(z.read(f'ppt/slides/slide{i}.xml'))
   for sp in tree.iter(P+'sp'):
    nv=sp.find(P+'nvSpPr/'+P+'cNvPr');name=nv.get('name') if nv is not None else '';value=''.join(x.text or '' for x in sp.iter(A+'t'))
    if not value or name=='slide-no':continue
    sizes=[int(x.get('sz'))/100 for x in sp.iter() if x.tag in (A+'rPr',A+'defRPr') and x.get('sz')];scales=[int(x.get('fontScale','100000'))/100000 for x in sp.iter(A+'normAutofit')];effective=min(sizes)*min(scales or [1]);minimum=50 if name=='cover-title' else 35 if name=='slide-title' else 24 if name.endswith('-title') or name.startswith('label-') else 16;assert effective>=minimum,(i,name,effective);assert name!='slide-title' or '\n' not in value
  xml='\n'.join(z.read(n).decode('utf-8','ignore') for n in z.namelist() if n.endswith('.xml'))
  for i in range(1,15):
   note_tree=ET.fromstring(z.read(f'ppt/notesSlides/notesSlide{i}.xml'))
   note_text='\n'.join(x.text or '' for x in note_tree.iter(A+'t'))
   blocks=re.findall(r'\[Sources\](.*?)\[/Sources\]',note_text,re.S);assert len(blocks)==1,i;notes.extend(blocks)
 assert len(notes)==14 and '12 sign-off controls' not in xml and not re.search(r'Task 2 raw logs|03 diagrams|Task 5 reports|authoritative-materializations/source-[^/]+',xml) and 'evidence/task7-source/' in xml
 for block in notes:
  for line in [x.strip() for x in block.splitlines() if x.strip()]:
   rel,checksum,status=[x.strip() for x in line.split(' | ',2)];assert (root/rel).is_file() and sha(root/rel)==checksum and status
 return {'slides':14,'source_note_blocks':14}
def independent_security_scan(root,src):
 secret_patterns={'private_key_marker':re.compile(rb'BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY'),'aws_access_key_marker':re.compile(rb'AKIA[0-9A-Z]{16}'),'github_token_marker':re.compile(rb'(?:ghp|github_pat)_[0-9A-Za-z_]{20,}'),'google_api_key_marker':re.compile(rb'AIza[0-9A-Za-z_-]{30,}'),'generic_secret_assignment':re.compile(rb'(?i)(client_secret|password|private_key)\s*[:=]')}
 participant_patterns={'email_address':re.compile(rb'(?i)\b[A-Z0-9._%+-]+@[A-Z0-9.-]+\.[A-Z]{2,}\b'),'thai_phone_number':re.compile(rb'(?<!\d)(?:\+66|0)[689]\d{8}(?!\d)'),'thai_national_id_shape':re.compile(rb'(?<!\d)\d{13}(?!\d)'),'participant_value_assignment':re.compile(rb'(?i)(?:participant|farmer|profile)[_-]?(?:id|name)\s*[:=]\s*[\"\'][^\"\']{2,}[\"\']')}
 domains=('source','source_archives','office','manifests');scope={d:{'content_items_scanned':0,'bytes_scanned':0,'skipped_items':0} for d in domains};secret={d:{k:0 for k in secret_patterns} for d in domains};participants={d:{k:0 for k in participant_patterns} for d in domains};assignments={}
 def consume(domain,logical,data):
  scope[domain]['content_items_scanned']+=1;scope[domain]['bytes_scanned']+=len(data)
  for key,pattern in secret_patterns.items():
   hits=len(pattern.findall(data));secret[domain][key]+=hits
   if domain=='source' and key=='generic_secret_assignment' and hits:assignments[logical]=hits
  for key,pattern in participant_patterns.items():participants[domain][key]+=len(pattern.findall(data))
 excluded={'.git','build','Pods','.dart_tool'}
 for path in src.rglob('*'):
  if not path.is_file() or any(x in excluded for x in path.parts) or path.stat().st_size>2_000_000:
   if path.is_file():scope['source']['skipped_items']+=1
   continue
  try:data=path.read_bytes()
  except OSError:scope['source']['skipped_items']+=1;continue
  consume('source',path.relative_to(src).as_posix(),data)
 for archive_path in sorted((root/'archives').glob('*')):
  if archive_path.name.endswith(('.tar.gz','.tgz','.tar')):
   with tarfile.open(archive_path) as archive:
    for member in archive.getmembers():
     if member.isfile():consume('source_archives',f'{archive_path.name}::{member.name}',archive.extractfile(member).read())
     else:scope['source_archives']['skipped_items']+=1
 for manifest in sorted((root/'manifests').glob('*.json')):
  if manifest.name.startswith('task7_'):scope['manifests']['skipped_items']+=1;continue
  consume('manifests',manifest.relative_to(root).as_posix(),manifest.read_bytes())
 for office in sorted((root/'artifacts').glob('*')):
  if office.suffix.lower() not in {'.docx','.xlsx','.pptx'}:continue
  if office.name in EDITABLE:scope['office']['skipped_items']+=1;continue
  with zipfile.ZipFile(office) as archive:
   for name in archive.namelist():
    if name.endswith(('.xml','.rels','.txt','.csv')):consume('office',f'{office.name}::{name}',archive.read(name))
    else:scope['office']['skipped_items']+=1
 triage=[]
 for relative,occurrences in sorted(assignments.items()):
  if relative.startswith('docs/user_manual_v1_1_1_android/screenshots/') and relative.endswith('_ui.xml'):category='android_ui_boolean_attribute';owner=False
  elif relative=='android/key.properties.example':category='credential_placeholder_example';owner=True
  else:category='unclassified_generic_assignment';owner=True
  triage.append({'path':relative,'occurrences':occurrences,'category':category,'owner_review_required':owner})
 return scope,secret,participants,triage
def verify_security(root):
 p=json.loads((root/'manifests/task7_offline_security_inspection.json').read_text());src=source_root(root);scope,secret,participants,triage=independent_security_scan(root,src)
 assert p['method']=='offline_source_backed_fallback' and not p['sealed_codex_security_report'] and len(p['findings'])==5 and p['files_scanned']==619
 assert p['scan_scope_counts']==scope and p['candidate_marker_counts_by_domain']==secret and p['participant_identifier_counts']==participants
 assert p['candidate_marker_counts']==secret['source'] and secret['source']['generic_secret_assignment']==978
 report_rows=p['generic_secret_assignment_triage']['files'];assert [(x['path'],x['occurrences'],x['category'],x['owner_review_required']) for x in report_rows]==[(x['path'],x['occurrences'],x['category'],x['owner_review_required']) for x in triage]
 pending=sum(x['occurrences'] for x in triage if x['owner_review_required']);assert pending==p['generic_secret_assignment_triage']['pending_owner_review_occurrences']==p['generic_secret_assignment_triage']['untriaged']==2
 assert sum(x['occurrences'] for x in triage if x['category']=='android_ui_boolean_attribute')==976 and sum(x['occurrences'] for x in triage if x['category']=='credential_placeholder_example')==2
 dump=json.dumps(p);assert not re.search(r'AIza[0-9A-Za-z_-]{30,}|BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY|AKIA[0-9A-Z]{16}',dump)
 return {'security_observations':5,'source_files_scanned':619,'generic_assignments_triaged':978,'generic_assignments_pending_owner':2,'archive_members_scanned':scope['source_archives']['content_items_scanned'],'office_members_scanned':scope['office']['content_items_scanned'],'manifests_scanned':scope['manifests']['content_items_scanned']}
def verify_visual(root):
 ep=root/'manifests/task7_expected_visual_renders.json';expected=json.loads(ep.read_text())
 external=root/'manual-decisions/task7_external_visual_decision.json'
 assert external.is_file(),'external visual decision absent'
 record=json.loads((root/'manifests/task7_visual_qa_manifest.json').read_text());decision=json.loads(external.read_text())
 assert expected['count']==116 and expected['counts']=={'docx_page':20,'workbook_sheet':24,'pptx_slide':14,'pdf_page':58};assert record['source']=='external manual decision' and record['decision']=='PASS' and record['expected_manifest_sha256']==sha(ep) and record['external_decision_sha256']==sha(external) and decision['expected_manifest_sha256']==sha(ep)
 assert set(decision['inspected_paths'])=={x['path'] for x in expected['renders']}
 for x in expected['renders']:assert (root/x['path']).is_file() and sha(root/x['path'])==x['sha256']
 return {'visual_surfaces':116}
def verify(root):
 Aroot=root/'artifacts';assert all((Aroot/n).is_file() and (Aroot/f'{Path(n).stem}.pdf').is_file() for n in EDITABLE)
 for n in [x for x in EDITABLE if x.endswith('.docx')]:verify_metadata(Aroot/n)
 pages={n:len(PdfReader(Aroot/n).pages) for n in PDF_PAGES};assert pages==PDF_PAGES and sum(pages.values())==58
 out={'status':'passed_with_human_actions','editable_artifacts':9,'pdf_artifacts':9,'pdf_pages':58};out.update(verify_documents(root));out.update(verify_workbooks(root));out.update(verify_deck(root));out.update(verify_security(root));out.update(verify_visual(root));return out
if __name__=='__main__':print(json.dumps(verify(Path(sys.argv[1] if len(sys.argv)>1 else os.environ.get('SOOKTA_HANDOVER_ROOT','/private/tmp/fsookta-final-handover'))),indent=2,sort_keys=True))
