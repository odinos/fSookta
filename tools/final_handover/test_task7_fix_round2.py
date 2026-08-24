import json
import os
import re
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(os.environ.get("SOOKTA_HANDOVER_ROOT", "/private/tmp/fsookta-final-handover"))
ART = ROOT / "artifacts"
MAN = ROOT / "manifests"
CONTROLLED = {
    "Implemented - Evidence Available",
    "Partially Implemented",
    "N/A with Rationale",
    "Open - Pending Owner",
    "Open - Pending Researcher",
    "Open - Pending Owner/Researcher",
    "Pending Signature",
}


class Task7FixRound2Tests(unittest.TestCase):
    def test_security_scan_records_content_actually_scanned(self):
        report = json.loads((MAN / "task7_offline_security_inspection.json").read_text())
        scope = report["scan_scope_counts"]
        for domain in ("source", "source_archives", "office", "manifests"):
            self.assertGreater(scope[domain]["content_items_scanned"], 0, domain)
            self.assertGreater(scope[domain]["bytes_scanned"], 0, domain)
        self.assertIn("participant_identifier_counts", report)
        self.assertIn("scan_algorithm", report)

    def test_generic_assignment_triage_is_per_file_and_pending_is_honest(self):
        report = json.loads((MAN / "task7_offline_security_inspection.json").read_text())
        triage = report["generic_secret_assignment_triage"]
        self.assertEqual(sum(row["occurrences"] for row in triage["files"]), 978)
        self.assertEqual(sum(row["occurrences"] for row in triage["files"] if row["category"] == "android_ui_boolean_attribute"), 976)
        self.assertEqual(sum(row["occurrences"] for row in triage["files"] if row["category"] == "credential_placeholder_example"), 2)
        self.assertEqual(triage["pending_owner_review_occurrences"], 2)
        self.assertEqual(triage["untriaged"], 2)

    def test_publication_semantic_bindings(self):
        wb = load_workbook(ART / "11_Publication_Tables.xlsx", data_only=False)
        methods = wb["Paper 2 Methods"]
        rows = {methods.cell(r, 1).value: methods.cell(r, 5).value for r in range(5, methods.max_row + 1)}
        self.assertRegex(rows["Static analysis"], r"^evidence/flutter_analyze_")
        self.assertRegex(rows["Automated tests"], r"^evidence/flutter_test_")
        self.assertRegex(rows["Android build"], r"^evidence/build_android_")
        self.assertRegex(rows["iOS build"], r"^evidence/build_ios_")
        self.assertEqual(rows["Case matrices"], "artifacts/07_Master_Test_and_Verification_Package.xlsx")
        citations = wb["Citation Mapping"]
        cite = {citations.cell(r, 1).value: citations.cell(r, 5).value for r in range(5, citations.max_row + 1)}
        self.assertTrue(cite["CIT-04"].startswith("evidence/task7-source/"))
        self.assertEqual(cite["CIT-05"], "artifacts/08_UAT_Field_Test_and_Usability_Package.xlsx")

    def test_no_ephemeral_materialization_citations(self):
        wb = load_workbook(ART / "11_Publication_Tables.xlsx", data_only=False)
        values = "\n".join(str(c.value or "") for ws in wb for row in ws.iter_rows() for c in row)
        with zipfile.ZipFile(ART / "10_Knowledge_Transfer_Deck.pptx") as archive:
            values += "\n" + "\n".join(archive.read(name).decode("utf-8", "ignore") for name in archive.namelist() if name.endswith(".xml"))
        self.assertNotRegex(values, r"authoritative-materializations/source-[^/]+")
        self.assertIn("evidence/task7-source/", values)

    def test_source_identity_is_tree_and_archive_bound(self):
        identity = json.loads((MAN / "task7_source_identity.json").read_text())
        self.assertEqual(identity["commit"], "bf8867a2083357cb9d60915bf6c2233801f923d8")
        self.assertEqual(identity["git_tree_id"], "b4ed5fd0c061492c74dba356ed8a114b5f6621ba")
        self.assertRegex(identity["authoritative_tar_sha256"], r"^[0-9a-f]{64}$")
        self.assertGreaterEqual(len(identity["stable_aliases"]), 5)
        for row in identity["stable_aliases"]:
            self.assertTrue((ROOT / row["alias_path"]).is_file())

    def test_every_control_status_uses_vocabulary(self):
        for name in ("09_Security_and_Access_Control_Matrices.xlsx", "11_Publication_Tables.xlsx"):
            wb = load_workbook(ART / name, data_only=False)
            for ws in wb:
                headers = {ws.cell(4, c).value: c for c in range(1, ws.max_column + 1)}
                if "Status" in headers:
                    values = {ws.cell(r, headers["Status"]).value for r in range(5, ws.max_row + 1)}
                    self.assertTrue(values <= CONTROLLED, (name, ws.title, values - CONTROLLED))
        sec = load_workbook(ART / "09_Security_and_Access_Control_Matrices.xlsx")
        self.assertIn("Open - Pending Owner/Researcher", {sec["Risk Register"].cell(r, 5).value for r in range(5, 10)})

    def test_xlsx_metadata_is_neutral(self):
        for name in ("09_Security_and_Access_Control_Matrices.xlsx", "11_Publication_Tables.xlsx"):
            with zipfile.ZipFile(ART / name) as archive:
                props = "\n".join(archive.read(x).decode("utf-8", "ignore") for x in archive.namelist() if x.startswith("docProps/"))
            self.assertIn("SookTa Project", props)
            self.assertNotRegex(props, r"(?i)artifact-tool|Microsoft Excel|python|openpyxl|Walnut")


if __name__ == "__main__":
    unittest.main()
