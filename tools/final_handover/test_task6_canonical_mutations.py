#!/usr/bin/env python3
import json
import shutil
import tempfile
import unittest
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader, PdfWriter

try:
 from . import verify_task6_artifacts as verifier
except ImportError:
 import verify_task6_artifacts as verifier

ROOT=Path('/private/tmp/fsookta-final-handover')
ARTIFACTS=['06_Development_Audit_Trail.xlsx','06_Development_Audit_Trail.pdf','07_Master_Test_and_Verification_Package.xlsx','07_Master_Test_and_Verification_Package.pdf','08_UAT_Field_Test_and_Usability_Package.xlsx','08_UAT_Field_Test_and_Usability_Package.pdf','07_Final_Algorithm_Verification_Report.docx','07_Final_Algorithm_Verification_Report.pdf','08_UAT_and_Field_Test_Technical_Report.docx','08_UAT_and_Field_Test_Technical_Report.pdf']
MANIFESTS=['task6_expected_visual_renders.json','task6_visual_qa_manifest.json','task6_pdf_page_counts.json']

class CanonicalArtifactMutationTests(unittest.TestCase):
 def sandbox(self):
  td=tempfile.TemporaryDirectory(); root=Path(td.name)
  (root/'artifacts').mkdir(); (root/'manifests').mkdir(); (root/'working/task6').mkdir(parents=True)
  for name in ARTIFACTS: shutil.copy2(ROOT/'artifacts'/name,root/'artifacts'/name)
  for name in MANIFESTS: shutil.copy2(ROOT/'manifests'/name,root/'manifests'/name)
  shutil.copy2(ROOT/'working/task6/task6_corrected_payload.json',root/'working/task6/task6_corrected_payload.json')
  shutil.copy2(ROOT/'evidence_map.json',root/'evidence_map.json')
  for name in ('evidence','authoritative-materializations','renders'): (root/name).symlink_to(ROOT/name,target_is_directory=True)
  return td,root
 def reject(self,mutate):
  td,root=self.sandbox()
  try:
   mutate(root)
   with self.assertRaises((AssertionError,ValueError,KeyError)): verifier.verify(root)
  finally: td.cleanup()
 def mutate_book(self,root,name,sheet,cell,value,validation=False):
  p=root/'artifacts'/name; wb=load_workbook(p,data_only=False); s=wb[sheet]; s[cell]=value
  if validation: s.data_validations.dataValidation=[]
  wb.save(p)
 def test_rejects_algorithm_test_name_mutation(self): self.reject(lambda r:self.mutate_book(r,'07_Master_Test_and_Verification_Package.xlsx','Algorithm Reference','C5','mutated name'))
 def test_rejects_automated_category_mutation(self): self.reject(lambda r:self.mutate_book(r,'07_Master_Test_and_Verification_Package.xlsx','Automated Suite','D5','UI / regression'))
 def test_rejects_cross_sheet_case_mapping_mutation(self): self.reject(lambda r:self.mutate_book(r,'07_Master_Test_and_Verification_Package.xlsx','Threshold Boundaries','A5','AUTO-999'))
 def test_rejects_requirement_mutation(self): self.reject(lambda r:self.mutate_book(r,'06_Development_Audit_Trail.xlsx','Requirement Traceability','B5','Complete'))
 def test_rejects_historical_version_date_round_mutation(self): self.reject(lambda r:self.mutate_book(r,'08_UAT_Field_Test_and_Usability_Package.xlsx','Historical Evidence','D11','1.3.6+21'))
 def test_rejects_generic_firebase_pass(self): self.reject(lambda r:self.mutate_book(r,'07_Master_Test_and_Verification_Package.xlsx','Network Error Applicability','H5','PASS'))
 def test_rejects_sus_formula_and_validation_mutation(self): self.reject(lambda r:self.mutate_book(r,'08_UAT_Field_Test_and_Usability_Package.xlsx','SUS Response Form','D16','=0',validation=True))
 def test_rejects_defect_git_ref_mutation(self): self.reject(lambda r:self.mutate_book(r,'07_Master_Test_and_Verification_Package.xlsx','Bugs Corrections Retest','D5','0'*40))
 def test_rejects_report_count_mutation(self):
  def mutate(root):
   p=root/'artifacts/07_Final_Algorithm_Verification_Report.docx'; d=Document(p)
   for para in d.paragraphs:
    if 'exactly 135 completed' in para.text:
     for run in para.runs: run.text=run.text.replace('exactly 135 completed','exactly 136 completed')
   d.save(p)
  self.reject(mutate)
 def test_rejects_visual_manifest_hash_mutation(self):
  def mutate(root):
   p=root/'manifests/task6_expected_visual_renders.json'; x=json.loads(p.read_text()); x['renders'][0]['sha256']='0'*64; p.write_text(json.dumps(x))
  self.reject(mutate)
 def test_rejects_unmodeled_control_summary_cell_mutation(self):
  self.reject(lambda r:self.mutate_book(r,'06_Development_Audit_Trail.xlsx','Control Summary','B5',999))
 def test_rejects_appended_unverified_docx_content(self):
  def mutate(root):
   p=root/'artifacts/08_UAT_and_Field_Test_Technical_Report.docx'; d=Document(p); d.add_paragraph('MUTATED UNVERIFIED CONTENT'); d.save(p)
  self.reject(mutate)
 def test_rejects_same_page_count_pdf_rewrite(self):
  def mutate(root):
   p=root/'artifacts/07_Final_Algorithm_Verification_Report.pdf'; reader=PdfReader(p); writer=PdfWriter()
   for page in reader.pages: writer.add_page(page)
   with p.open('wb') as stream: writer.write(stream)
   self.assertEqual(len(reader.pages),len(PdfReader(p).pages))
  self.reject(mutate)
 def test_rejects_zero_or_mutated_subject_git_history(self):
  def change(root,field,value):
   p=root/'working/task6/task6_corrected_payload.json'; data=json.loads(p.read_text()); data['git_history'][9][field]=value; p.write_text(json.dumps(data))
  self.reject(lambda r:change(r,'commit','0'*40))
  self.reject(lambda r:change(r,'subject','MUTATED SUBJECT'))

if __name__=='__main__': unittest.main()
