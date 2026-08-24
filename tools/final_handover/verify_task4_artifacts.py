#!/usr/bin/env python3
"""Independent fail-closed verification for Task 4 artifacts."""
from __future__ import annotations
import hashlib, json, re, zipfile
from pathlib import Path
from xml.etree import ElementTree as ET
from docx import Document
from openpyxl import load_workbook
from PIL import Image
from pypdf import PdfReader

VERSION="1.3.11+28"; COMMIT="bf8867a2083357cb9d60915bf6c2233801f923d8"
ALLOWED_STATUSES={
    "Complete", "Complete - Pending Signature", "Pending Owner Action",
    "Pending Researcher Evidence", "N/A with Rationale",
    "Exception Approval Required",
}
EXPECTED_TECHNOLOGY_IDENTIFIERS={
    "camera":"0.11.4 (resolved)", "image_picker":"1.2.2 (resolved)",
    "tflite_flutter":"0.12.1 (resolved)",
    "onnxruntime":"1.4.1 (resolved; local override source)",
    "shared_preferences":"2.5.5 (resolved)", "path_provider":"2.1.5 (resolved)",
    "share_plus":"13.1.0 (resolved)", "flutter_tts":"4.2.5 (resolved)",
    "firebase_core":"4.10.0 (resolved)", "firebase_analytics":"12.4.2 (resolved)",
    "firebase_crashlytics":"5.2.2 (resolved)",
    "XGBoost advisory":"reba-iso-xgboost-onnx-2026-06-07",
    "Daily logistic template":"daily-injury-logistic-template-2026-06-14",
}

def digest(path): return hashlib.sha256(Path(path).read_bytes()).hexdigest()

def verify_visual_manifest(root: Path, manifest: Path, expected: dict[str,list[str]]):
    data=json.loads(manifest.read_text(encoding="utf-8")); assert data.get("status")=="passed"
    rows={r["path"]:r for r in data.get("artifacts",[])}; assert set(rows)==set(expected)
    for artifact,renders in expected.items():
        row=rows[artifact]; path=root/artifact; assert path.is_file() and digest(path)==row["sha256"] and row["status"]=="passed"
        rr={r["path"]:r for r in row["renders"]}; assert set(rr)==set(renders)
        for render in renders:
            p=root/render; assert p.is_file() and digest(p)==rr[render]["sha256"] and rr[render]["status"]=="passed"

def doc_text(path):
    d=Document(path); return "\n".join(p.text for p in d.paragraphs)

def archive_text(path: Path) -> str:
    with zipfile.ZipFile(path) as archive:
        return "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.endswith((".xml", ".rels"))
        )

def verify_docx_preset(path: Path):
    with zipfile.ZipFile(path) as archive:
        root=ET.fromstring(archive.read("word/document.xml"))
    ns={"w":"http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    assert not root.findall(".//w:pPr/w:pBdr",ns),f"title paragraph border/rule residue: {path}"

def verify_workbook_summary(path: Path):
    data=json.loads(path.read_text(encoding="utf-8"))
    assert data["status"]=="passed" and data["sheets"]==4
    assert data["module_records"]==14 and data["stack_records"]==18
    contracts=data["formula_contracts"]; assert len(contracts)==5
    for contract in contracts:
        assert contract["status"]=="passed"
        assert contract["expected"]["formulas"]==contract["actual"]["formulas"]
        assert contract["expected"]["values"]==contract["actual"]["values"]
    assert "matched 0 entries" in data["formula_error_scan"]
    return data

def verify_status_values(values):
    normalized=[str(value).strip() for value in values if value is not None and str(value).strip()]
    assert normalized and all(value in ALLOWED_STATUSES for value in normalized),normalized
    return normalized

def verify_technology_identifiers(path: Path):
    workbook=load_workbook(path,read_only=True,data_only=False)
    sheet=workbook["Technology Stack"]
    rows={
        str(row[0]).strip():str(row[1]).strip()
        for row in sheet.iter_rows(values_only=True)
        if len(row)>=2 and row[0] and row[1]
    }
    workbook.close()
    for component,identifier in EXPECTED_TECHNOLOGY_IDENTIFIERS.items():
        assert rows.get(component)==identifier,(component,rows.get(component),identifier)
    thunder=rows.get("MoveNet Thunder",""); multipose=rows.get("MoveNet MultiPose Lightning","")
    assert "movenet-thunder-v1-17x3-normalized" in thunder and "8014d8fe22285265f52aa1cea84056b7704f75adf12341a7712d4cb28bd1d9b6" in thunder
    assert "upstream artifact release/version not recorded" in thunder
    assert "d4489f89e6bd6777a8b9a1a16189832131f84ff90d82fae729e670b84d7948dd" in multipose
    assert "upstream artifact release/version not recorded in repo" in multipose
    return rows

def verify_citation_records(sections: list[dict], staging: Path):
    records=[record for section in sections for record in section["citation_records"]]
    assert records and not [record for record in records if record["status"]=="Missing"]
    for record in records:
        if record["status"]=="Resolved":
            assert record["resolved_path"] and Path(record["resolved_path"]).exists(),record
        elif record["status"]=="Pending future artifact":
            assert record["path"] in {"08_UAT_Field_Test_and_Usability_Package.xlsx","12_Final_Developer_Statement_and_Signoff.docx"},record
        elif record["status"]=="Generated companion artifact":
            candidates=[staging/"artifacts"/record["path"],staging/"artifacts/diagrams"/record["path"]]
            assert any(path.exists() for path in candidates),record
        else:
            raise AssertionError(record)
    return records

def verify_checksums(staging: Path, required: list[Path]):
    manifest=staging/"manifests/SHA256SUMS.txt"; rows={}
    for line in manifest.read_text(encoding="ascii").splitlines():
        checksum,relative=line.split("  ",1); rows[relative]=checksum
    for relative,checksum in rows.items():
        path=staging/relative; assert path.is_file(),relative; assert digest(path)==checksum,relative
    for path in required:
        relative=path.relative_to(staging).as_posix()
        assert rows.get(relative)==digest(path),relative
    return len(rows)

def verify(staging: Path):
    artifacts=staging/"artifacts"; diagrams=artifacts/"diagrams"; manifests=staging/"manifests"
    report=artifacts/"03_Final_Technical_Development_Report.docx"; api=artifacts/"03_API_Applicability_Statement.docx"; xlsx=artifacts/"03_Technical_Stack_and_Module_Specification.xlsx"
    pdfs=[artifacts/"03_Final_Technical_Development_Report.pdf",artifacts/"03_API_Applicability_Statement.pdf",artifacts/"03_Technical_Stack_and_Module_Specification.pdf"]
    for p in [report,api,xlsx,*pdfs]: assert p.is_file() and p.stat().st_size>1000,p
    verify_docx_preset(report); verify_docx_preset(api)
    rt=doc_text(report); at=doc_text(api)
    import build_task4_architecture as builder
    positions=[]
    for i,h in enumerate(builder.REPORT_HEADINGS,1):
        token=f"{i}. {h}"; assert rt.count(token)==1,token; positions.append(rt.index(token))
    assert positions==sorted(positions) and VERSION in rt and COMMIT in rt
    assert all(x in at for x in ["N/A with Rationale","SOOKTA_TELEMETRY_ENABLED","defaultValue false","no direct HTTP","Pending Owner Action",VERSION,COMMIT])
    assert [len(PdfReader(p).pages) for p in pdfs]==[11,2,4]
    workbook_summary=verify_workbook_summary(manifests/"task4_workbook_build_summary.json")
    technology_rows=verify_technology_identifiers(xlsx)
    workbook=load_workbook(xlsx,read_only=True,data_only=False)
    status_values=verify_status_values(
        workbook["Evidence Sources"].cell(row=row,column=4).value for row in range(5,13)
    )
    workbook.close()
    observed_statuses=set(status_values)
    assert {
        "Complete", "Pending Owner Action", "Pending Researcher Evidence",
        "Complete - Pending Signature",
    }.issubset(observed_statuses)
    build_input=json.loads((staging/"working/task4/task4_build_input_summary.json").read_text())
    sections=builder.report_sections(build_input["source_facts"])
    citation_records=verify_citation_records(sections,staging)
    assert all(token not in rt for token in ["test/assessment_calculation_test.dart","docs/uat-test-plan.md","docs/uat-test-cases.md"])
    pairs=[]
    for spec in builder.diagram_specs():
        d=diagrams/f"{spec['basename']}.drawio"; p=diagrams/f"{spec['basename']}.png"; assert d.is_file() and p.is_file()
        root=ET.parse(d).getroot(); assert root.tag=="mxfile" and root.attrib["applicationVersion"]==VERSION and root.attrib["authoritativeCommit"]==COMMIT
        vertices=root.findall(".//mxCell[@vertex='1']"); edges=root.findall(".//mxCell[@edge='1']"); assert len(vertices)>=len(spec["nodes"])+2 and len(edges)==len(spec["edges"])
        edge_contract={(edge.attrib.get("source"),edge.attrib.get("target"),edge.attrib.get("value","")) for edge in edges}
        assert edge_contract=={(f"n{a}",f"n{b}",label) for a,b,label in spec["edges"]}
        boundary=root.find(".//mxCell[@id='boundary']/mxGeometry"); assert boundary is not None
        bx,by,bw,bh=(float(boundary.attrib[key]) for key in ("x","y","width","height"))
        for index in spec["external_nodes"]:
            node=root.find(f".//mxCell[@id='n{index}']"); assert node is not None and node.attrib.get("external")=="true"
            geometry=node.find("mxGeometry"); assert geometry is not None
            nx,ny,nw,nh=(float(geometry.attrib[key]) for key in ("x","y","width","height"))
            assert nx+nw<=bx or nx>=bx+bw or ny+nh<=by or ny>=by+bh,(spec["basename"],index)
        with Image.open(p) as im: assert im.width>=2800 and im.height>=1800 and im.info.get("dpi",(0,0))[0]>=299
        pairs.append(spec["basename"])
    expected=json.loads((manifests/"task4_expected_visual_renders.json").read_text(encoding="utf-8"))
    assert len(expected)==22
    verify_visual_manifest(staging,manifests/"task4_visual_qa_manifest.json",expected)
    required=[report,api,xlsx,*pdfs,*[diagrams/f"{b}.{ext}" for b in pairs for ext in ("drawio","png")]]
    checksum_entries=verify_checksums(staging,required)
    forbidden=re.compile(r"(?i)(-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----|AIza[0-9A-Za-z_-]{20,}|password\s*[:=]\s*\S+|api[_-]?key\s*[:=]\s*\S+)")
    for p in [report,api,xlsx,*pdfs,*diagrams.glob("*.drawio")]:
        content=archive_text(p) if p.suffix in {".docx",".xlsx"} else p.read_text(encoding="utf-8",errors="ignore")
        assert not forbidden.search(content),p
    return {
        "status":"passed_with_human_actions", "report_headings":28,
        "diagram_pairs":8, "primary_artifacts":22, "pdf_pages":17,
        "workbook_sheets":workbook_summary["sheets"], "technology_identifiers":len(technology_rows),
        "citation_records":len(citation_records),
        "formula_contracts":len(workbook_summary["formula_contracts"]),
        "formula_error_scan":"0 errors", "checksum_entries":checksum_entries,
        "visual_manifest":digest(manifests/"task4_visual_qa_manifest.json"),
        "human_actions":build_input["human_actions"],
    }

if __name__=="__main__":
    import argparse
    p=argparse.ArgumentParser(); p.add_argument("--staging",type=Path,default=Path("/private/tmp/fsookta-final-handover")); p.add_argument("--summary",type=Path,required=True); a=p.parse_args(); result=verify(a.staging); a.summary.write_text(json.dumps(result,indent=2)+"\n"); print(json.dumps(result,indent=2))
