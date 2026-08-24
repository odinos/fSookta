#!/usr/bin/env python3
"""Build source-grounded Task 4 DOCX and editable diagram artifacts."""

from __future__ import annotations

import argparse
import hashlib
import json
import math
import re
from pathlib import Path
from xml.sax.saxutils import escape, quoteattr

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor
from PIL import Image, ImageDraw, ImageFont

VERSION = "1.3.11+28"
COMMIT = "bf8867a2083357cb9d60915bf6c2233801f923d8"
TREE = "b4ed5fd0c061492c74dba356ed8a114b5f6621ba"
STAGING = Path("/private/tmp/fsookta-final-handover")
REPORT_HEADINGS = [
    "Project background and objectives", "Original vs final requirements",
    "Development methodology/process", "Development timeline and milestones",
    "System architecture", "Technology stack", "Application modules and user workflow",
    "Image acquisition and processing workflow", "Ergonomic assessment algorithm",
    "AI/model description and role", "Recommendation engine and source of rules",
    "Database architecture and data flow", "Prototype development and iterative modifications",
    "Verification and testing methodology", "Functional test results",
    "Algorithm/reference verification", "Integration/system/regression tests",
    "Compatibility/device tests", "Performance tests", "UAT and field-test technical support",
    "Bugs, corrective actions and retesting", "Final application specification",
    "Security, privacy and data management", "Technical challenges and solutions",
    "Technical limitations and known issues", "Scalability/maintainability",
    "Future development recommendations", "Final acceptance and handover statement",
]


def sha(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def inspect_source(root: Path) -> dict:
    dart_files = sorted((root / "lib").rglob("*.dart")) if (root / "lib").exists() else []
    joined = "\n".join(path.read_text(encoding="utf-8", errors="replace") for path in dart_files)
    pubspec = (root / "pubspec.yaml").read_text(encoding="utf-8", errors="replace")
    backend_patterns = [
        r"package:http/", r"package:dio/", r"\bHttpClient\s*\(", r"\bWebSocket\s*\.",
        r"package:graphql", r"package:firebase_database", r"package:cloud_firestore",
        r"package:firebase_functions", r"package:firebase_storage",
    ]
    backend_hits = [pattern for pattern in backend_patterns if re.search(pattern, joined)]
    declared_version = re.search(r"(?m)^version:\s*([^\s]+)", pubspec)
    telemetry_declared = "firebase_analytics:" in pubspec or "firebase_crashlytics:" in pubspec
    opt_in = bool(re.search(r"SOOKTA_TELEMETRY_ENABLED[\s\S]{0,160}defaultValue:\s*false", joined))
    local_storage = "shared_preferences:" in pubspec and "SharedPreferences" in joined
    local_models = "tflite_flutter:" in pubspec and "onnxruntime:" in pubspec
    lock_path = root / "pubspec.lock"
    resolved_packages = {}
    if lock_path.is_file():
        current = None
        for line in lock_path.read_text(encoding="utf-8", errors="replace").splitlines():
            if line.startswith("  ") and not line.startswith("    ") and line.endswith(":"):
                current = line.strip()[:-1]
            elif current and line.startswith("    version:"):
                resolved_packages[current] = line.split(":", 1)[1].strip().strip('"')
    xgboost_metadata = root / "assets/models/xgboost_model_metadata.json"
    daily_metadata = root / "assets/ml/daily_injury_logistic_model.json"
    schema_metadata = root / "assets/models/joint_feature_schema.json"
    xgboost = json.loads(xgboost_metadata.read_text(encoding="utf-8")) if xgboost_metadata.is_file() else {}
    daily = json.loads(daily_metadata.read_text(encoding="utf-8")) if daily_metadata.is_file() else {}
    schema = json.loads(schema_metadata.read_text(encoding="utf-8")) if schema_metadata.is_file() else {}
    thunder_asset = root / "assets/ml/movenet_thunder.tflite"
    multipose_asset = root / "assets/ml/movenet_multipose_lightning.tflite"
    return {
        "application_version": declared_version.group(1) if declared_version else "Unresolved",
        "authoritative_commit": COMMIT,
        "git_tree_id": TREE,
        "assessment_backend_api_present": bool(backend_hits),
        "backend_client_hits": backend_hits,
        "telemetry_declared": telemetry_declared,
        "telemetry_opt_in_default_off": opt_in,
        "local_storage_present": local_storage,
        "local_models_present": local_models,
        "api_status": "N/A with Rationale" if not backend_hits else "API documentation required",
        "scanned_dart_files": len(dart_files),
        "source_paths": [str(p.relative_to(root)) for p in dart_files],
        "source_root": str(root),
        "staging_root": str(STAGING),
        "resolved_packages": resolved_packages,
        "model_versions": {
            "xgboost": xgboost.get("version", "Unresolved"),
            "daily_logistic": daily.get("version", "Unresolved"),
            "movenet_schema": schema.get("schemaId", "Unresolved"),
            "movenet_schema_version": schema.get("version", "Unresolved"),
            "movenet_thunder_sha256": sha(thunder_asset) if thunder_asset.is_file() else "Unresolved",
            "movenet_multipose_sha256": sha(multipose_asset) if multipose_asset.is_file() else "Unresolved",
            "movenet_upstream_version": "Not recorded in repository metadata",
        },
    }


def minimum_facts_fixture() -> dict:
    return {
        "application_version": VERSION, "authoritative_commit": COMMIT, "git_tree_id": TREE,
        "assessment_backend_api_present": False, "telemetry_declared": True,
        "telemetry_opt_in_default_off": True, "local_storage_present": True,
        "local_models_present": True, "api_status": "N/A with Rationale", "scanned_dart_files": 40,
        "source_root": "/private/tmp/fsookta-final-handover/authoritative-materializations/source-g6tsggy7/source",
        "staging_root": str(STAGING), "resolved_packages": {}, "model_versions": {},
    }


SECTION_DATA = [
    ("Current behavior", "SookTa is an offline-first Flutter application supporting ergonomic risk assessment and research follow-up for coffee-farming work. The final technical objective is traceable local assessment, explainable output, maintainable source, and privacy-bounded export; it is not a medical device claim.", ["README.md", "lib/app/sookta_app.dart"]),
    ("Human-owned scope closure", "The implemented baseline is documented from source and reproducible evidence. Original contract scope, later-request authorization, and commercial acceptance cannot be inferred from code and remain Pending Owner Action until contract and approval records are signed.", ["/private/tmp/fsookta-final-handover/requirements.json", "01_Final_Release_and_Scope_Closure.docx"]),
    ("Historical evidence plus current state", "Repository history and versioned audit notes show iterative implementation and corrective review. Current-state claims in this report are re-grounded to the authoritative final commit; historical documents retain their own versions and dates.", ["README.md", "docs/final-revision-p1-audit-20260623.md"]),
    ("Evidence-bounded timeline", "Milestones can be reconstructed from Git and dated repository evidence, but contractual dates, effort, and approval milestones require owner confirmation. This report therefore separates technical chronology from human-owned schedule acceptance.", ["docs/system-baseline-v1.3.6+21-20260629.md", "final_source_metadata.txt"]),
    ("Current architecture", "Flutter presentation screens call local domain services for pose inference, deterministic assessment, advisory models, recommendation, persistence, and CSV generation. Assessment execution does not depend on a backend service; optional Firebase telemetry is outside the decision path and defaults off.", ["lib/screens/", "lib/core/services/", "03_system_context.drawio"]),
    ("Pinned stack", "The baseline uses Flutter 3.41.9, Dart 3.11.5, Flutter SDK constraints from pubspec.yaml, TFLite MoveNet, ONNX Runtime, SharedPreferences, path-provider, camera/image-picker, share-plus, TTS, and opt-in Firebase Analytics/Crashlytics. Resolved package detail is recorded in the companion workbook.", ["pubspec.yaml", "pubspec.lock", "verification_environment.json"]),
    ("Current workflow", "The user selects language/profile, chooses an activity, captures or selects media, completes task inputs, reviews calculated risk and recommendations, saves history, and explicitly exports/shares CSV when needed. Screen and service boundaries are listed in the module workbook.", ["lib/screens/onboarding/", "lib/screens/main/", "03_user_navigation.drawio"]),
    ("Local image path", "Camera/gallery images and extracted video frames are decoded locally; MoveNet input is letterboxed to 256 by 256, 17 keypoints are produced, and low-confidence or unreadable input returns an unavailable result for UI handling. Persisted media is copied into application documents storage.", ["lib/core/services/pose_image_preprocessor.dart", "lib/core/services/pose_estimation_service.dart", "lib/core/services/local_image_store.dart"]),
    ("Deterministic primary result", "Pose-derived and user-supplied inputs feed table-based REBA scoring for every activity. Applied ISO 11228 lifting or push/pull calculations supplement relevant tasks, and the combined result preserves the higher applicable risk. This is software/reference verification, not clinical validation.", ["lib/core/services/ergo_calculator.dart", "lib/core/models/evaluation_models.dart", "03_assessment_algorithm.drawio"]),
    ("Separated model roles", "MoveNet Thunder performs local pose estimation; XGBoost/ONNX provides an advisory alert attached to the deterministic result; the daily logistic JSON is a template-coefficient prediction input. None of these roles may be described as clinical validation, and advisory output does not replace REBA/ISO.", ["lib/core/services/pose_estimation_service.dart", "lib/core/services/xgboost_advisory_service.dart", "lib/core/services/daily_injury_prediction_service.dart"]),
    ("Source-grounded rules", "Recommendation output is deterministic and keyed by activity, risk tier, and body area. The service deduplicates items and bounds recommendations per category; bilingual strings and reference keys are source-controlled. Researcher approval of rule wording remains a human action.", ["lib/core/services/risk_recommendation_service.dart", "lib/core/models/assessment_reference_sources.dart"]),
    ("Local persistence", "Profiles, active profile, drafts, schema version, backup, and history are serialized in SharedPreferences. Images and generated exports use application document or temporary directories. No remote database schema, server migration, or backend data API is present in the assessed source.", ["lib/app/app_state.dart", "lib/core/services/local_image_store.dart", "lib/core/services/assessment_export_service.dart"]),
    ("Historical evolution", "Versioned audit notes and Git history document prototype-to-final changes. They are supporting historical evidence only; unavailable original design artifacts and undocumented feedback must be supplied or approved as substitutions by the owner/researcher.", ["docs/final-revision-p1-audit-20260623.md", "docs/final-revision-p2-audit-20260624.md"]),
    ("Reproduced final baseline", "Verification used a fresh materialization of the authoritative commit, dependency resolution without version upgrades, static analysis, the complete automated suite, and platform build attempts. Commands, environment, timestamps, and exits are preserved as raw evidence.", ["flutter_analyze_1.3.11+28.log", "flutter_test_1.3.11+28.log", "verification_environment.json"]),
    ("Final-version automated evidence", "The final automated test run records 135 tests passed with zero failures. This claim is limited to the captured test suite and environment; it does not substitute for physical-device, field, usability, or acceptance evidence.", ["flutter_test_1.3.11+28.log", "metadata_validation.json"]),
    ("Reference verification boundary", "Algorithm cases in the source test suite verify scoring and boundary behavior against encoded reference expectations. The term validation here means software/algorithm verification only; external validity, clinical effectiveness, and participant outcomes require researcher evidence.", ["test/ergo_calculator_test.dart", "test/assessment_readiness_test.dart"]),
    ("Automated coverage", "Unit, widget, export, model, persistence, and workflow tests form the regression baseline. Their pass status is tied to the final commit and reproduced log; external systems and physical sensors remain outside this automated evidence.", ["test/", "integration_test/", "flutter_test_1.3.11+28.log"]),
    ("Mixed evidence", "Shared Flutter logic supports Android and iOS, and technical platform builds were produced. Production signing/store ownership is unverified, while current physical-device camera, gallery, TTS, share sheet, and model-inference checks remain Pending Owner/Researcher Evidence.", ["build_android_1.3.11+28.log", "build_ios_1.3.11+28.log", "docs/reviews/local-platform-ml-review-2026-07-29.md"]),
    ("No unsupported benchmark", "No approved final-version performance benchmark with acceptance thresholds was found. Automated execution and successful technical builds demonstrate operability in the captured environment only; startup, inference latency, memory, battery, and export timing need a device protocol and owner-approved criteria.", ["verification_environment.json", "flutter_test_1.3.11+28.log"]),
    ("Human-owned field evidence", "Repository UAT/field records remain historical at their recorded versions. Final 1.3.11+28 participant observations, device matrix, SUS responses, research interpretation, and sign-off must be supplied by authorized researchers and cannot be fabricated from software tests.", ["docs/uat-last-phase-20260712.md", "docs/uat-production-platform-parity-20260719.md", "08_UAT_Field_Test_and_Usability_Package.xlsx"]),
    ("Corrective evidence", "Versioned audit notes and regression tests document fixes and retests. A final defect register must continue to distinguish reproduced final-baseline results from historical findings, and any newly discovered defect requires a failing regression test before correction.", ["docs/final-revision-p1-audit-20260623.md", "test/"]),
    ("Frozen technical baseline", "The deliverable baseline is version 1.3.11+28 at the authoritative commit and tree recorded above. It includes bilingual UI, local profile/history, media assessment, REBA/ISO, advisory ML, recommendations, economic-impact context, TTS, and explicit CSV export; store-ready signing remains unverified.", ["pubspec.yaml", "README.md", "final_source_metadata.txt"]),
    ("Privacy boundary", "Assessment records and media are stored locally unless the user explicitly exports/shares. Firebase Analytics/Crashlytics is present as optional telemetry but defaults off through SOOKTA_TELEMETRY_ENABLED; no credentials, participant media, or signing material are included in this package.", ["lib/core/services/firebase_telemetry_service.dart", "README.md", "03_API_Applicability_Statement.docx"]),
    ("Implemented responses", "The architecture isolates local inference, deterministic scoring, optional advisory output, persistence, export, and fail-safe Firebase initialization. This reduces coupling and preserves assessment operation when telemetry is unavailable; historical issues are not silently rewritten as final claims.", ["lib/main.dart", "lib/core/services/", "docs/reviews/local-platform-ml-review-2026-07-29.md"]),
    ("Known evidence limits", "Production signing and store acceptance, current physical-device compatibility, performance thresholds, final UAT/SUS, researcher interpretation, license conclusions, contractual scope approval, and final acceptance remain human-owned. Local preferences are not a server-grade multi-user database.", ["02_Repository_Access_Checklist.xlsx", "03_API_Applicability_Statement.docx"]),
    ("Maintainable boundaries", "Feature code is separated into app state, models, services, screens, widgets, assets, and platform folders. Maintainability depends on keeping tests current, pinning dependencies, reviewing local schema changes, preserving model provenance, and maintaining privacy-safe export behavior.", ["lib/app/", "lib/core/", "lib/screens/", "test/"]),
    ("Recommended future work", "Before adding cloud sync or remote APIs, define authentication, authorization, consent, retention, encryption, API versioning, audit logging, breach handling, migration, and updated store disclosures. Also reproduce device performance/UAT evidence and complete dependency license review.", ["README.md", "03_API_Applicability_Statement.docx"]),
    ("Pending authorized acceptance", "Technically derivable architecture artifacts are prepared against the pinned baseline. Final acceptance remains Pending Owner Action and Pending Researcher Evidence until repository/account transfers, signing evidence, legal/license review, physical-device/UAT evidence, contractual exceptions, and authorized signatures are completed.", ["12_Final_Developer_Statement_and_Signoff.docx", "02_Repository_Access_Checklist.xlsx"]),
]


PENDING_REFERENCES = {
    "08_UAT_Field_Test_and_Usability_Package.xlsx",
    "12_Final_Developer_Statement_and_Signoff.docx",
}
TASK4_COMPANION_REFERENCES = {
    "03_system_context.drawio", "03_user_navigation.drawio",
    "03_assessment_algorithm.drawio", "03_API_Applicability_Statement.docx",
}


def resolve_citation(reference: str, facts: dict) -> dict:
    if reference in PENDING_REFERENCES:
        return {"path": reference, "status": "Pending future artifact", "resolved_path": None}
    source_root = Path(facts.get("source_root", ""))
    staging_root = Path(facts.get("staging_root", STAGING))
    candidates = []
    path = Path(reference)
    if path.is_absolute():
        candidates.append(path)
    else:
        candidates.extend([
            source_root / reference,
            staging_root / reference,
            staging_root / "artifacts" / reference,
            staging_root / "artifacts" / "diagrams" / reference,
            staging_root / "evidence" / reference,
        ])
    for candidate in candidates:
        if candidate.exists():
            return {"path": reference, "status": "Resolved", "resolved_path": str(candidate)}
    if reference in TASK4_COMPANION_REFERENCES:
        return {"path": reference, "status": "Generated companion artifact", "resolved_path": None}
    return {"path": reference, "status": "Missing", "resolved_path": None}


def report_sections(facts: dict) -> list[dict]:
    result = []
    for index, (heading, label, body, sources) in enumerate((REPORT_HEADINGS[i], *SECTION_DATA[i]) for i in range(28)):
        records = [resolve_citation(source, facts) for source in sources]
        missing = [record["path"] for record in records if record["status"] == "Missing"]
        if missing:
            raise ValueError(f"Unresolved report citations: {missing}")
        source_text = "; ".join(f"{record['path']} [{record['status']}]" for record in records)
        inspected = ""
        if index == 0:
            inspected = (
                f" Inspected facts: {facts['scanned_dart_files']} Dart files; "
                f"assessment API status {facts['api_status']}; optional telemetry default-off="
                f"{str(facts['telemetry_opt_in_default_off']).lower()}."
            )
        grounded = f"{label}. {body}{inspected} Baseline: {VERSION}; authoritative commit {COMMIT}. Evidence: {source_text}."
        result.append({"heading": heading, "body": grounded, "sources": sources, "citation_records": records})
    return result


def diagram_specs() -> list[dict]:
    default_positions = [(100,190),(510,190),(920,190),(100,430),(510,430),(920,430),(310,650),(720,650)]
    def spec(basename, title, nodes, edges, sources, *, external_nodes=None, positions=None, boundary=(40,110,1320,660), box_size=(300,100)):
        labeled_edges = [(a, b, "") if len(edge) == 2 else edge for edge in edges for a, b in [edge[:2]]]
        return {"basename": basename, "title": title, "nodes": nodes, "edges": labeled_edges, "sources": sources, "local_boundary": True, "external_nodes": external_nodes or [], "positions": positions or default_positions, "boundary": boundary, "box_size": box_size}
    return [
        spec("03_system_context", "System Context", ["Field user / researcher", "SookTa mobile app", "Device camera/gallery", "Local app storage", "Explicit CSV share", "Optional Firebase telemetry (external)"], [(0,1,"uses"),(2,1,"media input"),(1,3,"local persistence"),(1,4,"user export"),(1,5,"opt-in telemetry only")], ["lib/main.dart","lib/app/sookta_app.dart"], external_nodes=[0,2,5], positions=[(30,270),(460,330),(30,470),(880,250),(880,450),(880,650)], boundary=(430,180,800,410)),
        spec("03_runtime_data_flow", "Runtime Data Flow", ["Profile + activity", "Media + task inputs", "Local pose inference", "REBA / ISO", "Recommendations + impact", "History", "Explicit export"], [(0,1),(1,2),(2,3),(3,4),(4,5),(5,6)], ["lib/screens/main/","lib/core/services/"]),
        spec("03_image_processing_flow", "Image Processing Flow", ["Camera / gallery / video", "Decode media", "Letterbox 256x256", "MoveNet Thunder", "17 keypoints + confidence", "Pose-to-input mapping", "Unavailable-result handling"], [(0,1),(1,2),(2,3),(3,4),(4,5),(1,6),(4,6)], ["lib/core/services/pose_estimation_service.dart","lib/core/services/pose_image_preprocessor.dart"]),
        spec("03_assessment_algorithm", "Assessment Algorithm", ["MoveNet 51 joint features", "REBA tables A/B/C", "Manual + activity inputs", "ISO lifting / push-pull when applicable", "Higher deterministic risk", "Deterministic score + risk tier", "XGBoost advisory inference", "Final result + advisory"], [(0,1,"pose-derived inputs"),(2,1,"manual inputs"),(2,3,"applicable task inputs"),(1,4,"REBA result"),(3,4,"ISO result"),(4,5,"primary result"),(0,6,"raw joint features"),(5,7,"attach primary"),(6,7,"attach advisory")], ["lib/core/services/ergo_calculator.dart","lib/core/ergonomics_risk_prediction/data/feature_extraction/movenet_joint_feature_extractor.dart","lib/core/services/xgboost_advisory_service.dart"], positions=[(50,170),(380,170),(380,410),(710,410),(710,170),(1040,170),(50,650),(1040,650)], box_size=(260,100)),
        spec("03_recommendation_flow", "Recommendation Flow", ["Activity", "Risk tier", "Body-area risk", "Source-key catalog", "Deduplicate + category cap", "Bilingual recommendations", "User selects actions"], [(0,3),(1,3),(2,3),(3,4),(4,5),(5,6)], ["lib/core/services/risk_recommendation_service.dart","lib/core/models/assessment_reference_sources.dart"]),
        spec("03_local_storage_and_export", "Local Storage and Data Export", ["Profiles / active profile", "Drafts + history", "SharedPreferences", "Temporary / captured image", "Persisted image in application documents", "Assessment record", "CSV in application documents", "OS share sheet (user action)"], [(0,2,"serialize"),(1,2,"serialize"),(3,4,"copy file"),(5,6,"write CSV"),(6,7,"user-selected share")], ["lib/app/app_state.dart","lib/core/services/local_image_store.dart","lib/core/services/assessment_export_service.dart"], positions=[(70,170),(70,410),(390,290),(70,650),(390,650),(720,410),(720,650),(1040,650)]),
        spec("03_user_navigation", "User Navigation", ["Splash / language", "Setup / farmer profile", "Home", "Assessment menu", "Camera / form", "Results + recommendations", "History", "Profile / export / help"], [(0,1),(1,2),(2,3),(3,4),(4,5),(5,6),(2,7),(6,7)], ["lib/screens/onboarding/","lib/screens/main/"]),
        spec("03_module_dependencies", "Module and Dependency View", ["Presentation screens", "App state", "Domain models", "Assessment services", "Media + ML services", "Persistence + export", "Platform plugins", "Optional telemetry"], [(0,1),(0,2),(0,3),(3,4),(1,5),(5,6),(4,6),(0,7)], ["lib/app/","lib/core/","lib/screens/","pubspec.yaml"]),
    ]


def drawio_xml(spec: dict) -> str:
    width, height = 1400, 900
    cells = ['<mxCell id="0"/>', '<mxCell id="1" parent="0"/>']
    bx, by, bw, bh = spec["boundary"]
    cells.append(f'<mxCell id="boundary" value="LOCAL / OFFLINE ASSESSMENT BOUNDARY" style="rounded=1;dashed=1;strokeColor=#4A86E8;fillColor=none;fontColor=#174EA6;fontStyle=1;verticalAlign=top;spacingTop=8;" vertex="1" parent="1"><mxGeometry x="{bx}" y="{by}" width="{bw}" height="{bh}" as="geometry"/></mxCell>')
    coords = spec["positions"]; box_width,box_height=spec["box_size"]
    for i, label in enumerate(spec["nodes"]):
        x,y=coords[i]
        external = i in spec["external_nodes"]
        fill, stroke = ("#FCE8E6", "#D93025") if external else ("#E8F0FE", "#4A86E8")
        cells.append(f'<mxCell id="n{i}" value={quoteattr(label)} external={quoteattr(str(external).lower())} style="rounded=1;whiteSpace=wrap;html=1;fillColor={fill};strokeColor={stroke};fontColor=#202124;fontSize=15;spacing=10;" vertex="1" parent="1"><mxGeometry x="{x}" y="{y}" width="{box_width}" height="{box_height}" as="geometry"/></mxCell>')
    for i,(a,b,label) in enumerate(spec["edges"]):
        cells.append(f'<mxCell id="e{i}" value={quoteattr(label)} style="edgeStyle=orthogonalEdgeStyle;rounded=0;orthogonalLoop=1;jettySize=auto;html=1;endArrow=block;endFill=1;strokeWidth=2;strokeColor=#5F6368;fontSize=11;labelBackgroundColor=#FFFFFF;" edge="1" parent="1" source="n{a}" target="n{b}"><mxGeometry relative="1" as="geometry"/></mxCell>')
    sources = "; ".join(spec["sources"])
    cells.append(f'<mxCell id="legend" value={quoteattr("Legend: blue = local component; red = external actor/service; arrows = labeled data/control flow. Sources: " + sources)} style="rounded=1;whiteSpace=wrap;html=1;fillColor=#F8F9FA;strokeColor=#DADCE0;fontColor=#5F6368;fontSize=11;spacing=8;" vertex="1" parent="1"><mxGeometry x="40" y="790" width="1320" height="70" as="geometry"/></mxCell>')
    return f'''<?xml version="1.0" encoding="UTF-8"?>\n<mxfile host="app.diagrams.net" modified="2026-08-24T00:00:00.000Z" agent="Codex" version="24.7.17" type="device" applicationVersion="{VERSION}" authoritativeCommit="{COMMIT}" sourcePaths={quoteattr(sources)}><diagram id={quoteattr(spec['basename'])} name="Page-1"><mxGraphModel dx="1400" dy="900" grid="1" gridSize="10" guides="1" tooltips="1" connect="1" arrows="1" fold="1" page="1" pageScale="1" pageWidth="1400" pageHeight="900" math="0" shadow="0"><root>{''.join(cells)}</root></mxGraphModel></diagram></mxfile>\n'''


def _font(size: int, bold: bool = False):
    names = ["/System/Library/Fonts/Supplemental/Arial Bold.ttf" if bold else "/System/Library/Fonts/Supplemental/Arial.ttf", "/System/Library/Fonts/Helvetica.ttc"]
    for name in names:
        if Path(name).exists(): return ImageFont.truetype(name, size)
    return ImageFont.load_default()


def edge_boundary_points(source, target):
    def center(box): return ((box[0]+box[2])/2, (box[1]+box[3])/2)
    sc, tc = center(source), center(target); dx, dy = tc[0]-sc[0], tc[1]-sc[1]
    if dx == 0 and dy == 0: return (round(sc[0]),round(sc[1])),(round(tc[0]),round(tc[1]))
    def boundary(box, origin, vx, vy):
        half_w=(box[2]-box[0])/2; half_h=(box[3]-box[1])/2
        tx=half_w/abs(vx) if vx else float("inf"); ty=half_h/abs(vy) if vy else float("inf"); t=min(tx,ty)
        return (round(origin[0]+vx*t),round(origin[1]+vy*t))
    return boundary(source,sc,dx,dy), boundary(target,tc,-dx,-dy)


def render_diagram(spec: dict, output: Path) -> None:
    scale=2
    image=Image.new("RGB", (2800,1800), "white")
    d=ImageDraw.Draw(image)
    title_font, body_font, small_font = _font(50,True), _font(29,True), _font(22)
    d.text((80,50), f"{spec['title']} | SookTa {VERSION}", font=title_font, fill="#202124")
    bx,by,bw,bh=spec["boundary"]
    boundary=(bx*2,by*2,(bx+bw)*2,(by+bh)*2)
    d.rounded_rectangle(boundary, radius=28, outline="#4A86E8", width=5)
    d.text((boundary[0]+30,boundary[1]+25), "LOCAL / OFFLINE ASSESSMENT BOUNDARY", font=small_font, fill="#174EA6")
    coords=[(x*2,y*2) for x,y in spec["positions"]]
    box_width,box_height=spec["box_size"]; box_width*=2; box_height=box_height*2-10
    boxes=[(x,y,x+box_width,y+box_height) for x,y in coords[:len(spec["nodes"])]]
    edge_labels=[]
    for a,b,label in spec["edges"]:
        start,end=edge_boundary_points(boxes[a],boxes[b]); d.line((start,end),fill="#5F6368",width=7)
        angle=math.atan2(end[1]-start[1],end[0]-start[0]); length=34
        p1=(end[0]-length*math.cos(angle-0.55),end[1]-length*math.sin(angle-0.55)); p2=(end[0]-length*math.cos(angle+0.55),end[1]-length*math.sin(angle+0.55))
        d.polygon([end,p1,p2],fill="#5F6368")
        if label: edge_labels.append(((start[0]+end[0])//2,(start[1]+end[1])//2,label))
    for i,label in enumerate(spec["nodes"]):
        x,y=coords[i]; box=boxes[i]
        external=i in spec["external_nodes"]
        d.rounded_rectangle(box, radius=24, fill="#FCE8E6" if external else "#E8F0FE", outline="#D93025" if external else "#4A86E8", width=4)
        words=label.split(); lines=[]; line=""
        for word in words:
            trial=(line+" "+word).strip()
            if d.textbbox((0,0),trial,font=body_font)[2] > box_width-80 and line: lines.append(line); line=word
            else: line=trial
        lines.append(line)
        total=len(lines)*38
        for j,text in enumerate(lines):
            tw=d.textbbox((0,0),text,font=body_font)[2]
            d.text((x+(box_width-tw)/2,y+(box_height-total)/2+j*38),text,font=body_font,fill="#202124")
    label_font=_font(18)
    for x,y,label in edge_labels:
        bounds=d.textbbox((0,0),label,font=label_font); label_width=bounds[2]-bounds[0]
        d.rounded_rectangle((x-label_width/2-8,y-15,x+label_width/2+8,y+15),radius=6,fill="white")
        d.text((x-label_width/2,y-12),label,font=label_font,fill="#3C4043")
    legend=f"Legend: blue = local component; red = external actor/service; arrows = labeled data/control flow. Sources: {'; '.join(spec['sources'])} | Git {COMMIT}"
    d.rounded_rectangle((70,1600,2730,1740),radius=18,fill="#F8F9FA",outline="#DADCE0",width=3)
    d.multiline_text((105,1625),legend,font=small_font,fill="#5F6368",spacing=7)
    output.parent.mkdir(parents=True,exist_ok=True); image.save(output,dpi=(300,300),optimize=True)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120):
    tcPr=cell._tc.get_or_add_tcPr(); tcMar=tcPr.first_child_found_in("w:tcMar")
    if tcMar is None: tcMar=OxmlElement("w:tcMar"); tcPr.append(tcMar)
    for m,v in (("top",top),("start",start),("bottom",bottom),("end",end)):
        node=tcMar.find(qn(f"w:{m}"))
        if node is None: node=OxmlElement(f"w:{m}"); tcMar.append(node)
        node.set(qn("w:w"),str(v)); node.set(qn("w:type"),"dxa")


def configure_doc(doc: Document):
    sec=doc.sections[0]; sec.top_margin=sec.bottom_margin=sec.left_margin=sec.right_margin=Inches(1); sec.header_distance=sec.footer_distance=Inches(.492)
    styles=doc.styles
    for name,size,color,before,after in [("Normal",11,"000000",0,8),("Heading 1",20,"000000",20,6),("Heading 2",16,"000000",18,6),("Heading 3",14,"434343",16,4)]:
        st=styles[name]; st.font.name="Arial"; st.font.size=Pt(size); st.font.color.rgb=RGBColor.from_string(color); st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after); st.paragraph_format.line_spacing=1.15
        st._element.get_or_add_rPr().rFonts.set(qn("w:ascii"),"Arial"); st._element.get_or_add_rPr().rFonts.set(qn("w:hAnsi"),"Arial")
    footer=sec.footer.paragraphs[0]; footer.alignment=WD_ALIGN_PARAGRAPH.CENTER
    run=footer.add_run(f"SookTa {VERSION} | Technical handover | {COMMIT[:12]}"); run.font.name="Arial"; run.font.size=Pt(8); run.font.color.rgb=RGBColor(0x55,0x55,0x55)


def add_title(doc: Document, title: str, subtitle: str):
    p=doc.add_paragraph(); p.paragraph_format.space_after=Pt(3)
    r=p.add_run(title); r.font.name="Arial"; r.font.size=Pt(26); r.font.color.rgb=RGBColor(0,0,0); r.bold=False
    p2=doc.add_paragraph(subtitle); p2.style=doc.styles["Normal"]
    p3=doc.add_paragraph(f"Application {VERSION} | Authoritative commit {COMMIT} | Tree {TREE}"); p3.style=doc.styles["Normal"]
    doc.add_paragraph("Document status: Complete - Pending Signature / Human-Owned Evidence where stated.")


def add_evidence_table(doc: Document, rows: list[tuple[str,str]]):
    table=doc.add_table(rows=1,cols=2); table.autofit=False; table.alignment=0
    table.columns[0].width=Inches(1.7); table.columns[1].width=Inches(4.8)
    hdr=table.rows[0].cells; hdr[0].text="Evidence field"; hdr[1].text="Value"
    for cell in hdr:
        for run in cell.paragraphs[0].runs: run.bold=True; run.font.name="Arial"
        set_cell_margins(cell)
    for key,value in rows:
        cells=table.add_row().cells; cells[0].text=key; cells[1].text=value
        for cell in cells: set_cell_margins(cell)
    return table


def build_report(output: Path, facts: dict):
    doc=Document(); configure_doc(doc); add_title(doc,"Final Technical Development Report","Architecture, implementation, verification boundaries, and handover status")
    add_evidence_table(doc,[("Evidence basis",f"Final source {COMMIT}; Task 2 raw logs"),("Interpretation boundary","Software/algorithm verification only; no clinical or external-validity claim"),("Acceptance","Pending authorized owner/researcher actions and signatures")])
    doc.add_page_break()
    for idx,item in enumerate(report_sections(facts),1):
        doc.add_heading(f"{idx}. {item['heading']}",level=1)
        doc.add_paragraph(item["body"])
        p=doc.add_paragraph(); r=p.add_run("Evidence references: "); r.bold=True; p.add_run("; ".join(item["sources"]))
        if idx in {5,7,8,9,11,12}:
            relevant={5:"03_system_context.png",7:"03_user_navigation.png",8:"03_image_processing_flow.png",9:"03_assessment_algorithm.png",11:"03_recommendation_flow.png",12:"03_local_storage_and_export.png"}[idx]
            p=doc.add_paragraph(f"Figure reference: {relevant} (editable Draw.io source paired in diagrams folder).")
        if idx != 28: doc.add_paragraph()
    output.parent.mkdir(parents=True,exist_ok=True); doc.save(output)


def build_api_statement(output: Path, facts: dict):
    doc=Document(); configure_doc(doc); add_title(doc,"API Applicability Statement","Evidence-based backend/API applicability for the final production assessment path")
    doc.add_heading("Decision",level=1)
    doc.add_paragraph("Status: N/A with Rationale for backend endpoint and server API reference documentation. Code inspection found no direct HTTP/Dio/GraphQL/WebSocket client or Firebase database/functions/storage dependency in the application assessment path. The assessment flow remains local/offline.")
    doc.add_heading("Scope of source inspection",level=1)
    add_evidence_table(doc,[("Scanned baseline",f"{VERSION} at {COMMIT}"),("Dart files scanned",str(facts['scanned_dart_files'])),("Direct backend client hits",json.dumps(facts['backend_client_hits'])),("Local persistence","SharedPreferences and app document directories"),("Local inference","TFLite MoveNet and ONNX advisory model")])
    doc.add_heading("Production assessment path",level=1)
    doc.add_paragraph("Profile/activity inputs and media are processed by local Dart services. Pose inference loads bundled model assets; deterministic REBA/ISO produces the primary result; recommendations are generated from source-controlled rules; history and export files are written locally. Evidence: lib/core/services/pose_estimation_service.dart; lib/core/services/ergo_calculator.dart; lib/core/services/risk_recommendation_service.dart; lib/app/app_state.dart; lib/core/services/assessment_export_service.dart.")
    doc.add_heading("Optional Firebase telemetry is separate",level=1)
    doc.add_paragraph("Firebase Core, Analytics, and Crashlytics are declared for optional app-quality telemetry. FirebaseTelemetryService uses SOOKTA_TELEMETRY_ENABLED with defaultValue false, disables collection when not opted in, and the app continues when Firebase initialization fails. Telemetry is not a backend assessment API and does not calculate or store the assessment result.")
    doc.add_heading("N/A rationale and limits",level=1)
    doc.add_paragraph("Endpoint catalog, request/response schema, authentication method, API versioning, server SLA, and remote database ER/migration documents are N/A because no such production assessment backend was found in the pinned source. This statement does not prove that optional telemetry never transmits data when explicitly enabled, and it does not authorize future cloud behavior.")
    doc.add_heading("Evidence required if remote APIs are introduced",level=1)
    for text in ["API inventory with owners, base URLs, environments, versions, and data classification.","OpenAPI or equivalent request/response/error schema, authentication and authorization design.","Consent, privacy disclosure, retention/deletion, encryption, audit logging, and incident response.","Offline/failure/retry behavior, migration/rollback plan, tests, performance limits, and store disclosure updates."]:
        doc.add_paragraph(text,style="List Bullet")
    doc.add_heading("Human actions",level=1)
    doc.add_paragraph("Owner/researcher approval is required before enabling telemetry in production or introducing any remote service. Firebase project ownership/access and final privacy/store disclosure evidence remain Pending Owner Action. Authorized signature and date remain pending.")
    doc.add_paragraph("Owner/Researcher: ____________________  Signature: ____________________  Date: __________")
    output.parent.mkdir(parents=True,exist_ok=True); doc.save(output)


def main():
    parser=argparse.ArgumentParser(); parser.add_argument("--source",type=Path,required=True); parser.add_argument("--staging",type=Path,default=STAGING); args=parser.parse_args()
    facts=inspect_source(args.source)
    if facts["application_version"] != VERSION: raise SystemExit(f"Baseline mismatch: {facts['application_version']}")
    if facts["assessment_backend_api_present"]: raise SystemExit("Direct backend client found; API N/A classification is unsafe")
    artifacts=args.staging/"artifacts"; diagrams=artifacts/"diagrams"; working=args.staging/"working"/"task4"; manifests=args.staging/"manifests"
    for p in (artifacts,diagrams,working,manifests): p.mkdir(parents=True,exist_ok=True)
    build_report(working/"03_Final_Technical_Development_Report.raw.docx",facts)
    build_api_statement(working/"03_API_Applicability_Statement.raw.docx",facts)
    diagram_records=[]
    for spec in diagram_specs():
        drawio=diagrams/f"{spec['basename']}.drawio"; png=diagrams/f"{spec['basename']}.png"
        drawio.write_text(drawio_xml(spec),encoding="utf-8"); render_diagram(spec,png)
        diagram_records.append({"basename":spec["basename"],"title":spec["title"],"drawio":str(drawio),"drawio_sha256":sha(drawio),"png":str(png),"png_sha256":sha(png),"png_dimensions":[2800,1800],"dpi":300,"sources":spec["sources"]})
    payload={"schema_version":1,"baseline":{"version":VERSION,"commit":COMMIT,"tree":TREE},"source_facts":facts,"report_headings":REPORT_HEADINGS,"diagrams":diagram_records,"human_actions":["Owner/researcher confirmation of contract scope and historical milestone approvals","Physical-device compatibility and performance evidence for 1.3.11+28","Final UAT/SUS evidence and research interpretation","Firebase/store ownership and production privacy disclosure approval","Authorized acceptance and signatures"]}
    (working/"task4_build_input_summary.json").write_text(json.dumps(payload,ensure_ascii=False,indent=2)+"\n",encoding="utf-8")
    print(json.dumps({"status":"built_raw","diagrams":len(diagram_records),"report_sections":28,"facts":facts},indent=2))


if __name__ == "__main__": main()
