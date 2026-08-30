#!/usr/bin/env python3
import json
import os
import shutil
import subprocess
import tempfile
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT=Path('/private/tmp/fsookta-final-handover')
HERE=ROOT/'working/task6'
PY='/Users/kpc/.cache/codex-runtimes/codex-primary-runtime/dependencies/python/bin/python3'
NODE='/Users/kpc/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/bin/node'
NODE_PATH='/Users/kpc/.cache/codex-runtimes/codex-primary-runtime/dependencies/node/node_modules'

class CanonicalEntrypointTest(unittest.TestCase):
 def test_clean_staging_entrypoints_rebuild_from_source_truth(self):
  with tempfile.TemporaryDirectory() as td:
   root=Path(td); tools=root/'tools'; tools.mkdir(); (root/'artifacts').mkdir(); (root/'working/task6').mkdir(parents=True)
   for name in ('task6_source.py','build_task6_audit_test_uat.py','build_task6_documents.py','build_task6_workbooks.mjs','build_task6_canonical_workbooks.mjs'):
    shutil.copy2(HERE/name,tools/name)
   (tools/'node_modules').symlink_to(NODE_PATH,target_is_directory=True)
   for name in ('evidence','authoritative-materializations'):
    (root/name).symlink_to(ROOT/name,target_is_directory=True)
   shutil.copy2(ROOT/'evidence_map.json',root/'evidence_map.json')
   shutil.copy2(ROOT/'working/task6/task6_build_input.json',root/'working/task6/task6_build_input.json')
   env={**os.environ,'FSOOKTA_HANDOVER_ROOT':str(root),'NODE_PATH':NODE_PATH}
   subprocess.run([PY,'build_task6_audit_test_uat.py'],cwd=tools,env=env,check=True,capture_output=True,text=True)
   subprocess.run([NODE,'build_task6_workbooks.mjs'],cwd=tools,env=env,check=True,capture_output=True,text=True)
   payload=json.loads((root/'working/task6/task6_corrected_payload.json').read_text())
   self.assertEqual(135,len(payload['tests'])); self.assertEqual(55,sum(x['category']=='Algorithm / reference' for x in payload['tests']))
   master=load_workbook(root/'artifacts/07_Master_Test_and_Verification_Package.xlsx',data_only=False)
   self.assertEqual(16,len(master.sheetnames)); self.assertEqual(55,master['Algorithm Reference'].max_row-4); self.assertEqual(7,master['Threshold Boundaries'].max_row-4)
   self.assertTrue((root/'artifacts/07_Final_Algorithm_Verification_Report.docx').is_file()); self.assertTrue((root/'artifacts/08_UAT_and_Field_Test_Technical_Report.docx').is_file())

if __name__=='__main__': unittest.main()
