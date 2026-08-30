#!/usr/bin/env python3
"""Verify Task 3 archive, DOCX, XLSX, PDF, and checksum deliverables."""

from __future__ import annotations

import argparse
import hashlib
import io
import json
import posixpath
import re
import subprocess
import tarfile
import tempfile
import zipfile
from datetime import datetime
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from pypdf import PdfReader

from build_task3_release import COMMIT, ROOT_NAME, TREE, VERSION


STAGING = Path("/private/tmp/fsookta-final-handover")
NS_W = "http://schemas.openxmlformats.org/wordprocessingml/2006/main"
NS_CP = "http://schemas.openxmlformats.org/package/2006/metadata/core-properties"
NS_DC = "http://purl.org/dc/elements/1.1/"
NS_X = "http://schemas.openxmlformats.org/spreadsheetml/2006/main"
W = f"{{{NS_W}}}"
X = f"{{{NS_X}}}"
NS_R = "http://schemas.openxmlformats.org/officeDocument/2006/relationships"
NS_PR = "http://schemas.openxmlformats.org/package/2006/relationships"

VERIFY_DENIED_SEGMENTS = {
    ".dart_tool", ".gradle", ".idea", ".pub-cache", ".vscode", ".worktrees",
    "build", "cache", "caches", "coverage", "deriveddata", "node_modules",
    "outputs", "pods", "tmp", "vendor",
}
VERIFY_DENIED_SUFFIXES = {
    ".jks", ".keystore", ".key", ".mobileprovision", ".p8", ".p12", ".pem", ".pfx",
}
VERIFY_SENSITIVE_NAMES = {
    ".env", ".npmrc", ".pypirc", "credentials.json", "firebase_app_id_file.json",
    "firebase_options.dart", "google-services.json", "googleservice-info.plist",
    "key.properties", "secrets.json", "service-account.json", "service_account.json",
}
TEXT_SCAN_SUFFIXES = {
    "", ".cfg", ".conf", ".dart", ".env", ".gradle", ".json", ".md", ".plist",
    ".properties", ".sh", ".swift", ".txt", ".xml", ".yaml", ".yml",
}


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def require(condition: bool, message: str) -> None:
    if not condition:
        raise AssertionError(message)


def git(repo: Path, *args: str, text: bool = True) -> str | bytes:
    return subprocess.check_output(["git", "-C", str(repo), *args], text=text)


def independent_archive_path_violation(path: str) -> str | None:
    """Return a verifier-owned deny reason without calling the builder predicate."""
    normalized = path.replace("\\", "/")
    lower = normalized.lower()
    parts = [part for part in lower.split("/") if part]
    name = parts[-1] if parts else ""
    if set(parts) & VERIFY_DENIED_SEGMENTS:
        return "denied generated/cache/vendor directory segment"
    if name in VERIFY_SENSITIVE_NAMES or name == "local.properties" or name.startswith(".env."):
        return "sensitive configuration filename"
    if Path(name).suffix.lower() in VERIFY_DENIED_SUFFIXES:
        return "credential/signing filename suffix"
    if re.search(r"(^|/)(secrets?|credentials?|tokens?)(/|\.|$)", lower):
        return "secret/credential/token path"
    if lower.startswith(("data/research/", "docs/qa/", "docs/uat_evidence_", "docs/user_manual_v1_1_1_android/")):
        return "research/QA media path"
    if name in {
        "sookta_reba_logic_evidence_package_20260603.zip",
        "sookta_research_training_dataset_reba_iso11228.xlsx",
    }:
        return "excluded research evidence package"
    return None


def secret_content_violation(path: str, content: bytes) -> str | None:
    """High-confidence independent content scan for text-like archived files."""
    suffix = Path(path).suffix.lower()
    if suffix not in TEXT_SCAN_SUFFIXES or len(content) > 2_000_000:
        return None
    text = content.decode("utf-8", errors="ignore")
    patterns = [
        (r"-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----", "private-key block"),
        (r"AIza[0-9A-Za-z_-]{25,}", "Google API key"),
        (r"AKIA[0-9A-Z]{16}", "AWS access key"),
        (r"gh[pousr]_[0-9A-Za-z]{30,}", "GitHub token"),
        (r'"(?:client_secret|private_key|api_key)"\s*:\s*"(?!REDACTED|EXAMPLE|PLACEHOLDER)[^"\r\n]{12,}"', "secret-bearing JSON value"),
        (r"(?im)^\s*(?:password|token|secret)\s*[=:]\s*(?!REDACTED|EXAMPLE|PLACEHOLDER)\S{12,}\s*$", "secret-bearing assignment"),
    ]
    for pattern, label in patterns:
        if re.search(pattern, text):
            return label
    return None


def verify_archive(repo: Path, staging: Path) -> dict:
    archive = staging / "archives" / f"SookTa-{VERSION}-source-snapshot.tar.gz"
    manifest_path = staging / "manifests" / "source_snapshot_manifest.json"
    manifest = json.loads(manifest_path.read_text(encoding="utf-8"))
    require(manifest["application_version"] == VERSION, "archive manifest version mismatch")
    require(manifest["authoritative_commit"] == COMMIT, "archive manifest commit mismatch")
    require(manifest["git_tree_id"] == TREE, "archive manifest tree mismatch")
    require(git(repo, "rev-parse", f"{COMMIT}^{{tree}}").strip() == TREE, "repository tree mismatch")
    require(sha256(archive) == manifest["sha256"], "archive digest mismatch")
    require(archive.stat().st_size == manifest["byte_size"], "archive byte-size mismatch")

    with tarfile.open(archive, "r:gz") as source:
        members = source.getmembers()
        files = [m for m in members if m.isfile()]
        require(len(files) == manifest["file_count"], "archive file-count mismatch")
        paths: list[str] = []
        for member in files:
            require(member.name.startswith(f"{ROOT_NAME}/"), f"unexpected archive root: {member.name}")
            relative = member.name[len(ROOT_NAME) + 1 :]
            require(relative and ".." not in Path(relative).parts, f"unsafe archive path: {member.name}")
            violation = independent_archive_path_violation(relative)
            require(violation is None, f"independent exclusion failure ({violation}): {relative}")
            require(member.mtime == 1787443200, f"non-deterministic mtime: {relative}")
            require(member.uid == 0 and member.gid == 0, f"non-deterministic ownership: {relative}")
            stream = source.extractfile(member)
            require(stream is not None, f"unable to scan archived file: {relative}")
            content_violation = secret_content_violation(relative, stream.read())
            require(content_violation is None, f"secret-like archive content ({content_violation}): {relative}")
            paths.append(relative)
        require(paths == manifest["paths"], "archive member order/path manifest mismatch")
        require(paths == sorted(paths), "archive paths are not sorted")
        required_absent = {
            "android/app/google-services.json",
            "ios/Runner/GoogleService-Info.plist",
            "lib/firebase_options.dart",
            "docs/Sookta_REBA_Logic_Evidence_Package_20260603.zip",
            "docs/Sookta_Research_Training_Dataset_REBA_ISO11228.xlsx",
        }
        require(required_absent.isdisjoint(paths), "secret/research exclusion failure")
        pubspec = source.extractfile(f"{ROOT_NAME}/pubspec.yaml")
        require(pubspec is not None and f"version: {VERSION}" in pubspec.read().decode("utf-8"), "archived version mismatch")

    baseline_tar = tarfile.open(fileobj=io.BytesIO(git(repo, "archive", "--format=tar", COMMIT, text=False)), mode="r:")
    baseline_hashes = {}
    for member in baseline_tar.getmembers():
        if member.isfile() and member.name in paths:
            stream = baseline_tar.extractfile(member)
            require(stream is not None, f"unable to read baseline member {member.name}")
            baseline_hashes[member.name] = hashlib.sha256(stream.read()).hexdigest()
    require(set(baseline_hashes) == set(paths), "archive/baseline path reconciliation failed")
    with tempfile.TemporaryDirectory(prefix="fsookta-task3-unpack-", dir="/private/tmp") as temp:
        temp_path = Path(temp)
        with tarfile.open(archive, "r:gz") as source:
            source.extractall(temp_path, filter="data")
        unpacked = sorted(p.relative_to(temp_path / ROOT_NAME).as_posix() for p in (temp_path / ROOT_NAME).rglob("*") if p.is_file())
        require(unpacked == paths, "unpacked path reconciliation failed")
        for relative in unpacked:
            require(sha256(temp_path / ROOT_NAME / relative) == baseline_hashes[relative], f"content mismatch: {relative}")
    return {
        "path": str(archive),
        "sha256": manifest["sha256"],
        "byte_size": manifest["byte_size"],
        "file_count": manifest["file_count"],
        "version": VERSION,
        "commit": COMMIT,
        "tree": TREE,
        "exclusions": "passed",
        "unpack_reconciliation": "365/365 files match authoritative commit" if manifest["file_count"] == 365 else f"{manifest['file_count']}/{manifest['file_count']} files match authoritative commit",
    }


def attr(node: ET.Element | None, name: str) -> str | None:
    return None if node is None else node.get(f"{W}{name}")


def title_residue_violations(document_xml: ET.Element, styles_xml: ET.Element | None = None) -> list[str]:
    violations: list[str] = []
    title = next(
        (
            paragraph
            for paragraph in document_xml.findall(f".//{W}p")
            if "".join(node.text or "" for node in paragraph.findall(f".//{W}t")).strip()
        ),
        None,
    )
    if title is None:
        return ["missing title paragraph"]
    if title.find(f"{W}pPr/{W}pBdr") is not None:
        violations.append("title paragraph border")
    for underline in title.findall(f".//{W}u"):
        if underline.get(f"{W}val", "single") not in {"none", "0", "false"}:
            violations.append("title underline/rule residue")
            break
    style_ref = title.find(f"{W}pPr/{W}pStyle")
    if styles_xml is not None and style_ref is not None:
        style_id = style_ref.get(f"{W}val")
        style = styles_xml.find(f".//{W}style[@{W}styleId='{style_id}']") if style_id else None
        if style is not None and style.find(f"{W}pPr/{W}pBdr") is not None:
            violations.append("title style paragraph border")
        if style is not None:
            for underline in style.findall(f".//{W}u"):
                if underline.get(f"{W}val", "single") not in {"none", "0", "false"}:
                    violations.append("title style underline/rule residue")
                    break
    return violations


def verify_docx(path: Path) -> dict:
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        document_xml = ET.fromstring(package.read("word/document.xml"))
        styles_xml = ET.fromstring(package.read("word/styles.xml"))
        core_xml = ET.fromstring(package.read("docProps/core.xml"))
        require("docProps/custom.xml" not in names, f"custom properties present: {path.name}")
        creators = core_xml.findall(f".//{{{NS_DC}}}creator") + core_xml.findall(f".//{{{NS_CP}}}lastModifiedBy")
        require(all(not (node.text or "").strip() for node in creators), f"author metadata not scrubbed: {path.name}")
        residue = title_residue_violations(document_xml, styles_xml)
        require(not residue, f"title border/underline residue: {path.name}: {residue}")

    doc = Document(path)
    first = next(p for p in doc.paragraphs if p.text.strip())
    require(first.style.name == "Normal", f"title uses a named title style: {path.name}")
    require(first.text in {"Final Release and Scope Closure", "Repository Ownership and IP Handover"}, f"unexpected title: {path.name}")
    require(first.runs and first.runs[0].font.name == "Arial", f"title font mismatch: {path.name}")
    require(first.runs[0].font.size and round(first.runs[0].font.size.pt) == 26, f"title size mismatch: {path.name}")
    require(first.runs[0].underline is not True, f"underlined title: {path.name}")

    sect = document_xml.find(f".//{W}sectPr")
    require(sect is not None, f"missing section properties: {path.name}")
    page = sect.find(f"{W}pgSz")
    margins = sect.find(f"{W}pgMar")
    require(attr(page, "w") == "12240" and attr(page, "h") == "15840", f"page preset mismatch: {path.name}")
    for edge in ("top", "right", "bottom", "left"):
        require(attr(margins, edge) == "1440", f"margin preset mismatch ({edge}): {path.name}")
    require(attr(margins, "header") == "708" and attr(margins, "footer") == "708", f"header/footer preset mismatch: {path.name}")

    normal = styles_xml.find(f".//{W}style[@{W}styleId='Normal']")
    require(normal is not None, f"missing Normal style: {path.name}")
    fonts = normal.find(f".//{W}rFonts")
    size = normal.find(f".//{W}sz")
    spacing = normal.find(f".//{W}spacing")
    require(attr(fonts, "ascii") == "Arial" and attr(fonts, "hAnsi") == "Arial", f"Normal font preset mismatch: {path.name}")
    require(attr(size, "val") == "22", f"Normal size preset mismatch: {path.name}")
    require(attr(spacing, "after") == "160" and attr(spacing, "line") == "276", f"Normal spacing preset mismatch: {path.name}")

    headings = document_xml.findall(f".//{W}pStyle[@{W}val='Heading1']")
    require(len(headings) >= 5, f"insufficient heading structure: {path.name}")
    list_paragraphs = [
        paragraph
        for paragraph in document_xml.findall(f".//{W}p")
        if paragraph.find(f"{W}pPr/{W}numPr") is not None
    ]
    require(len(list_paragraphs) >= 3, f"true OOXML lists not found: {path.name}")
    literal_list_markers = [p.text for p in doc.paragraphs if re.match(r"^\s*(?:[•●]|-\s)", p.text)]
    require(not literal_list_markers, f"literal fake list markers present: {path.name}")

    table_count = 0
    for table in document_xml.findall(f".//{W}tbl"):
        table_count += 1
        props = table.find(f"{W}tblPr")
        grid = table.find(f"{W}tblGrid")
        require(props is not None and grid is not None, f"missing table geometry: {path.name}")
        table_width = props.find(f"{W}tblW")
        table_indent = props.find(f"{W}tblInd")
        grid_widths = [int(attr(col, "w") or "0") for col in grid.findall(f"{W}gridCol")]
        require(attr(table_width, "w") == "9360" and attr(table_width, "type") == "dxa", f"table width mismatch: {path.name}")
        require(attr(table_indent, "w") == "0", f"table indent mismatch: {path.name}")
        require(sum(grid_widths) == 9360, f"table grid sum mismatch: {path.name}")
        for row in table.findall(f"{W}tr"):
            cells = row.findall(f"{W}tc")
            require(len(cells) == len(grid_widths), f"table cell/grid mismatch: {path.name}")
            for cell, expected in zip(cells, grid_widths):
                cell_width = cell.find(f"{W}tcPr/{W}tcW")
                require(attr(cell_width, "w") == str(expected), f"cell width mismatch: {path.name}")
    require(table_count >= 2, f"expected tables missing: {path.name}")

    all_text = "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(cell.text for table in doc.tables for row in table.rows for cell in row.cells)
    require("Pending Owner Action" in all_text or "human action required" in all_text.lower(), f"human-action status missing: {path.name}")
    require("PRIVATE KEY-----" not in all_text and not re.search(r"AIza[0-9A-Za-z_-]{25,}", all_text), f"secret-like value present: {path.name}")
    return {
        "path": str(path),
        "sha256": sha256(path),
        "byte_size": path.stat().st_size,
        "paragraphs": len(doc.paragraphs),
        "tables": table_count,
        "headings": len(headings),
        "true_list_paragraphs": len(list_paragraphs),
        "preset": "google_docs_default passed",
        "table_geometry": "passed",
        "privacy": "passed",
        "title_sanitizer": "passed",
    }


def workbook_sheet_names(package: zipfile.ZipFile) -> list[str]:
    workbook = ET.fromstring(package.read("xl/workbook.xml"))
    return [node.get("name") or "" for node in workbook.findall(f".//{X}sheet")]


def workbook_sheet_targets(package: zipfile.ZipFile) -> dict[str, str]:
    workbook = ET.fromstring(package.read("xl/workbook.xml"))
    relationships = ET.fromstring(package.read("xl/_rels/workbook.xml.rels"))
    targets = {
        node.get("Id"): node.get("Target")
        for node in relationships.findall(f".//{{{NS_PR}}}Relationship")
    }
    result: dict[str, str] = {}
    for sheet in workbook.findall(f".//{X}sheet"):
        relationship_id = sheet.get(f"{{{NS_R}}}id")
        target = targets.get(relationship_id)
        require(bool(target), f"missing worksheet relationship for {sheet.get('name')}")
        if target.startswith("/"):
            package_target = posixpath.normpath(target).lstrip("/")
        else:
            package_target = posixpath.normpath(posixpath.join("xl", target))
        result[sheet.get("name") or ""] = package_target
    return result


def verify_xlsx_formula_expectations(path: Path, expected: dict[str, dict[str, dict[str, str]]]) -> dict:
    checked = 0
    with zipfile.ZipFile(path) as package:
        targets = workbook_sheet_targets(package)
        for sheet_name, cells in expected.items():
            require(sheet_name in targets, f"expected formula sheet missing: {path.name}:{sheet_name}")
            root = ET.fromstring(package.read(targets[sheet_name]))
            by_address = {cell.get("r"): cell for cell in root.findall(f".//{X}c")}
            for address, spec in cells.items():
                cell = by_address.get(address)
                require(cell is not None, f"expected formula cell missing: {path.name}:{sheet_name}!{address}")
                formula = cell.find(f"{X}f")
                value = cell.find(f"{X}v")
                require(formula is not None and formula.text == spec["formula"], f"expected formula mismatch: {path.name}:{sheet_name}!{address}")
                require(value is not None and value.text is not None, f"missing cached formula result: {path.name}:{sheet_name}!{address}")
                require(value.text == str(spec["value"]), f"expected formula result mismatch: {path.name}:{sheet_name}!{address}: {value.text}")
                checked += 1
    return {"checked_cells": checked, "status": "passed"}


def parse_ndjson(value: str) -> list[dict]:
    records = []
    for line in value.splitlines():
        if line.strip():
            record = json.loads(line)
            require(isinstance(record, dict), "artifact-tool inspection record is not an object")
            records.append(record)
    require(bool(records), "artifact-tool inspection output is empty")
    return records


def verify_workbook_build_summary(path: Path) -> dict:
    summary = json.loads(path.read_text(encoding="utf-8"))
    checked = 0
    for key in ("license", "access"):
        inspection = summary[key]["inspection"]
        scan_records = parse_ndjson(inspection["formula_error_scan"])
        matches = [record for record in scan_records if record.get("kind") != "notice" or "matched 0 entries" not in record.get("message", "")]
        require(not matches, f"artifact-tool formula error scan matched cells: {key}: {matches}")
        contract = inspection.get("inspection_contract")
        require(isinstance(contract, dict) and contract.get("status") == "passed", f"missing/pending inspection contract: {key}")
        require(contract.get("formula_error_matches") == 0, f"formula error match count is nonzero: {key}")
        checks = contract.get("checks")
        require(isinstance(checks, list) and checks, f"expected inspection checks missing: {key}")
        for check in checks:
            require(check.get("status") == "passed", f"inspection check did not pass: {key}: {check}")
            require(check.get("expected_values") == check.get("inspected_values"), f"inspection expected-value mismatch: {key}: {check.get('range')}")
            require(check.get("expected_formulas") == check.get("inspected_formulas"), f"inspection expected-formula mismatch: {key}: {check.get('range')}")
            checked += 1
    return {"path": str(path), "sha256": sha256(path), "checks": checked, "status": "passed"}


def verify_visual_qa_manifest(staging: Path, expected: dict[str, list[str]]) -> dict:
    manifest_path = staging / "manifests" / "task3_visual_qa_manifest.json"
    require(manifest_path.is_file(), "visual QA manifest is absent")
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    require(payload.get("schema_version") == 1, "visual QA manifest schema mismatch")
    require(payload.get("status") == "passed", "visual QA manifest status is not passed")
    require(bool(str(payload.get("reviewer_role", "")).strip()), "visual QA reviewer role missing")
    require(bool(str(payload.get("notes", "")).strip()), "visual QA notes missing")
    inspected_at = str(payload.get("inspected_at", ""))
    try:
        datetime.fromisoformat(inspected_at)
    except ValueError as error:
        raise AssertionError("visual QA inspected-at timestamp is invalid") from error
    entries = payload.get("artifacts")
    require(isinstance(entries, list), "visual QA artifact entries missing")
    by_path = {entry.get("path"): entry for entry in entries}
    require(set(by_path) == set(expected), "visual QA artifact set is incomplete or unexpected")
    render_references = 0
    unique_renders: set[str] = set()
    for artifact_relative, expected_renders in expected.items():
        entry = by_path[artifact_relative]
        artifact_path = staging / artifact_relative
        require(artifact_path.is_file(), f"visual QA artifact missing: {artifact_relative}")
        require(entry.get("sha256") == sha256(artifact_path), f"visual QA artifact hash is stale: {artifact_relative}")
        require(entry.get("status") == "passed" and bool(str(entry.get("notes", "")).strip()), f"visual QA artifact review incomplete: {artifact_relative}")
        renders = entry.get("renders")
        require(isinstance(renders, list), f"visual QA render records missing: {artifact_relative}")
        render_by_path = {render.get("path"): render for render in renders}
        require(set(render_by_path) == set(expected_renders), f"visual QA render set incomplete: {artifact_relative}")
        for render_relative in expected_renders:
            render = render_by_path[render_relative]
            render_path = staging / render_relative
            require(render_path.is_file(), f"visual QA render missing: {render_relative}")
            require(render.get("sha256") == sha256(render_path), f"visual QA render hash is stale: {render_relative}")
            require(render.get("status") == "passed" and bool(str(render.get("notes", "")).strip()), f"visual QA render review incomplete: {render_relative}")
            render_references += 1
            unique_renders.add(render_relative)
    return {
        "path": str(manifest_path),
        "sha256": sha256(manifest_path),
        "inspected_at": inspected_at,
        "reviewer_role": payload["reviewer_role"],
        "artifacts": len(expected),
        "render_references": render_references,
        "unique_render_files": len(unique_renders),
        "status": "passed",
    }


def verify_xlsx(
    path: Path,
    expected_sheets: list[str],
    paper_size: str,
    formula_expectations: dict[str, dict[str, dict[str, str]]],
) -> dict:
    with zipfile.ZipFile(path) as package:
        names = set(package.namelist())
        require("docProps/core.xml" not in names and "docProps/custom.xml" not in names, f"workbook metadata present: {path.name}")
        sheets = workbook_sheet_names(package)
        require(sheets == expected_sheets, f"sheet-name mismatch: {path.name}: {sheets}")
        formulas = 0
        cached_errors: list[str] = []
        print_settings = 0
        for name in sorted(n for n in names if re.fullmatch(r"xl/worksheets/sheet\d+\.xml", n)):
            root = ET.fromstring(package.read(name))
            formulas += len(root.findall(f".//{X}f"))
            for cell in root.findall(f".//{X}c"):
                value = cell.find(f"{X}v")
                if value is not None and value.text in {"#REF!", "#DIV/0!", "#VALUE!", "#NAME?", "#N/A"}:
                    cached_errors.append(f"{name}:{cell.get('r')}:{value.text}")
            setup = root.find(f"{X}pageSetup")
            require(setup is not None, f"missing page setup: {path.name}:{name}")
            require(setup.get("orientation") == "landscape" and setup.get("fitToWidth") == "1", f"print fit mismatch: {path.name}:{name}")
            require(setup.get("paperSize") == paper_size, f"paper size mismatch: {path.name}:{name}")
            print_settings += 1
        expected_formula_count = sum(len(cells) for cells in formula_expectations.values())
        require(formulas == expected_formula_count, f"formula count mismatch: {path.name}: {formulas} != {expected_formula_count}")
        require(not cached_errors, f"cached formula errors: {cached_errors}")
        workbook_xml = package.read("xl/workbook.xml").decode("utf-8")
        require("_xlnm.Print_Titles" in workbook_xml, f"continuation-page header missing: {path.name}")
    formula_contract = verify_xlsx_formula_expectations(path, formula_expectations)
    return {
        "path": str(path),
        "sha256": sha256(path),
        "byte_size": path.stat().st_size,
        "sheets": sheets,
        "sheet_count": len(sheets),
        "formula_cells": formulas,
        "formula_error_scan": "0 errors",
        "formula_contract": formula_contract,
        "privacy": "passed",
        "print_settings": f"{print_settings}/{print_settings} sheets landscape fit-to-width",
    }


def verify_pdf(path: Path, expected_pages: int) -> dict:
    reader = PdfReader(path)
    require(len(reader.pages) == expected_pages, f"PDF page-count mismatch: {path.name}")
    require(all(float(page.mediabox.width) > 0 and float(page.mediabox.height) > 0 for page in reader.pages), f"invalid PDF page: {path.name}")
    return {"path": str(path), "sha256": sha256(path), "byte_size": path.stat().st_size, "pages": len(reader.pages)}


def expected_formula_contracts(staging: Path) -> dict[str, dict[str, dict[str, dict[str, str]]]]:
    dependencies = json.loads((staging / "working" / "task3" / "dependencies.json").read_text(encoding="utf-8"))["records"]
    license_end = len(dependencies) + 4
    license_cells: dict[str, dict[str, dict[str, str]]] = {"License Register": {}, "Review Summary": {}}
    for index, record in enumerate(dependencies, start=5):
        license_cells["License Register"][f"L{index}"] = {
            "formula": f'IF(E{index}="Human verification required","Pending legal review","Recorded")',
            "value": "Pending legal review" if record["license_identifier"] == "Human verification required" else "Recorded",
        }
    pending = sum(record["license_identifier"] == "Human verification required" for record in dependencies)
    runtime = sum(record["classification"] == "Runtime" for record in dependencies)
    license_cells["Review Summary"] = {
        "B5": {"formula": f"COUNTA('License Register'!A5:A{license_end})", "value": str(len(dependencies))},
        "B6": {"formula": f'COUNTIF(\'License Register\'!L5:L{license_end},"Pending legal review")', "value": str(pending)},
        "B7": {"formula": f'COUNTIF(\'License Register\'!L5:L{license_end},"Recorded")', "value": str(len(dependencies) - pending)},
        "B8": {"formula": f'COUNTIF(\'License Register\'!H5:H{license_end},"Runtime")', "value": str(runtime)},
    }
    access_cells = {
        "Status Summary": {
            "B5": {"formula": "COUNTA('Access Checklist'!A5:A18)", "value": "14"},
            "B6": {"formula": 'COUNTIF(\'Access Checklist\'!E5:E18,"Pending Owner Action")', "value": "14"},
            "B7": {"formula": 'COUNTIF(\'Access Checklist\'!E5:E18,"Pending Researcher Evidence")', "value": "0"},
            "B8": {"formula": 'COUNTIF(\'Access Checklist\'!E5:E18,"Exception Approval Required")', "value": "0"},
            "B9": {"formula": 'COUNTIF(\'Access Checklist\'!E5:E18,"Completed - Evidence Attached")', "value": "0"},
        }
    }
    return {"license": license_cells, "access": access_cells}


def task3_visual_expectations() -> dict[str, list[str]]:
    release_pages = [f"renders/task3/docx/01_release/page-{page}.png" for page in range(1, 6)]
    repository_pages = [f"renders/task3/docx/02_repository/page-{page}.png" for page in range(1, 6)]
    return {
        "artifacts/01_Final_Release_and_Scope_Closure.docx": release_pages,
        "artifacts/01_Final_Release_and_Scope_Closure.pdf": release_pages,
        "artifacts/02_Repository_Ownership_and_IP_Handover.docx": repository_pages,
        "artifacts/02_Repository_Ownership_and_IP_Handover.pdf": repository_pages,
        "artifacts/02_Third_Party_License_Register.xlsx": [
            "renders/task3/xlsx/02_Third_Party_License_Register__License_Register.png",
            "renders/task3/xlsx/02_Third_Party_License_Register__Review_Summary.png",
            "renders/task3/xlsx/02_Third_Party_License_Register__Sources_Notes.png",
        ],
        "artifacts/02_Third_Party_License_Register.pdf": [
            f"renders/task3/pdf/02_Third_Party_License_Register/page-{page}.png" for page in range(1, 10)
        ],
        "artifacts/02_Repository_Access_Checklist.xlsx": [
            "renders/task3/xlsx/02_Repository_Access_Checklist__Access_Checklist.png",
            "renders/task3/xlsx/02_Repository_Access_Checklist__Status_Summary.png",
            "renders/task3/xlsx/02_Repository_Access_Checklist__Secure_Handover_Notes.png",
        ],
        "artifacts/02_Repository_Access_Checklist.pdf": [
            f"renders/task3/pdf/02_Repository_Access_Checklist/page-{page}.png" for page in range(1, 5)
        ],
    }


def write_checksums(staging: Path, paths: list[Path]) -> Path:
    output = staging / "manifests" / "SHA256SUMS.txt"
    lines = [f"{sha256(path)}  {path.relative_to(staging).as_posix()}" for path in sorted(paths, key=lambda p: p.relative_to(staging).as_posix())]
    output.write_text("\n".join(lines) + "\n", encoding="ascii")
    for line in output.read_text(encoding="ascii").splitlines():
        digest, relative = line.split("  ", 1)
        require(sha256(staging / relative) == digest, f"checksum round-trip failed: {relative}")
    return output


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repository", type=Path, required=True)
    parser.add_argument("--staging", type=Path, default=STAGING)
    parser.add_argument("--summary", type=Path, required=True)
    args = parser.parse_args()
    staging = args.staging
    artifacts = staging / "artifacts"

    archive = verify_archive(args.repository, staging)
    docx_paths = [artifacts / "01_Final_Release_and_Scope_Closure.docx", artifacts / "02_Repository_Ownership_and_IP_Handover.docx"]
    docs = [verify_docx(path) for path in docx_paths]
    formula_contracts = expected_formula_contracts(staging)
    workbooks = [
        verify_xlsx(
            artifacts / "02_Third_Party_License_Register.xlsx",
            ["License Register", "Review Summary", "Sources & Notes"],
            "8",
            formula_contracts["license"],
        ),
        verify_xlsx(
            artifacts / "02_Repository_Access_Checklist.xlsx",
            ["Access Checklist", "Status Summary", "Secure Handover Notes"],
            "9",
            formula_contracts["access"],
        ),
    ]
    workbook_build_summary = verify_workbook_build_summary(staging / "manifests" / "task3_workbook_build_summary.json")
    pdf_expectations = {
        "01_Final_Release_and_Scope_Closure.pdf": 5,
        "02_Repository_Ownership_and_IP_Handover.pdf": 5,
        "02_Third_Party_License_Register.pdf": 9,
        "02_Repository_Access_Checklist.pdf": 4,
    }
    pdfs = [verify_pdf(artifacts / name, pages) for name, pages in pdf_expectations.items()]
    visual_qa = verify_visual_qa_manifest(staging, task3_visual_expectations())
    safe_manifests = [
        staging / "manifests" / "source_snapshot_manifest.json",
        staging / "manifests" / "task3_workbook_build_summary.json",
        staging / "manifests" / "task3_visual_qa_manifest.json",
    ]
    checksum_paths = [Path(archive["path"]), *docx_paths, *(artifacts / name for name in ["02_Third_Party_License_Register.xlsx", "02_Repository_Access_Checklist.xlsx"]), *(artifacts / name for name in pdf_expectations), *safe_manifests]
    checksums = write_checksums(staging, checksum_paths)
    dependencies = json.loads((staging / "working" / "task3" / "dependencies.json").read_text(encoding="utf-8"))["records"]
    pending_license_rows = sum(record["license_identifier"] == "Human verification required" for record in dependencies)
    payload = {
        "schema_version": 1,
        "baseline": {"version": VERSION, "commit": COMMIT, "tree": TREE},
        "archive": archive,
        "docx": docs,
        "xlsx": workbooks,
        "pdf": pdfs,
        "checksum_manifest": {"path": str(checksums), "sha256": sha256(checksums), "entries": len(checksum_paths)},
        "safe_manifest_hashes": {
            "source_snapshot_manifest": sha256(safe_manifests[0]),
            "workbook_build_summary": sha256(safe_manifests[1]),
            "visual_qa_manifest": sha256(safe_manifests[2]),
        },
        "workbook_build_summary": workbook_build_summary,
        "formula_error_scan": "0 errors across both workbooks",
        "structural_preset_table_geometry_privacy": "passed",
        "visual_inspection": visual_qa,
        "human_actions": [
            "Original contract scope mapping and later-request approvals remain owner/researcher actions.",
            "Production Android/iOS signing ownership and store access remain unverified.",
            f"Third-party license identifiers/notices require owner/legal review ({pending_license_rows} pending rows).",
            "Repository/service access transfer evidence and rights/acceptance signatures remain pending.",
        ],
        "status": "passed_with_human_actions",
    }
    args.summary.parent.mkdir(parents=True, exist_ok=True)
    args.summary.write_text(json.dumps(payload, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(payload, indent=2))


if __name__ == "__main__":
    main()
