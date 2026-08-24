#!/usr/bin/env python3
import json
import unittest
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path('/private/tmp/fsookta-final-handover')


class CanonicalRound2Tests(unittest.TestCase):
    @classmethod
    def setUpClass(cls):
        cls.payload = json.loads((ROOT / 'working/task6/task6_corrected_payload.json').read_text())
        cls.master = load_workbook(ROOT / 'artifacts/07_Master_Test_and_Verification_Package.xlsx', data_only=False)
        cls.uat = load_workbook(ROOT / 'artifacts/08_UAT_Field_Test_and_Usability_Package.xlsx', data_only=False)

    def test_algorithm_sheet_is_exact_case_id_join(self):
        expected = [(x['case_id'], x['name'], x['category']) for x in self.payload['tests'] if x['category'] == 'Algorithm / reference']
        sheet = self.master['Algorithm Reference']
        actual = [(sheet.cell(r, 1).value, sheet.cell(r, 3).value, sheet.cell(r, 4).value) for r in range(5, sheet.max_row + 1)]
        self.assertEqual(expected, actual)

    def test_threshold_sheet_is_exact_case_id_join(self):
        expected = [(x['case_id'], x['name'], x['category']) for x in self.payload['tests'] if x['category'] == 'Invalid / boundary']
        sheet = self.master['Threshold Boundaries']
        actual = [(sheet.cell(r, 1).value, sheet.cell(r, 3).value, sheet.cell(r, 4).value) for r in range(5, sheet.max_row + 1)]
        self.assertEqual(expected, actual)

    def test_historical_round_and_unstated_version(self):
        h3 = self.payload['historical_uat'][2]
        h7 = self.payload['historical_uat'][6]
        self.assertEqual('2026-06-06', h3['date'])
        self.assertEqual('r2', h3['round_revision'])
        self.assertEqual('not stated', h7['version'])

    def test_automated_pass_schema_is_complete(self):
        required = {'case_id','requirement_id','precondition_input','expected','actual','status','baseline','tester_category','timestamp','method','raw_path','sha256'}
        for row in self.payload['result_rows']:
            if row['status'] == 'PASS':
                self.assertTrue(required <= row.keys(), row['case_id'])
                self.assertTrue(all(row[k] not in (None, '') for k in required - {'requirement_id'}), row['case_id'])

    def test_network_firebase_is_not_generic_pass(self):
        sheet = self.master['Network Error Applicability']
        rows = [[sheet.cell(r, c).value for c in range(1, sheet.max_column + 1)] for r in range(5, sheet.max_row + 1)]
        firebase = [r for r in rows if any('Firebase' in str(v) for v in r)]
        self.assertTrue(firebase)
        self.assertTrue(all('PASS' not in str(v) for row in firebase for v in row))


if __name__ == '__main__':
    unittest.main()
