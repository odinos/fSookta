#!/usr/bin/env python3
"""Use the recalculated LibreOffice PDF as the visual sheet truth for security workbook."""
import os,subprocess,tempfile
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(os.environ.get('SOOKTA_HANDOVER_ROOT','/private/tmp/fsookta-final-handover'));pdf=ROOT/'artifacts/09_Security_and_Access_Control_Matrices.pdf';xlsx=ROOT/'artifacts/09_Security_and_Access_Control_Matrices.xlsx';out=ROOT/'renders/task7/sheets/09_Security_and_Access_Control_Matrices';out.mkdir(parents=True,exist_ok=True)
for p in out.glob('*.png'):p.unlink()
with tempfile.TemporaryDirectory(prefix='task7-sheet-pages-') as tmp:
 prefix=Path(tmp)/'page';subprocess.run(['pdftoppm','-png','-r','150',str(pdf),str(prefix)],check=True,capture_output=True)
 pages=sorted(Path(tmp).glob('page-*.png'));names=load_workbook(xlsx,read_only=True).sheetnames;assert len(pages)==len(names)==15
 for i,(page,name) in enumerate(zip(pages,names),1):page.replace(out/f'{i:02d}_{"".join(c if c.isalnum() or c in "_-" else "_" for c in name)}.png')
print('security_sheet_renders=15 source=recalculated_pdf')
