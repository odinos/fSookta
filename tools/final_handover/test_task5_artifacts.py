#!/usr/bin/env python3
"""Behavior-first tests for Task 5 AI and data handover artifacts."""

from __future__ import annotations

import tempfile
import unittest
import copy
from pathlib import Path
import sys

sys.path.insert(0, str(Path(__file__).resolve().parent))

import build_task5_ai_data as builder
import record_task5_visual_qa as recorder
import verify_task5_artifacts as verifier
from docx import Document


SOURCE = Path(
    "/private/tmp/fsookta-final-handover/authoritative-materializations/"
    "source-g6tsggy7/source"
)


class SourceGroundingTests(unittest.TestCase):
    def test_inventory_separates_current_legacy_and_test_only_roles(self) -> None:
        facts = builder.inspect_source(SOURCE)
        roles = {row["role_id"]: row for row in facts["model_algorithm_inventory"]}
        self.assertEqual(len(roles), 10)
        self.assertEqual(roles["xgboost_onnx"]["training_class"], "Project-trained")
        self.assertEqual(roles["daily_logistic"]["training_class"], "Template coefficients")
        self.assertEqual(roles["movenet_thunder"]["training_class"], "Pretrained; not fine-tuned in project")
        self.assertEqual(roles["reba"]["authority"], "Primary deterministic")
        self.assertEqual(roles["xgboost_onnx"]["authority"], "Advisory only")
        self.assertEqual(roles["movenet_multipose"]["authority"], "Single-person eligibility gate")
        self.assertIn("evaluation_form_screen.dart", roles["movenet_multipose"]["citation"])
        self.assertEqual(roles["legacy_logistic_weights"]["current_reference_status"], "Legacy/test-only; not packaged")
        self.assertEqual(roles["deprecated_risk_alert"]["current_reference_status"], "Deprecated service; JSON/fromJson test only; absent from current pubspec bundle")
        self.assertIn("does not verify rootBundle", roles["deprecated_risk_alert"]["limitations"])
        self.assertEqual(len(roles["legacy_logistic_weights"]["binary_sha256"]), 64)
        self.assertEqual(len(roles["deprecated_risk_alert"]["binary_sha256"]), 64)

    def test_training_evidence_preserves_holdout_limits_and_missing_raw_metrics(self) -> None:
        facts = builder.inspect_source(SOURCE)
        evidence = {row["evidence_id"]: row for row in facts["training_evidence"]}
        xgb = evidence["xgboost_training"]
        self.assertEqual(xgb["total_samples"], 388)
        self.assertEqual(xgb["training_samples"], 298)
        self.assertEqual(xgb["holdout_samples"], 90)
        self.assertEqual(xgb["holdout_risk_accuracy"], 0.6667)
        self.assertEqual(xgb["raw_metrics_status"], "Pending Owner Action")
        self.assertIn("not present", xgb["raw_metrics_note"])
        self.assertFalse(evidence["daily_logistic_template"]["research_trained"])
        self.assertEqual(xgb["split_method"], "GroupShuffleSplit")
        self.assertEqual(xgb["test_size"], 0.22)
        self.assertEqual(xgb["random_seed"], 42)
        self.assertEqual(xgb["dataset_status"], "Pending Researcher Evidence")
        self.assertEqual(xgb["xgb_parameters"]["n_estimators"], 96)
        self.assertEqual(xgb["xgb_parameters"]["tree_method"], "hist")

    def test_persistence_and_csv_contracts_are_extracted_from_source(self) -> None:
        facts = builder.inspect_source(SOURCE)
        self.assertIn("sookta.dataSchemaVersion", facts["preference_keys"])
        self.assertEqual(facts["data_schema_version"], 2)
        self.assertGreaterEqual(len(facts["preference_keys"]), 10)
        self.assertIn("transaction_id", facts["all_history_csv_headers"])
        self.assertIn("Export Generated At", facts["all_history_csv_headers"])
        self.assertIn("REBA_before", facts["all_history_csv_headers"])
        self.assertIn("completion_status", facts["all_history_csv_headers"])
        persisted = {row["field"] for row in facts["persisted_record_fields"]}
        self.assertIn("profileId", persisted)
        self.assertIn("scoreBefore", persisted)
        self.assertIn("selectedImagePaths", persisted)
        self.assertIn("appVersion", persisted)
        self.assertIn("sookta.backup.schema.<version>.<timestamp>", facts["preference_keys"])
        app_versions = [row for row in facts["persisted_record_fields"] if row["field"] == "appVersion"]
        self.assertGreaterEqual(len({row["owner"] for row in app_versions}), 2)

    def test_source_contracts_cover_exact_runtime_behavior(self) -> None:
        facts = builder.inspect_source(SOURCE)
        triggers = facts["recommendation_triggers"]
        self.assertEqual(len(triggers), 24)
        lookup = {(r["activity"], r["risk_tier"]): r for r in triggers}
        self.assertEqual(lookup[("transplanting", "low")]["activity_key"], "act_transplant_ref_low")
        self.assertEqual(lookup[("fertilizing", "medium")]["activity_key"], "act_fert_ref_medium")
        self.assertEqual(lookup[("harvesting", "veryHigh")]["activity_key"], "act_harvest_ref_high")
        self.assertEqual(lookup[("transport", "veryHigh")]["weight_key"], "act_ref_weight_high")
        self.assertEqual(facts["daily_runtime_tiers"], [
            {"high_risk_count": "0-1", "level": "Low"},
            {"high_risk_count": "2-3", "level": "Watch"},
            {"high_risk_count": "4-5", "level": "High"},
            {"high_risk_count": "6-7", "level": "Critical"},
        ])
        self.assertFalse(facts["daily_probability_thresholds"]["used_by_predict_for_records"])

    def test_export_and_telemetry_contracts_are_semantically_complete(self) -> None:
        facts = builder.inspect_source(SOURCE)
        rows = facts["export_schema_rows"]
        self.assertEqual(len(rows), 84)
        self.assertEqual([row["field"] for row in rows], facts["all_history_csv_headers"])
        self.assertTrue(all(row["description"] and row["type"] and row["validation"] for row in rows))
        self.assertFalse(any("source-defined" in row["type"].lower() for row in rows))
        history = next(row for row in facts["persisted_schema_rows"] if row["field"] == "sookta.history")
        language = next(row for row in facts["persisted_schema_rows"] if row["field"] == "sookta.language")
        self.assertIn("Sensitive", history["privacy"])
        self.assertNotIn("Sensitive participant", language["privacy"])
        self.assertTrue(all(row["nullable"] in {"Yes", "No"} for row in facts["persisted_schema_rows"]))
        self.assertFalse(any("source-defined" in str(value).lower() for row in facts["persisted_schema_rows"] for value in row.values()))
        events = facts["firebase_telemetry"]["events"]
        self.assertEqual(events["assessment_calculated"], ["activity", "job_type", "primary_method", "risk_level", "score", "image_count", "uses_iso11228"])
        self.assertEqual(events["assessment_saved"], ["activity", "before_risk", "after_risk", "before_score", "after_score", "suggestion_count"])
        self.assertTrue(facts["firebase_telemetry"]["default_off"])
        self.assertTrue(facts["firebase_telemetry"]["crashlytics_context"])

    def test_persisted_contracts_follow_exact_dart_declarations_and_serializer_owner(self) -> None:
        facts = builder.inspect_source(SOURCE)
        rows = {(row.get("owner"), row["field"]): row for row in facts["persisted_schema_rows"]}
        expected = {
            ("EvaluationHistoryRecord", "id"):("int", "No", "N/A"),
            ("EvaluationHistoryRecord", "farmerBmiCategory"):("String?", "Yes", "N/A"),
            ("EvaluationHistoryRecord", "timeOnTaskSeconds"):("int?", "Yes", "seconds"),
            ("EvaluationDraft", "selectedImagePaths"):("List<String>", "No", "N/A"),
            ("RebaInputData", "trunkTwist"):("bool", "No", "N/A"),
            ("PoseRebaFrameAnalysis", "imageIndex"):("int", "No", "N/A"),
            ("PoseRebaFrameAnalysis", "timestampMs"):("int?", "Yes", "milliseconds"),
            ("PoseRebaFrameAnalysis", "neckFlexionDeg"):("double?", "Yes", "degrees"),
            ("MotionAnalysisSummary", "sampleRateFps"):("double", "No", "frames/second"),
            ("AssessmentBreakdown", "isoMethod"):("AssessmentMethod?", "Yes", "N/A"),
            ("AssessmentBreakdown", "isoResult"):("ErgoResult?", "Yes", "N/A"),
            ("AssessmentBreakdown", "xgboostProbability"):("double?", "Yes", "probability 0..1"),
            ("ErgoResult", "techScore"):("double", "No", "score/ratio"),
            ("ErgoInputData", "durationHours"):("double", "No", "hours"),
            ("ErgoInputData", "workDaysPerWeek"):("double", "No", "days/week"),
        }
        for key, contract in expected.items():
            row = rows[key]
            self.assertEqual((row["type"], row["nullable"], row["unit"]), contract, key)
            self.assertIn(key[0], row["persisted_location"])
        self.assertIn("AssessmentBreakdown.ergoInput", rows[("ErgoInputData", "durationHours")]["persisted_location"])
        self.assertIn("EvaluationHistoryRecord.assessmentBreakdown", rows[("AssessmentBreakdown", "isoResult")]["persisted_location"])
        self.assertNotIn(("AssessmentBreakdown", "techScore"), rows)

    def test_persisted_semantics_units_and_privacy_are_source_exact(self) -> None:
        facts = builder.inspect_source(SOURCE)
        rows = {(row.get("owner"), row["field"]): row for row in facts["persisted_schema_rows"]}
        expected = {
            ("ErgoInputData", "liftFrequency"):("lifts/minute", ">= 0 lifts/minute"),
            ("ErgoInputData", "horizontalDist"):("cm", ">= 0 cm"),
            ("ErgoInputData", "verticalHeight"):("cm", ">= 0 cm"),
            ("ErgoInputData", "initialForce"):("N", ">= 0 N"),
            ("ErgoInputData", "sustainForce"):("N", ">= 0 N"),
            ("EvaluationDraft", "horizontalDistanceText"):("cm", "Numeric text in cm"),
            ("EvaluationDraft", "verticalHeightText"):("cm", "Numeric text in cm"),
            ("EvaluationDraft", "initialForce"):("N", ">= 0 N"),
            ("EvaluationDraft", "sustainForce"):("N", ">= 0 N"),
            ("UserProfile", "incomePerYear"):("THB/year", "Numeric text >= 0 THB/year or blank"),
            ("UserProfile", "age"):("years", "Numeric text >= 0 years or blank"),
            ("UserProfile", "height"):("cm", "Numeric text > 0 cm or blank"),
            ("UserProfile", "weight"):("kg", "Numeric text > 0 kg or blank"),
            ("EvaluationHistoryRecord", "farmerAge"):("years", "Numeric text >= 0 years or blank"),
            ("EvaluationHistoryRecord", "farmerHeight"):("cm", "Numeric text > 0 cm or blank"),
            ("EvaluationHistoryRecord", "farmerWeight"):("kg", "Numeric text > 0 kg or blank"),
            ("EvaluationHistoryRecord", "farmerBmi"):("kg/m2", "> 0 kg/m2 when present"),
            ("EvaluationHistoryRecord", "aiRiskPercent"):("%", "Integer 0..100 when present"),
            ("EvaluationHistoryRecord", "expertReba"):("REBA score points", ">= 0 when present"),
        }
        for key, contract in expected.items():
            self.assertEqual((rows[key]["unit"], rows[key]["allowed"]), contract, key)
        self.assertIn("multiplies by 60", rows[("ErgoInputData", "liftFrequency")]["description"])
        self.assertEqual(rows[("ErgoResult", "limitValue")]["unit"], "kg, N, or encoded REBA limit (context-dependent)")
        self.assertIn("lifting RWL", rows[("ErgoResult", "limitValue")]["description"])
        self.assertEqual(rows[("UserProfile", "incomePerYear")]["privacy"], "Sensitive financial data")
        self.assertEqual(rows[("UserProfile", "location")]["privacy"], "Sensitive participant/profile data")
        self.assertNotEqual(rows[("EvaluationDraft", "durationHours")]["allowed"], "0..1")

    def test_export_dictionary_exact_semantics_and_fallbacks(self) -> None:
        facts = builder.inspect_source(SOURCE)
        rows = {row["field"]: row for row in facts["export_schema_rows"]}
        self.assertEqual((rows["BMI Category"]["type"], rows["BMI Category"]["allowed"]),
                         ("text", "Underweight; Normal weight; Above Asian BMI range; localized Thai values; '-' when unavailable"))
        self.assertEqual((rows["assessment_time"]["type"], rows["assessment_time"]["allowed"]),
                         ("HH:mm:ss text", "00:00:00..23:59:59 local device time"))
        self.assertEqual((rows["time_on_task_seconds"]["type"], rows["time_on_task_seconds"]["nullable"], rows["time_on_task_seconds"]["unit"]),
                         ("integer", "Yes", "seconds"))
        self.assertEqual(rows["time_on_task_seconds"]["missing"], "Blank string when record.timeOnTaskSeconds is null")
        self.assertEqual(rows["ISO 11228 Risk Level"]["type"], "text")
        self.assertEqual(rows["ISO_risk_before"]["type"], "text")
        self.assertEqual(rows["Economic Impact Formula"]["type"], "text")
        self.assertEqual(rows["REBA_reduction_percent"]["type"], "percentage text")

    def test_independent_verifier_rejects_any_full_dictionary_contract_drift(self) -> None:
        facts = builder.inspect_source(SOURCE)
        mutated = copy.deepcopy(facts)
        next(row for row in mutated["persisted_schema_rows"] if row["owner"] == "MotionAnalysisSummary" and row["field"] == "sampleRateFps")["unit"] = "Hz"
        with self.assertRaises(AssertionError):
            verifier.verify_source_contracts(mutated, SOURCE)
        mutated = copy.deepcopy(facts)
        next(row for row in mutated["export_schema_rows"] if row["field"] == "Age")["type"] = "integer"
        with self.assertRaises(AssertionError):
            verifier.verify_source_contracts(mutated, SOURCE)

    def test_export_trend_and_multipose_and_firebase_call_sites_are_exact(self) -> None:
        facts = builder.inspect_source(SOURCE)
        trend = next(row for row in facts["export_schema_rows"] if row["field"] == "trend_level")
        self.assertEqual(trend["allowed"], "Low; Medium; High; Very high (localized when Thai export is selected)")
        self.assertNotIn("Critical", trend["derivation"])
        multipose = {r["role_id"]: r for r in facts["model_algorithm_inventory"]}["movenet_multipose"]
        self.assertIn("person count", multipose["output"].lower())
        self.assertNotIn("candidate poses", multipose["output"].lower())
        telemetry = facts["firebase_telemetry"]
        self.assertEqual(telemetry["generic_call_sites"]["pose_analysis_failed"], ["platform", "error_code"])
        self.assertEqual(telemetry["generic_call_site_paths"]["pose_analysis_failed"], "lib/screens/main/evaluation_form_screen.dart")
        self.assertIn("lib/app/sookta_app.dart", telemetry["analytics_observer_call_sites"])
        self.assertIn("lib/screens/main/training_data_export_screen.dart", telemetry["wrapper_call_sites"]["export_created"])
        self.assertTrue(telemetry["log_app_open"])
        self.assertTrue(telemetry["analytics_observer_navigation"])

    def test_independent_verifier_extracts_source_contracts_without_builder_facts(self) -> None:
        contracts = verifier.extract_authoritative_source_contracts(SOURCE)
        self.assertEqual(contracts["dart_fields"][("EvaluationHistoryRecord", "id")], "int")
        self.assertEqual(contracts["dart_fields"][("AssessmentBreakdown", "isoResult")], "ErgoResult?")
        self.assertEqual(contracts["trend_levels"], ["Low", "Medium", "High", "Very high"])
        self.assertEqual(contracts["telemetry_events"]["pose_analysis_failed"], ["platform", "error_code"])
        self.assertIn("lib/screens/main/training_data_export_screen.dart", contracts["telemetry_call_sites"]["export_created"])
        self.assertEqual(contracts["generic_call_site_paths"]["pose_analysis_failed"], "lib/screens/main/evaluation_form_screen.dart")
        self.assertEqual(contracts["analytics_observer_call_sites"], ["lib/app/sookta_app.dart"])
        self.assertEqual(contracts["xgb"]["test_size"], 0.22)
        self.assertEqual(contracts["xgb"]["random_seed"], 42)
        self.assertIn(("EvaluationHistoryRecord", "id"), contracts["serialized_fields"])
        self.assertIn(("AssessmentBreakdown", "isoResult"), contracts["serialized_fields"])
        self.assertIn(("ErgoResult", "techScore"), contracts["serialized_fields"])
        self.assertNotIn(("AssessmentBreakdown", "techScore"), contracts["serialized_fields"])


class ContentContractTests(unittest.TestCase):
    def test_ai_report_avoids_manual_page_break_that_can_create_blank_page(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            output = Path(temp) / "ai.docx"
            builder.build_ai_report(output, builder.inspect_source(SOURCE))
            document = Document(output)
            self.assertFalse(any('w:type="page"' in paragraph._p.xml for paragraph in document.paragraphs))

    def test_signature_table_rows_are_compact_and_cannot_split(self) -> None:
        document = Document()
        builder._add_signature_table(document)
        table = document.tables[0]
        for row in table.rows:
            self.assertTrue(row._tr.xpath("./w:trPr/w:cantSplit"))
            for cell in row.cells:
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        self.assertLessEqual(run.font.size.pt, 8)

    def test_bullet_helper_applies_readable_paragraph_spacing(self) -> None:
        document = Document()
        builder._add_bullets(document, ["One", "Two"])
        self.assertTrue(all(p.paragraph_format.space_after.pt >= 4 for p in document.paragraphs))
        self.assertEqual([p.text for p in document.paragraphs], ["• One", "\n• Two"])

    def test_required_workbook_sheets_are_exact(self) -> None:
        self.assertEqual(len(builder.ALGORITHM_WORKBOOK_SHEETS), 15)
        self.assertEqual(len(builder.DATA_WORKBOOK_SHEETS), 9)
        self.assertIn("Recommendation Messages", builder.ALGORITHM_WORKBOOK_SHEETS)
        self.assertIn("Export Schema Order", builder.DATA_WORKBOOK_SHEETS)

    def test_claim_guard_rejects_unqualified_validation_claims(self) -> None:
        verifier.assert_claim_boundaries(
            "Software verification passed. No clinical validation or external validity is claimed."
        )
        verifier.assert_claim_boundaries(
            "Do not state that the model is clinically validated or externally validated."
        )
        with self.assertRaises(AssertionError):
            verifier.assert_claim_boundaries("The model is clinically validated and production-ready.")
        with self.assertRaises(AssertionError):
            verifier.assert_claim_boundaries("All participants consented and data was de-identified.")

    def test_status_vocabulary_is_fail_closed(self) -> None:
        verifier.verify_status_values(["Complete", "Pending Owner Action", "N/A with Rationale"])
        with self.assertRaises(AssertionError):
            verifier.verify_status_values(["Almost complete"])

    def test_synthetic_examples_cannot_contain_real_capture_identity(self) -> None:
        verifier.assert_synthetic_examples(["SYN-001", "Synthetic Farmer", "2026-08-24T09:00:00+07:00"])
        with self.assertRaises(AssertionError):
            verifier.assert_synthetic_examples(["FSK-944631", "ddd"])


class VisualManifestTests(unittest.TestCase):
    def test_visual_expectations_cover_all_ten_primary_artifacts(self) -> None:
        expected = recorder.task5_visual_expectations()
        self.assertEqual(len(expected), 10)
        self.assertTrue(all(renders for renders in expected.values()))

    def test_checksum_update_preserves_existing_entries(self) -> None:
        with tempfile.TemporaryDirectory() as temp:
            root = Path(temp)
            (root / "artifacts").mkdir()
            (root / "manifests").mkdir()
            prior = root / "artifacts/prior.txt"
            current = root / "artifacts/current.txt"
            prior.write_text("prior", encoding="utf-8")
            current.write_text("current", encoding="utf-8")
            recorder.update_checksums(root, [prior])
            recorder.update_checksums(root, [current])
            manifest = (root / "manifests/SHA256SUMS.txt").read_text(encoding="ascii")
            self.assertIn("artifacts/prior.txt", manifest)
            self.assertIn("artifacts/current.txt", manifest)


if __name__ == "__main__":
    unittest.main()
