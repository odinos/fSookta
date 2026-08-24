#!/usr/bin/env python3
import os
from pathlib import Path
from openpyxl import load_workbook
from openpyxl.worksheet.properties import PageSetupProperties

root=Path(os.environ.get('FSOOKTA_HANDOVER_ROOT','/private/tmp/fsookta-final-handover'))
for name in ['09_Security_and_Access_Control_Matrices.xlsx','11_Publication_Tables.xlsx']:
    path=root/'artifacts'/name
    wb=load_workbook(path)
    if name.startswith('09_'):
        wb['Control Summary']['B9']='=IF(COUNTIF(\'Risk Register\'!E5:E9,"Open*")>0,"OPEN ACTIONS","NO OPEN ROWS")'
    for ws in wb.worksheets:
        ws.sheet_properties.pageSetUpPr=PageSetupProperties(fitToPage=True,autoPageBreaks=False)
        ws.page_setup.orientation='landscape'
        ws.page_setup.paperSize=ws.PAPERSIZE_A4
        ws.page_setup.fitToWidth=1
        ws.page_setup.fitToHeight=1
        ws.sheet_view.showGridLines=False
        ws.freeze_panes='A5'
        ws.print_title_rows='1:4'
        ws.print_area=f'A1:{ws.cell(ws.max_row,ws.max_column).coordinate}'
        ws.page_margins.left=0.2;ws.page_margins.right=0.2;ws.page_margins.top=0.35;ws.page_margins.bottom=0.35
        ws.oddFooter.center.text=f'SookTa {name} | {ws.title}'
    wb.save(path)
    print(f'patched {name} sheets={len(wb.worksheets)}')
