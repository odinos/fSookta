#!/usr/bin/env python3
"""Independent fail-closed verifier for Task 5 handover artifacts."""

from __future__ import annotations

import argparse
import hashlib
import json
import re
import zipfile
from pathlib import Path

from openpyxl import load_workbook
from pypdf import PdfReader

from build_task5_ai_data import (
    ALGORITHM_WORKBOOK_SHEETS, ALLOWED_STATUSES, COMMIT,
    DATA_WORKBOOK_SHEETS, VERSION,
)


def digest(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def archive_text(path: Path) -> str:
    if path.suffix.lower() not in {".docx", ".xlsx"}:
        return path.read_text(encoding="utf-8", errors="ignore")
    with zipfile.ZipFile(path) as archive:
        return "\n".join(
            archive.read(name).decode("utf-8", errors="ignore")
            for name in archive.namelist()
            if name.endswith((".xml", ".rels"))
        )


def doc_text(path: Path) -> str:
    from docx import Document
    document = Document(path)
    return "\n".join(p.text for p in document.paragraphs) + "\n" + "\n".join(
        " | ".join(cell.text for cell in row.cells)
        for table in document.tables for row in table.rows
    )


def verify_status_values(values) -> list[str]:
    normalized = [str(value).strip() for value in values if value is not None and str(value).strip()]
    assert normalized and all(value in ALLOWED_STATUSES for value in normalized), normalized
    return normalized


def assert_claim_boundaries(text: str) -> None:
    sentences = re.split(r"(?<=[.!?])\s+|\n+", text.lower())
    forbidden = [
        "clinically validated", "clinical validation completed", "externally validated",
        "production-ready", "all participants consented", "data was de-identified",
        "ethical approval obtained", "researcher accepted",
    ]
    for sentence in sentences:
        # A prohibited-claim register must be able to name the claim it bans.
        negated = any(marker in sentence for marker in (
            "do not state", "do not claim", "no clinical validation",
            "not clinically validated", "not externally validated",
            "is not claimed", "are not claimed", "not proven",
        ))
        if negated:
            continue
        for phrase in forbidden:
            assert phrase not in sentence, phrase


def assert_synthetic_examples(values: list[str]) -> None:
    joined = "\n".join(values)
    forbidden = ["FSK-944631", "ddd"]
    assert not any(token in joined for token in forbidden), joined


def verify_visual_manifest(staging: Path, manifest_path: Path, expected: dict[str, list[str]]) -> None:
    payload = json.loads(manifest_path.read_text(encoding="utf-8"))
    assert payload["status"] == "passed"
    rows = {row["path"]: row for row in payload["artifacts"]}
    assert set(rows) == set(expected)
    for artifact_relative, render_relatives in expected.items():
        row = rows[artifact_relative]
        artifact = staging / artifact_relative
        assert row["sha256"] == digest(artifact)
        renders = {render["path"]: render for render in row["renders"]}
        assert set(renders) == set(render_relatives)
        for relative in render_relatives:
            assert renders[relative]["status"] == "passed"
            assert renders[relative]["sha256"] == digest(staging / relative)


def verify_formula_summary(path: Path) -> dict:
    data = json.loads(path.read_text(encoding="utf-8"))
    assert data["status"] == "passed"
    assert data["formula_error_matches"] == 0
    assert data["formula_contracts"]
    for contract in data["formula_contracts"]:
        assert contract["status"] == "passed"
        assert contract["expected_formula"] == contract["actual_formula"]
        assert contract["expected_value"] == contract["actual_value"]
    return data


def verify_checksums(staging: Path, required: list[Path]) -> int:
    rows: dict[str, str] = {}
    for line in (staging / "manifests/SHA256SUMS.txt").read_text(encoding="ascii").splitlines():
        checksum, relative = line.split("  ", 1); rows[relative] = checksum
    for relative, checksum in rows.items():
        path = staging / relative
        assert path.is_file(), relative
        assert digest(path) == checksum, relative
    for path in required:
        relative = path.relative_to(staging).as_posix()
        assert rows.get(relative) == digest(path), relative
    return len(rows)


def verify(staging: Path) -> dict:
    artifacts = staging / "artifacts"; manifests = staging / "manifests"
    stems = [
        "04_AI_Algorithm_and_Model_Technical_Report",
        "04_Algorithm_Recommendation_and_Reference_Matrices",
        "04_AI_Training_and_Non_Training_Statement",
        "05_Local_Storage_and_Data_Management_Manual",
        "05_Data_Dictionary_and_Export_Schema",
    ]
    required: list[Path] = []
    for stem in stems:
        editable = artifacts / f"{stem}{'.xlsx' if 'Matrices' in stem or 'Dictionary' in stem else '.docx'}"
        pdf = artifacts / f"{stem}.pdf"
        assert editable.is_file() and editable.stat().st_size > 1000, editable
        assert pdf.is_file() and pdf.stat().st_size > 1000, pdf
        required.extend([editable, pdf])
    ai = doc_text(required[0]); statement = doc_text(required[4]); manual = doc_text(required[6])
    for text in [ai, statement, manual]:
        assert VERSION in text and COMMIT in text
        assert_claim_boundaries(text)
    for token in ["MoveNet Thunder", "XGBoost ONNX", "Template coefficients", "Primary deterministic", "Advisory only", "no clinical validation"]:
        assert token.lower() in ai.lower(), token
    for token in ["Prohibited claims", "Complete - Pending Signature", "Preparer", "Technical reviewer", "Owner", "Researcher"]:
        assert token in statement, token
    for token in ["SharedPreferences", "sookta.latestBackup", "schema version is 2", "N/A with Rationale", "default off", "retention", "deletion"]:
        assert token.lower() in manual.lower(), token
    algorithm = load_workbook(required[2], read_only=True, data_only=False)
    data = load_workbook(required[8], read_only=True, data_only=False)
    assert algorithm.sheetnames == ALGORITHM_WORKBOOK_SHEETS
    assert data.sheetnames == DATA_WORKBOOK_SHEETS
    statuses = []
    for workbook in [algorithm, data]:
        for sheet in workbook.worksheets:
            for row in sheet.iter_rows(values_only=True):
                for value in row:
                    if value in ALLOWED_STATUSES:
                        statuses.append(value)
    verify_status_values(statuses)
    assert {"Complete", "Pending Owner Action", "Pending Researcher Evidence", "N/A with Rationale", "Complete - Pending Signature"}.issubset(set(statuses))
    examples = []
    sheet = data["Export Schema Order"]
    headers = [cell.value for cell in next(sheet.iter_rows())]
    if "Synthetic example" in headers:
        index = headers.index("Synthetic example")
        examples = [str(row[index]) for row in sheet.iter_rows(min_row=2, values_only=True) if row[index] is not None]
        assert_synthetic_examples(examples)
    algorithm.close(); data.close()
    summary = verify_formula_summary(manifests / "task5_workbook_build_summary.json")
    expected = json.loads((manifests / "task5_expected_visual_renders.json").read_text(encoding="utf-8"))
    verify_visual_manifest(staging, manifests / "task5_visual_qa_manifest.json", expected)
    counts = json.loads((manifests / "task5_render_counts.json").read_text(encoding="utf-8"))
    assert len(PdfReader(required[1]).pages) == counts["ai_report"]
    assert len(PdfReader(required[5]).pages) == counts["training_statement"]
    assert len(PdfReader(required[7]).pages) == counts["data_manual"]
    assert len(PdfReader(required[3]).pages) == counts["algorithm_pdf"]
    assert len(PdfReader(required[9]).pages) == counts["data_pdf"]
    forbidden_secret = re.compile(r"(?i)(-----BEGIN (?:RSA |EC |OPENSSH )?PRIVATE KEY-----|AIza[0-9A-Za-z_-]{20,}|password\s*[:=]\s*\S+|api[_-]?key\s*[:=]\s*\S+)")
    forbidden_identity = re.compile(r"FSK-944631|>ddd<")
    for path in required:
        text = archive_text(path)
        assert not forbidden_secret.search(text), path
        assert not forbidden_identity.search(text), path
    checksum_entries = verify_checksums(staging, required)
    input_data = json.loads((staging / "working/task5/task5_build_input.json").read_text(encoding="utf-8"))
    missing = [row for row in input_data["source_citations"] if row["status"] != "Resolved"]
    assert not missing, missing
    return {
        "status": "passed_with_human_actions", "primary_artifacts": 10,
        "document_pdf_pages": counts["ai_report"] + counts["training_statement"] + counts["data_manual"],
        "workbook_sheets": summary["sheets"], "workbook_pdf_pages": counts["algorithm_pdf"] + counts["data_pdf"],
        "formula_contracts": len(summary["formula_contracts"]), "formula_error_scan": "0 errors",
        "visual_manifest": digest(manifests / "task5_visual_qa_manifest.json"),
        "checksum_entries": checksum_entries, "model_algorithm_roles": len(input_data["model_algorithm_inventory"]),
        "training_evidence_records": len(input_data["training_evidence"]),
        "data_dictionary_records": summary["data_dictionary_records"],
        "human_actions": input_data["human_actions"],
    }


def main() -> None:
    parser = argparse.ArgumentParser(); parser.add_argument("--staging", type=Path, default=Path("/private/tmp/fsookta-final-handover")); parser.add_argument("--summary", type=Path, required=True)
    args = parser.parse_args(); result = verify(args.staging)
    args.summary.write_text(json.dumps(result, indent=2) + "\n", encoding="utf-8")
    print(json.dumps(result, indent=2))


if __name__ == "__main__":
    main()
