#!/usr/bin/env python3
import json, os
from collections import Counter
from pathlib import Path
from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT=Path(os.environ.get('FSOOKTA_HANDOVER_ROOT','/private/tmp/fsookta-final-handover')); A=ROOT/'artifacts'
P=json.loads((ROOT/'working/task6/task6_corrected_payload.json').read_text())
BLUE=RGBColor(46,116,181); DARK=RGBColor(31,77,120); INK=RGBColor(11,37,69); GRAY=RGBColor(90,98,108)
def font(run,size=10,bold=False,color=INK,italic=False):
 run.font.name='Calibri'; run._element.get_or_add_rPr().rFonts.set(qn('w:ascii'),'Calibri'); run._element.rPr.rFonts.set(qn('w:hAnsi'),'Calibri'); run.font.size=Pt(size); run.bold=bold; run.italic=italic; run.font.color.rgb=color
def shade(cell,fill):
 tc=cell._tc; pr=tc.get_or_add_tcPr(); shd=OxmlElement('w:shd'); shd.set(qn('w:fill'),fill); pr.append(shd)
def margins(cell,top=80,bottom=80,start=120,end=120):
 pr=cell._tc.get_or_add_tcPr(); mar=pr.first_child_found_in('w:tcMar')
 if mar is None: mar=OxmlElement('w:tcMar'); pr.append(mar)
 for tag,val in (('top',top),('bottom',bottom),('start',start),('end',end)):
  x=OxmlElement('w:'+tag); x.set(qn('w:w'),str(val)); x.set(qn('w:type'),'dxa'); mar.append(x)
def table(doc,headers,rows,widths=None,size=8):
 t=doc.add_table(rows=1,cols=len(headers)); t.alignment=WD_TABLE_ALIGNMENT.LEFT; t.style='Table Grid'; t.autofit=False
 for i,h in enumerate(headers):
  c=t.rows[0].cells[i]; c.text=str(h); shade(c,'F2F4F7'); margins(c); c.vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.CENTER
  for r in c.paragraphs[0].runs: font(r,size,bold=True,color=DARK)
 trPr=t.rows[0]._tr.get_or_add_trPr(); repeat=OxmlElement('w:tblHeader'); repeat.set(qn('w:val'),'true'); trPr.append(repeat)
 for row in rows:
  cells=t.add_row().cells
  for i,v in enumerate(row):
   cells[i].text='' if v is None else str(v); margins(cells[i]); cells[i].vertical_alignment=WD_CELL_VERTICAL_ALIGNMENT.TOP
   for r in cells[i].paragraphs[0].runs: font(r,size,color=INK)
 if widths:
  for row in t.rows:
   for c,w in zip(row.cells,widths): c.width=Inches(w)
 return t
def setup(title,subtitle):
 d=Document(); s=d.sections[0]; s.page_width=Inches(8.5); s.page_height=Inches(11); s.top_margin=s.bottom_margin=s.left_margin=s.right_margin=Inches(1); s.header_distance=s.footer_distance=Inches(.492)
 styles=d.styles; n=styles['Normal']; n.font.name='Calibri'; n.font.size=Pt(10); n.paragraph_format.space_after=Pt(6); n.paragraph_format.line_spacing=1.10
 for name,size,before,after,color in [('Heading 1',16,16,8,BLUE),('Heading 2',13,12,6,BLUE),('Heading 3',12,8,4,DARK)]:
  st=styles[name]; st.font.name='Calibri'; st.font.size=Pt(size); st.font.bold=True; st.font.color.rgb=color; st.paragraph_format.space_before=Pt(before); st.paragraph_format.space_after=Pt(after)
 hp=s.header.paragraphs[0]; hp.text='SOOKTA | FINAL HANDOVER TECHNICAL EVIDENCE'; hp.alignment=WD_ALIGN_PARAGRAPH.LEFT
 for r in hp.runs: font(r,8,bold=True,color=GRAY)
 fp=s.footer.paragraphs[0]; fp.text='Controlled evidence package | Version 1.3.11+28'; fp.alignment=WD_ALIGN_PARAGRAPH.CENTER
 for r in fp.runs: font(r,8,color=GRAY)
 p=d.add_paragraph(); p.paragraph_format.space_after=Pt(4); r=p.add_run(title); font(r,23,bold=True,color=INK)
 p=d.add_paragraph(); p.paragraph_format.space_after=Pt(14); r=p.add_run(subtitle); font(r,13,color=GRAY)
 table(d,['Control','Value'],[['Baseline','1.3.11+28'],['Authoritative commit',P['baseline']['commit']],['Evidence date','2026-08-23 UTC'],['Claim boundary',P['claim_boundary']]],[1.7,4.8],8.5)
 return d
def para(d,text,bold=False,color=INK):
 p=d.add_paragraph(); r=p.add_run(text); font(r,10,bold,color); return p
def algorithm_report():
 d=setup('Final Algorithm Verification Report','Corrected test attribution, evidence tuples, requirement boundary, and historical retest trace')
 d.add_heading('Executive conclusion',level=1); para(d,'The final automated baseline contains exactly 135 completed Flutter tests plus one static-analysis PASS. The expanded reporter terminal line is completion proof only and is not counted as a test. No device, participant, Firebase runtime delivery, performance, usability, or acceptance PASS is inferred.')
 d.add_heading('Expanded reporter interpretation',level=1); para(d,'For reporter counter +N, the preceding displayed test is the test that completed. Therefore +135 All tests passed proves AUTO-135 - clears the draft after a successful evaluation save - completed. The summary line is excluded as a test case.')
 cats=Counter(x['category'] for x in P['tests']); table(d,['Category','Completed tests'],sorted(cats.items()),[4.8,1.7],9)
 d.add_heading('Hash-bound PASS evidence',level=1); pass_rows=[x for x in P['result_rows'] if x['status']=='PASS']; table(d,['Case range','Baseline / timestamp','Method','Raw evidence / SHA-256'],[[f"AUTO-001-AUTO-135",'1.3.11+28 / 2026-08-23T14:32:28Z','flutter test --reporter expanded',pass_rows[0]['evidence_tuple']['raw_path']+'\n'+pass_rows[0]['evidence_tuple']['sha256']],['STATIC-001','1.3.11+28 / '+pass_rows[-1]['evidence_tuple']['timestamp'],pass_rows[-1]['evidence_tuple']['method'],pass_rows[-1]['evidence_tuple']['raw_path']+'\n'+pass_rows[-1]['evidence_tuple']['sha256']]],[1.3,1.6,1.6,2.0],7.5)
 para(d,'Each PASS is joined by Case ID to the canonical result record containing requirement ID where applicable, precondition/input, expected, actual, status, baseline/version, tester category, timestamp/date, exact method/command, raw evidence path, and SHA-256. The independent verifier proves the join across every workbook test surface.')
 d.add_heading('Requirement status boundary',level=1); table(d,['Terminal status','Count'],sorted(P['requirement_status_counts'].items()),[4.8,1.7],9); para(d,'All 155 evidence_map records retain their original terminal status, authoritative sources, rationale, human action, evidence, and final-version claim. Task 2 and Task 6 evidence are additive layers only.')
 d.add_heading('Historical bugs, corrections, and retests',level=1); table(d,['ID / version','Defect and source path','Before / after commits','Retest and evidence hash'],[[x['defect_id']+'\n'+x['version'],x['description']+'\n'+x['source_path'],x['before_commit']+'\n'+x['after_commit'],x['retest']+'\n'+x['evidence_path']+'\n'+x['sha256']] for x in P['defects']],[1.1,1.8,1.7,1.9],7)
 d.add_heading('Human actions remaining',level=1)
 for item in P['human_gaps']:
  para(d,f"{item['action_id']} | {item['owner']} | {item['status']}: {item['action']} Closure evidence: {item['closure_evidence']}")
 d.add_page_break(); d.add_heading('Appendix A - Exact automated test order',level=1); para(d,'Rows are reconstructed independently from counter transitions in the raw expanded-reporter log.')
 table(d,['Case','Category','Source path','Exact test name'],[[x['case_id'],x['category'],x['source_path'],x['name']] for x in P['tests']],[.65,1.25,1.95,2.65],6.7)
 d.save(A/'07_Final_Algorithm_Verification_Report.docx')
def uat_report():
 d=setup('UAT and Field-Test Technical Report','Corrected historical metadata, controlled SUS instrument, and final human-evidence boundary')
 d.add_heading('Current final-baseline position',level=1); para(d,'No final participant UAT, SUS score, physical-device matrix, performance measurement, or authorized acceptance is claimed. Historical records remain historical even where a simulator, emulator, host test, build workaround, or temporary assessment bypass produced an observation.')
 d.add_heading('Historical evidence register',level=1)
 for x in P['historical_uat']:
  d.add_heading(f"{x['evidence_id']} | {x['date']} | {x['version']}",level=2)
  table(d,['Dimension','Source-exact metadata'],[['Round / revision',x['round_revision'] or 'not applicable'],['Device',x['device']],['Runner / harness',x['runner_harness']],['Build workaround',x['build_workaround']],['Assessment bypass',x['assessment_bypass']],['Blocker',x['blocker']],['Observation',x['observation']],['Evidence',x['evidence_path']],['SHA-256',x['sha256']],['Layer',x['evidence_layer']]],[1.55,4.95],8)
 d.add_heading('Controlled SUS instrument',level=1); para(d,'Use the standard 10-item System Usability Scale. Response anchors: 1 = Strongly disagree; 5 = Strongly agree. Enter whole numbers 1-5. Score only when all 10 responses exist: odd items contribute response - 1; even items contribute 5 - response; multiply the sum by 2.5.')
 table(d,['Item','Approved wording'],[[x['item'],x['wording']] for x in P['sus']['items']],[.55,5.95],8)
 d.add_heading('Final protocol and acceptance actions',level=1); para(d,'Researcher: obtain consent/ethics clearance, execute participant-coded tasks, collect observations and SUS responses, interpret results, and retain raw evidence. Owner: execute camera/gallery, TTS listening, share/export, offline inference, supported-device matrix, and performance measurements with thresholds. Authorized representatives: sign acceptance only after evidence review.')
 d.add_heading('Prohibited substitutions',level=1); para(d,'Host/widget tests do not substitute for physical-device permission, audio, performance, or usability evidence. Simulator/emulator results do not substitute for physical-device UAT. A clean temporary worktree is a build-path workaround, not an assessment bypass. Historical temporary assessment bypass observations cannot establish the final production baseline.')
 d.add_heading('Acceptance status',level=1); table(d,['Role','Required evidence','Current status'],[['Owner','Device/build IDs, timestamps, methods, raw paths, hashes','Pending Owner Action'],['Researcher','Consent, coded participants, protocol results, SUS responses and interpretation','Pending Researcher Evidence'],['Authorized approver','Signed acceptance against reviewed evidence','Pending Authorized Signature']],[1.3,3.7,1.5],8)
 d.save(A/'08_UAT_and_Field_Test_Technical_Report.docx')
def main():
 algorithm_report(); uat_report(); print(json.dumps({'status':'docs_corrected','docx':2}))

if __name__=='__main__': main()
