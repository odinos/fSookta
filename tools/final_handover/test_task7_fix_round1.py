import json
import os
import re
import unittest
import zipfile
from pathlib import Path
from xml.etree import ElementTree as ET

from docx import Document
from openpyxl import load_workbook


ROOT = Path(os.environ.get("SOOKTA_HANDOVER_ROOT", "/private/tmp/fsookta-final-handover"))
ART = ROOT / "artifacts"
HERE = Path(__file__).parent
A = "{http://schemas.openxmlformats.org/drawingml/2006/main}"
P = "{http://schemas.openxmlformats.org/presentationml/2006/main}"


def doc_text(name):
    doc = Document(ART / name)
    return "\n".join(p.text for p in doc.paragraphs) + "\n" + "\n".join(c.text for t in doc.tables for r in t.rows for c in r.cells)


class Task7FixRound1Tests(unittest.TestCase):
    def test_no_unsupported_media_or_history_deletion_claim(self):
        text = "\n".join(doc_text(name) for name in (
            "09_Security_Privacy_and_Data_Protection_Report.docx",
            "10_End_User_Manual.docx", "10_Research_Admin_Manual.docx",
            "10_Developer_Handover_Manual.docx"))
        wb = load_workbook(ART / "09_Security_and_Access_Control_Matrices.xlsx", data_only=False)
        text += "\n" + "\n".join(str(c.value or "") for ws in wb for row in ws.iter_rows() for c in row)
        self.assertNotIn("Linked path deletion", text)
        self.assertNotIn("Linked local paths may be deleted", text)
        self.assertIn("Farmer removal leaves existing history", text)
        self.assertIn("No File.delete path", text)

    def test_pptx_effective_typography_and_one_line_titles(self):
        with zipfile.ZipFile(ART / "10_Knowledge_Transfer_Deck.pptx") as zf:
            for index in range(1, 15):
                root = ET.fromstring(zf.read(f"ppt/slides/slide{index}.xml"))
                for shape in root.iter(P + "sp"):
                    nv = shape.find(P + "nvSpPr/" + P + "cNvPr")
                    name = nv.get("name") if nv is not None else ""
                    value = "".join(x.text or "" for x in shape.iter(A + "t"))
                    if not value or name == "slide-no":
                        continue
                    sizes = [int(x.get("sz")) / 100 for x in shape.iter() if x.tag in (A + "rPr", A + "defRPr") and x.get("sz")]
                    scales = [int(x.get("fontScale", "100000")) / 100000 for x in shape.iter(A + "normAutofit")]
                    effective = min(sizes) * min(scales or [1])
                    minimum = 50 if name == "cover-title" else 35 if name == "slide-title" else 24 if name.endswith("-title") or name.startswith("label-") else 16
                    self.assertGreaterEqual(effective, minimum, (index, name, effective, value))
                    if name == "slide-title":
                        self.assertNotIn("\n", value)

    def test_publication_rows_have_exact_resolvable_hash_contract(self):
        wb = load_workbook(ART / "11_Publication_Tables.xlsx", data_only=False)
        for sheet_name in ("Chapter 4 Tables", "Paper 2 Methods", "Paper 2 Results", "Variable Definitions", "Citation Mapping", "Figure Registry", "Evidence Index"):
            ws = wb[sheet_name]
            headers = [ws.cell(4, col).value for col in range(1, ws.max_column + 1)]
            for required in ("Exact source path", "SHA-256", "Version / status"):
                self.assertIn(required, headers, sheet_name)
            indexes = {name: headers.index(name) + 1 for name in headers}
            for row in range(5, ws.max_row + 1):
                source = ws.cell(row, indexes["Exact source path"]).value
                checksum = ws.cell(row, indexes["SHA-256"]).value
                status = ws.cell(row, indexes["Version / status"]).value
                self.assertTrue(source and "*" not in source and ";" not in source, (sheet_name, row, source))
                self.assertTrue((ROOT / source).is_file(), (sheet_name, row, source))
                self.assertRegex(checksum or "", r"^[0-9a-f]{64}$")
                self.assertTrue(status, (sheet_name, row))

    def test_manuals_use_exact_navigation_routes(self):
        for name in ("10_End_User_Manual.docx", "10_Research_Admin_Manual.docx"):
            text = doc_text(name)
            self.assertIn("Profile > Manage Farmers", text)
            self.assertIn("History > per-record CSV", text)
            self.assertIn("History > all-visible CSV", text)
            self.assertIn("Profile > Export Model Training Data", text)
            self.assertIn("two CSV files", text)
            self.assertIn("lib/screens/main/profile_tab.dart", text)
            self.assertIn("lib/screens/main/history_tab.dart", text)
            self.assertIn("lib/screens/main/training_data_export_screen.dart", text)

    def test_developer_manual_is_self_contained(self):
        text = doc_text("10_Developer_Handover_Manual.docx")
        for value in ("Flutter 3.41.9 stable", "Dart 3.11.5", "Xcode 26.6", "17F113", "macOS 26.6.2 arm64", "assets/models/xgboost_model.onnx", "dbedb2ab5ce57f3af0cd620e956f30ef34a64beaea493385afd3d27994002efc"):
            self.assertIn(value, text)

    def test_deck_sources_are_exact_and_no_future_artifact(self):
        with zipfile.ZipFile(ART / "10_Knowledge_Transfer_Deck.pptx") as zf:
            xml = "\n".join(zf.read(name).decode("utf-8", "ignore") for name in zf.namelist() if name.endswith(".xml"))
        self.assertNotIn("12 sign-off controls", xml)
        self.assertNotRegex(xml, r"Task 2 raw logs|03 diagrams|Task 5 reports")
        notes = re.findall(r"\[Sources\](.*?)\[/Sources\]", xml, re.S)
        self.assertEqual(len(notes), 14)
        for note in notes:
            self.assertRegex(note, r"[A-Za-z0-9_./+()-]+ \| [0-9a-f]{64} \| (?:1\.3\.11\+28|historical|controlled draft)")

    def test_security_scan_is_portable_expanded_and_triaged(self):
        builder = (HERE / "build_task7_documents.py").read_text()
        self.assertNotIn("source-gonwr134", builder)
        report = json.loads((ROOT / "manifests/task7_offline_security_inspection.json").read_text())
        self.assertIn("scan_scope_counts", report)
        self.assertIn("generic_secret_assignment_triage", report)
        self.assertEqual(report["generic_secret_assignment_triage"]["untriaged"], 0)
        self.assertGreater(report["scan_scope_counts"]["office_archive_members"], 0)
        self.assertGreater(report["scan_scope_counts"]["manifest_files"], 0)

    def test_office_metadata_is_neutral(self):
        for name in ("09_Security_Privacy_and_Data_Protection_Report.docx", "10_Knowledge_Transfer_Deck.pptx"):
            with zipfile.ZipFile(ART / name) as zf:
                props = "\n".join(zf.read(x).decode("utf-8", "ignore") for x in zf.namelist() if x.startswith("docProps/"))
            self.assertNotRegex(props, r"(?i)python-docx|Walnut Exporter|artifact-tool")
            self.assertIn("SookTa Project", props)

    def test_security_summary_formula_links_risk_register(self):
        wb = load_workbook(ART / "09_Security_and_Access_Control_Matrices.xlsx", data_only=False)
        self.assertEqual(len(wb.sheetnames), 15)
        formula = wb["Control Summary"]["B9"].value
        self.assertIn("'Risk Register'!E5:E9", formula)
        self.assertIn("COUNTIF", formula)

    def test_visual_recorder_consumes_external_decision(self):
        source = (HERE / "record_task7_visual_qa.py").read_text()
        self.assertIn("task7_external_visual_decision.json", source)
        self.assertNotIn('"decision": "pass"', source)


if __name__ == "__main__":
    unittest.main()
