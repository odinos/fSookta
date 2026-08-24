"""Mutation tests proving Task 7 verification is independent of its builders."""
from __future__ import annotations

import json
import os
import shutil
import sys
import tarfile
import tempfile
import unittest
import zipfile
from pathlib import Path

from openpyxl import load_workbook

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE))
import verify_task7_artifacts as verifier  # noqa: E402


ROOT = Path(os.environ.get("SOOKTA_HANDOVER_ROOT", "/private/tmp/fsookta-final-handover"))


def link_tree(source: Path, target: Path, copied: set[Path]) -> None:
    """Link a fixture tree, copying only files that a test is allowed to mutate."""
    target.mkdir(parents=True, exist_ok=True)
    for path in source.rglob("*"):
        relative = path.relative_to(source)
        destination = target / relative
        if path.is_dir():
            destination.mkdir(parents=True, exist_ok=True)
        elif relative in copied:
            destination.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(path, destination)
        else:
            destination.parent.mkdir(parents=True, exist_ok=True)
            destination.symlink_to(path)


class Fixture:
    def __init__(self, copies: dict[str, set[str]], linked_tops: tuple[str, ...]):
        self._temporary = tempfile.TemporaryDirectory(prefix="task7-mutation-")
        self.root = Path(self._temporary.name)
        for top in linked_tops:
            source = ROOT / top
            requested = {Path(value) for value in copies.get(top, set())}
            if requested:
                link_tree(source, self.root / top, requested)
            else:
                (self.root / top).symlink_to(source, target_is_directory=True)

    def close(self) -> None:
        self._temporary.cleanup()


def mutate_zip(path: Path, member: str, old: bytes, new: bytes) -> None:
    replacement = path.with_suffix(path.suffix + ".mutation")
    with zipfile.ZipFile(path) as source, zipfile.ZipFile(replacement, "w") as target:
        found = False
        for info in source.infolist():
            data = source.read(info.filename)
            if info.filename == member:
                if old not in data:
                    raise AssertionError(f"mutation token absent: {old!r}")
                data = data.replace(old, new, 1)
                found = True
            target.writestr(info, data)
    if not found:
        raise AssertionError(f"archive member absent: {member}")
    os.replace(replacement, path)


class Task7VerifierMutationTests(unittest.TestCase):
    def test_manual_navigation_claim_mutation_is_rejected(self):
        fixture = Fixture(
            {"artifacts": {"10_End_User_Manual.docx"}},
            ("artifacts", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts/10_End_User_Manual.docx"
        mutate_zip(path, "word/document.xml", b"Profile &gt; Manage Farmers", b"Profile &gt; Farmers")
        with self.assertRaises(AssertionError):
            verifier.verify_documents(fixture.root)

    def test_publication_hash_cell_mutation_is_rejected(self):
        fixture = Fixture(
            {"artifacts": {"11_Publication_Tables.xlsx"}},
            ("artifacts", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts/11_Publication_Tables.xlsx"
        workbook = load_workbook(path)
        sheet = workbook["Evidence Index"]
        headers = {sheet.cell(4, column).value: column for column in range(1, sheet.max_column + 1)}
        sheet.cell(5, headers["SHA-256"]).value = "0" * 64
        workbook.save(path)
        with self.assertRaises(AssertionError):
            verifier.verify_workbooks(fixture.root)

    def test_security_formula_mutation_is_rejected(self):
        fixture = Fixture(
            {"artifacts": {"09_Security_and_Access_Control_Matrices.xlsx"}},
            ("artifacts", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts/09_Security_and_Access_Control_Matrices.xlsx"
        workbook = load_workbook(path)
        workbook["Control Summary"]["B9"] = "OPEN ACTIONS"
        workbook.save(path)
        with self.assertRaises(AssertionError):
            verifier.verify_workbooks(fixture.root)

    def test_deck_source_note_mutation_is_rejected(self):
        fixture = Fixture(
            {"artifacts": {"10_Knowledge_Transfer_Deck.pptx"}},
            ("artifacts", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts/10_Knowledge_Transfer_Deck.pptx"
        mutate_zip(path, "ppt/notesSlides/notesSlide1.xml", b"571746f51b3e80a1847cdc1e5b79ad880dd4809b533ddcbbb3448e796918e6c2", b"071746f51b3e80a1847cdc1e5b79ad880dd4809b533ddcbbb3448e796918e6c2")
        with self.assertRaises(AssertionError):
            verifier.verify_deck(fixture.root)

    def test_deck_effective_typography_mutation_is_rejected(self):
        fixture = Fixture(
            {"artifacts": {"10_Knowledge_Transfer_Deck.pptx"}},
            ("artifacts", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts/10_Knowledge_Transfer_Deck.pptx"
        mutate_zip(path, "ppt/slides/slide1.xml", b'sz="5100"', b'sz="1200"')
        with self.assertRaises(AssertionError):
            verifier.verify_deck(fixture.root)

    def test_deck_generator_metadata_mutation_is_rejected(self):
        fixture = Fixture(
            {"artifacts": {"10_Knowledge_Transfer_Deck.pptx"}},
            ("artifacts", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts/10_Knowledge_Transfer_Deck.pptx"
        mutate_zip(path, "docProps/core.xml", b"SookTa Project", b"Walnut Exporter")
        with self.assertRaises(AssertionError):
            verifier.verify_deck(fixture.root)

    def test_security_secret_and_triage_mutations_are_rejected(self):
        fixture = Fixture(
            {"manifests": {"task7_offline_security_inspection.json"}},
            ("manifests",),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "manifests/task7_offline_security_inspection.json"
        payload = json.loads(path.read_text())
        payload["generic_secret_assignment_triage"]["untriaged"] = 1
        payload["mutation_probe"] = "AKIA0000000000000000"
        path.write_text(json.dumps(payload))
        with self.assertRaises(AssertionError):
            verifier.verify_security(fixture.root)

    def test_source_alias_substitution_is_rejected(self):
        fixture = Fixture(
            {"evidence": {"task7-source/pubspec.yaml"}},
            ("authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "evidence/task7-source/pubspec.yaml"
        path.write_text(path.read_text() + "\n# substituted source with unchanged version\n")
        with self.assertRaises(AssertionError):
            verifier.source_root(fixture.root)

    def test_archive_secret_and_participant_injection_is_rejected(self):
        name = "SookTa-1.3.11+28-source-snapshot.tar.gz"
        fixture = Fixture(
            {"archives": {name}},
            ("artifacts", "archives", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "archives" / name
        replacement = path.with_suffix(".mutation.tar.gz")
        with tarfile.open(path, "r:gz") as source, tarfile.open(replacement, "w:gz") as target:
            for member in source.getmembers():
                target.addfile(member, source.extractfile(member) if member.isfile() else None)
            payload = b"AKIA0000000000000000 participant@example.invalid"
            member = tarfile.TarInfo("mutation-probe.txt")
            member.size = len(payload)
            import io
            target.addfile(member, io.BytesIO(payload))
        os.replace(replacement, path)
        with self.assertRaises(AssertionError):
            verifier.verify_security(fixture.root)

    def test_office_secret_and_participant_injection_is_rejected(self):
        name = "08_UAT_Field_Test_and_Usability_Package.xlsx"
        fixture = Fixture(
            {"artifacts": {name}},
            ("artifacts", "archives", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts" / name
        mutate_zip(path, "xl/worksheets/sheet1.xml", b"</worksheet>", b"<!-- AKIA0000000000000000 participant@example.invalid --></worksheet>")
        with self.assertRaises(AssertionError):
            verifier.verify_security(fixture.root)

    def test_manifest_secret_and_participant_injection_is_rejected(self):
        name = "task6_verification_summary.json"
        fixture = Fixture(
            {"manifests": {name}},
            ("artifacts", "archives", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "manifests" / name
        payload = json.loads(path.read_text())
        payload["mutation_probe"] = "AKIA0000000000000000 participant@example.invalid"
        path.write_text(json.dumps(payload))
        with self.assertRaises(AssertionError):
            verifier.verify_security(fixture.root)

    def test_uncontrolled_status_mutation_is_rejected(self):
        fixture = Fixture(
            {"artifacts": {"11_Publication_Tables.xlsx"}},
            ("artifacts", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts/11_Publication_Tables.xlsx"
        workbook = load_workbook(path)
        workbook["Paper 2 Methods"]["D5"] = "Available"
        workbook.save(path)
        with self.assertRaises(AssertionError):
            verifier.verify_workbooks(fixture.root)

    def test_xlsx_generator_metadata_mutation_is_rejected(self):
        fixture = Fixture(
            {"artifacts": {"11_Publication_Tables.xlsx"}},
            ("artifacts", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts/11_Publication_Tables.xlsx"
        mutate_zip(path, "docProps/core.xml", b"SookTa Project", b"openpyxl")
        with self.assertRaises(AssertionError):
            verifier.verify_workbooks(fixture.root)

    def test_semantically_wrong_but_valid_evidence_path_is_rejected(self):
        fixture = Fixture(
            {"artifacts": {"11_Publication_Tables.xlsx"}},
            ("artifacts", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts/11_Publication_Tables.xlsx"
        workbook = load_workbook(path)
        sheet = workbook["Paper 2 Methods"]
        row = next(r for r in range(5, sheet.max_row + 1) if sheet.cell(r, 1).value == "Static analysis")
        wrong = "evidence/flutter_test_1.3.11+28.log"
        sheet.cell(row, 5).value = wrong
        sheet.cell(row, 6).value = verifier.sha(fixture.root / wrong)
        workbook.save(path)
        with self.assertRaises(AssertionError):
            verifier.verify_workbooks(fixture.root)

    def test_current_task7_docx_participant_injection_is_rejected(self):
        name = "10_End_User_Manual.docx"
        fixture = Fixture(
            {"artifacts": {name}},
            ("artifacts", "archives", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts" / name
        mutate_zip(path, "word/document.xml", b"</w:document>", b"<!-- participant@example.invalid --></w:document>")
        with self.assertRaises(AssertionError):
            verifier.verify_security(fixture.root)

    def test_current_task7_xlsx_participant_injection_is_rejected(self):
        name = "11_Publication_Tables.xlsx"
        fixture = Fixture(
            {"artifacts": {name}},
            ("artifacts", "archives", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts" / name
        mutate_zip(path, "xl/worksheets/sheet1.xml", b"</worksheet>", b"<!-- participant@example.invalid --></worksheet>")
        with self.assertRaises(AssertionError):
            verifier.verify_security(fixture.root)

    def test_current_task7_pptx_participant_injection_is_rejected(self):
        name = "10_Knowledge_Transfer_Deck.pptx"
        fixture = Fixture(
            {"artifacts": {name}},
            ("artifacts", "archives", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "artifacts" / name
        mutate_zip(path, "ppt/slides/slide1.xml", b"</p:sld>", b"<!-- participant@example.invalid --></p:sld>")
        with self.assertRaises(AssertionError):
            verifier.verify_security(fixture.root)

    def test_nonself_task7_manifest_participant_injection_is_rejected(self):
        name = "task7_expected_visual_renders.json"
        fixture = Fixture(
            {"manifests": {name}},
            ("artifacts", "archives", "authoritative-materializations", "evidence", "manifests"),
        )
        self.addCleanup(fixture.close)
        path = fixture.root / "manifests" / name
        payload = json.loads(path.read_text())
        payload["mutation_probe"] = "participant@example.invalid"
        path.write_text(json.dumps(payload))
        with self.assertRaises(AssertionError):
            verifier.verify_security(fixture.root)


if __name__ == "__main__":
    unittest.main()
