#!/usr/bin/env python3
"""Build deterministic Task 3 source and DOCX handover artifacts.

The final artifacts are written outside Git.  This builder deliberately uses
only repository metadata at the pinned authoritative commit and safe Task 1-2
JSON summaries; it never reads signing material, credentials, or participant
data.
"""

from __future__ import annotations

import argparse
import gzip
import hashlib
import io
import json
import os
import re
import subprocess
import tarfile
from datetime import datetime, timezone
from pathlib import Path
from typing import Iterable

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_ALIGN_VERTICAL, WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


VERSION = "1.3.11+28"
COMMIT = "bf8867a2083357cb9d60915bf6c2233801f923d8"
TREE = "b4ed5fd0c061492c74dba356ed8a114b5f6621ba"
EVIDENCE_DATE = "2026-08-23"
ARCHIVE_MTIME = int(datetime(2026, 8, 23, tzinfo=timezone.utc).timestamp())
ROOT_NAME = f"SookTa-{VERSION}"
STAGING = Path("/private/tmp/fsookta-final-handover")


EXCLUDED_PREFIXES = (
    "data/research/", "docs/qa/",
    "docs/uat_evidence_", "docs/user_manual_v1_1_1_android/",
)
EXCLUDED_SEGMENTS = {
    ".dart_tool", ".gradle", ".idea", ".pub-cache", ".vscode", ".worktrees",
    "build", "cache", "caches", "coverage", "deriveddata", "node_modules",
    "outputs", "pods", "tmp", "vendor",
}
EXCLUDED_EXACT = {
    "android/app/google-services.json",
    "ios/Runner/GoogleService-Info.plist",
    "ios/firebase_app_id_file.json",
    "lib/firebase_options.dart",
    "assets/test_fixtures/video/portrait_h264_4s.mp4",
    "docs/Sookta_REBA_Logic_Evidence_Package_20260603.zip",
    "docs/Sookta_Research_Training_Dataset_REBA_ISO11228.xlsx",
}
EXCLUDED_SUFFIXES = (
    ".jks", ".keystore", ".mobileprovision", ".p8", ".p12", ".pfx", ".pem", ".key",
)
SENSITIVE_CONFIG_NAMES = {
    ".env", ".npmrc", ".pypirc", "credentials.json", "firebase_app_id_file.json",
    "firebase_options.dart", "google-services.json", "googleservice-info.plist",
    "key.properties", "secrets.json", "service-account.json", "service_account.json",
}


def run_git(repo: Path, *args: str, text: bool = True) -> str | bytes:
    return subprocess.check_output(["git", "-C", str(repo), *args], text=text)


def excluded(path: str) -> bool:
    lower = path.lower()
    name = Path(path).name.lower()
    parts = {part.lower() for part in Path(path).parts}
    if path in EXCLUDED_EXACT or path.startswith(EXCLUDED_PREFIXES):
        return True
    if parts & EXCLUDED_SEGMENTS:
        return True
    if name in SENSITIVE_CONFIG_NAMES or name == "local.properties":
        return True
    if name.startswith(".env.") or lower.endswith(EXCLUDED_SUFFIXES):
        return True
    if re.search(r"(^|/)(secret|credentials?|tokens?)(/|\.|$)", lower):
        return True
    return False


def source_paths(repo: Path) -> list[str]:
    raw = run_git(repo, "ls-tree", "-r", "--name-only", COMMIT)
    paths = [part for part in raw.splitlines() if part]
    return [path for path in paths if not excluded(path)]


def build_archive(repo: Path, output: Path, manifest_output: Path) -> dict:
    source_tar = run_git(repo, "archive", "--format=tar", COMMIT, text=False)
    source_stream = tarfile.open(fileobj=io.BytesIO(source_tar), mode="r:")
    members = {member.name: member for member in source_stream.getmembers() if member.isfile() and not excluded(member.name)}
    paths = sorted(members)
    output.parent.mkdir(parents=True, exist_ok=True)
    uncompressed = io.BytesIO()
    with tarfile.open(fileobj=uncompressed, mode="w", format=tarfile.PAX_FORMAT) as tf:
        for path in paths:
            extracted = source_stream.extractfile(members[path])
            if extracted is None:
                raise RuntimeError(f"Unable to read archived file: {path}")
            content = extracted.read()
            info = tarfile.TarInfo(f"{ROOT_NAME}/{path}")
            info.size = len(content)
            info.mtime = ARCHIVE_MTIME
            info.uid = info.gid = 0
            info.uname = info.gname = ""
            info.mode = 0o755 if path in {"android/gradlew"} or path.endswith((".sh", ".py")) else 0o644
            tf.addfile(info, io.BytesIO(content))
    with output.open("wb") as stream:
        with gzip.GzipFile(filename="", mode="wb", fileobj=stream, mtime=0, compresslevel=9) as gz:
            gz.write(uncompressed.getvalue())
    digest = hashlib.sha256(output.read_bytes()).hexdigest()
    manifest = {
        "schema_version": 1,
        "archive": str(output),
        "application_version": VERSION,
        "authoritative_commit": COMMIT,
        "git_tree_id": TREE,
        "evidence_date": EVIDENCE_DATE,
        "archive_root": ROOT_NAME,
        "file_count": len(paths),
        "byte_size": output.stat().st_size,
        "sha256": digest,
        "determinism": {"tar_mtime_utc": "2026-08-23T00:00:00Z", "gzip_mtime": 0, "uid_gid": 0},
        "exclusions": [
            "build outputs and dependency caches (including Pods)",
            "secret/signing assets and credential-like files",
            "Firebase/local machine configuration files",
            "research datasets/media, QA screenshots/logs, and historic UAT media",
        ],
        "paths": paths,
    }
    manifest_output.parent.mkdir(parents=True, exist_ok=True)
    manifest_output.write_text(json.dumps(manifest, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return manifest


def parse_dependencies(repo: Path, output: Path) -> dict:
    lock = parse_pubspec_lock(run_git(repo, "show", f"{COMMIT}:pubspec.lock"))
    records = []
    for name, item in sorted(lock.items()):
        records.append(dependency_record(name, item))
    for plugin, version in gradle_plugins(repo):
        records.append({
            "dependency": plugin,
            "resolved_version": version,
            "purpose": "Android build/plugin toolchain",
            "source_url": "https://plugins.gradle.org/ or vendor documentation",
            "license_identifier": "Human verification required",
            "license_text_source": "Vendor/plugin distribution license",
            "platform": "Android",
            "classification": "Development",
            "dependency_scope": "Gradle plugin",
            "restriction_notes": "Build-time component; verify redistribution/notice terms.",
            "evidence_source": f"android/settings.gradle at Git {COMMIT}",
        })
    for pod, version in pod_versions(repo):
        records.append({
            "dependency": pod,
            "resolved_version": version,
            "purpose": "iOS native transitive/runtime dependency",
            "source_url": f"https://cocoapods.org/pods/{pod.split('/')[0]}",
            "license_identifier": "Human verification required",
            "license_text_source": "Podspec license field and installed pod LICENSE/NOTICE",
            "platform": "iOS",
            "classification": "Runtime",
            "dependency_scope": "CocoaPods resolved dependency",
            "restriction_notes": "Confirm podspec/license text and notice obligations before distribution.",
            "evidence_source": f"ios/Podfile.lock at Git {COMMIT}",
        })
    payload = {"schema_version": 1, "baseline": {"version": VERSION, "commit": COMMIT}, "records": records}
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(payload, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    return payload


def dependency_record(name: str, item: dict) -> dict:
    """Map one pub lock record without inventing source, license, or graph scope."""
    description = item.get("description") or {}
    source = str(item.get("source", "unknown"))
    scope = str(item.get("dependency", "transitive"))
    local_path = str(description.get("path", "")).strip()
    git_url = str(description.get("url", "")).strip()
    git_ref = str(description.get("ref", "")).strip()

    if source == "hosted":
        source_url = f"https://pub.dev/packages/{name}"
        license_source = f"{source_url} (License tab/package archive)"
        evidence_detail = source_url
        license_identifier = "Human verification required"
        review_status = "Human verification required"
    elif source == "sdk":
        source_url = "Flutter SDK"
        license_source = "Flutter SDK LICENSE and package source at the pinned toolchain version"
        evidence_detail = f"sdk:{description.get('name') or description.get('url') or name}"
        license_identifier = "Human verification required"
        review_status = "Human verification required"
    elif source == "git":
        repository = git_url or "Git repository not recorded"
        revision = git_ref or "resolved ref not recorded"
        suffix = f" (path {local_path})" if local_path else ""
        source_url = f"{repository} at {revision}{suffix}"
        license_source = f"License file in {repository} at {revision}{suffix}"
        evidence_detail = source_url
        license_identifier = "Human verification required"
        review_status = "Human verification required"
    elif source == "path":
        actual_path = local_path or "path not recorded"
        source_url = f"Repository path {actual_path} at authoritative commit"
        evidence_detail = source_url
        if actual_path.rstrip("/") == "third_party/onnxruntime_16kb":
            license_source = f"Git {COMMIT}:third_party/onnxruntime_16kb/LICENSE"
            license_identifier = "MIT"
            review_status = "Recorded"
        else:
            license_source = f"License file under repository path {actual_path}; human verification required"
            license_identifier = "Human verification required"
            review_status = "Human verification required"
    else:
        source_url = f"Source type {source or 'unknown'} in pubspec.lock; human verification required"
        license_source = f"Authoritative license source for {name} not resolved; human verification required"
        evidence_detail = source_url
        license_identifier = "Human verification required"
        review_status = "Human verification required"

    if scope == "direct main":
        classification = "Runtime"
    elif scope == "direct dev":
        classification = "Development"
    else:
        classification = "Unresolved - dependency graph review required"

    purpose = purpose_for(name)
    if scope == "transitive" and purpose == "Resolved transitive package support":
        purpose = "Resolved transitive package; runtime/development scope unresolved"
    restriction = "Confirm license text and notice obligations before external distribution."
    if review_status == "Recorded":
        restriction = "Retain the tracked copyright and permission notice in distributions; native runtime notices still require review."
    return {
        "dependency": name,
        "resolved_version": str(item.get("version", "Unspecified")),
        "purpose": purpose,
        "source_type": source,
        "source_url": source_url,
        "license_identifier": license_identifier,
        "license_text_source": license_source,
        "platform": platform_for(name),
        "classification": classification,
        "dependency_scope": scope,
        "restriction_notes": restriction,
        "evidence_source": f"pubspec.lock at Git {COMMIT}; {evidence_detail}",
        "review_status": review_status,
    }


def parse_pubspec_lock(text: str) -> dict[str, dict]:
    """Parse the stable package blocks needed from a Dart pub lockfile."""
    records: dict[str, dict] = {}
    current: str | None = None
    in_description = False
    for line in text.splitlines():
        package = re.match(r"^  ([A-Za-z0-9_+.-]+):$", line)
        if package:
            current = package.group(1)
            records[current] = {"description": {}}
            in_description = False
            continue
        if current is None:
            continue
        if re.match(r"^    description:$", line):
            in_description = True
            continue
        field = re.match(r"^    (dependency|source|version):\s+(.+)$", line)
        if field:
            value = field.group(2).strip().strip('"')
            records[current][field.group(1)] = value
            in_description = False
            continue
        desc_field = re.match(r"^      (name|url|path|ref|resolved-ref|sha256):\s+(.+)$", line)
        if in_description and desc_field:
            records[current]["description"][desc_field.group(1)] = desc_field.group(2).strip().strip('"')
    if not records:
        raise ValueError("pubspec.lock parser found no package records")
    return records


def purpose_for(name: str) -> str:
    mapping = {
        "camera": "Camera capture", "image_picker": "Gallery image selection", "image": "Image processing",
        "flutter_tts": "Text-to-speech", "tflite_flutter": "On-device TensorFlow Lite inference",
        "onnxruntime": "On-device ONNX inference", "shared_preferences": "Local settings/history storage",
        "share_plus": "User-initiated file sharing", "firebase_core": "Optional Firebase bootstrap",
        "firebase_analytics": "Opt-in analytics", "firebase_crashlytics": "Opt-in crash reporting",
        "ml_algo": "Local machine-learning utilities", "ml_dataframe": "ML dataframe support",
        "path_provider": "Platform storage locations", "integration_test": "Integration test harness",
        "flutter_test": "Unit/widget test harness", "flutter_lints": "Static-analysis rules",
    }
    return mapping.get(name, "Resolved transitive package support")


def platform_for(name: str) -> str:
    for suffix, value in (("_android", "Android"), ("_ios", "iOS"), ("_web", "Web"), ("_linux", "Linux"), ("_windows", "Windows"), ("_macos", "macOS"), ("_foundation", "iOS/macOS")):
        if suffix in name:
            return value
    return "Flutter/cross-platform"


def gradle_plugins(repo: Path) -> list[tuple[str, str]]:
    text = run_git(repo, "show", f"{COMMIT}:android/settings.gradle")
    return re.findall(r'id\s+"([^"]+)"\s+version\s+"([^"]+)"', text)


def parse_pod_lock_entry(line: str) -> tuple[str, str] | None:
    """Parse one top-level CocoaPods lock entry, including YAML-quoted scalars."""
    if not line.startswith("  - "):
        return None
    scalar = line[4:].strip()
    if scalar.endswith(":"):
        scalar = scalar[:-1].rstrip()
    if scalar.startswith(('"', "'")):
        quote = scalar[0]
        if not scalar.endswith(quote):
            raise ValueError(f"malformed quoted CocoaPods lock entry: {line!r}")
        if quote == '"':
            try:
                scalar = json.loads(scalar)
            except json.JSONDecodeError as exc:
                raise ValueError(f"invalid quoted CocoaPods lock entry: {line!r}") from exc
        else:
            scalar = scalar[1:-1].replace("''", "'")
    match = re.fullmatch(r"(.+?) \(([^()]+)\)", scalar)
    if not match:
        return None
    name, version = (value.strip() for value in match.groups())
    if not name or name.startswith(('"', "'")) or name.endswith(('"', "'")):
        raise ValueError(f"invalid CocoaPods dependency name: {name!r}")
    return name, version


def pod_versions(repo: Path) -> list[tuple[str, str]]:
    text = run_git(repo, "show", f"{COMMIT}:ios/Podfile.lock")
    pods = []
    in_pods = False
    seen = set()
    for line in text.splitlines():
        if line == "PODS:":
            in_pods = True
            continue
        if in_pods and line and not line.startswith(" "):
            break
        entry = parse_pod_lock_entry(line)
        if entry and entry[0] not in seen:
            seen.add(entry[0])
            pods.append(entry)
    return pods


def set_cell_shading(cell, fill: str) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd")) or OxmlElement("w:shd")
    shd.set(qn("w:fill"), fill)
    if shd.getparent() is None:
        tc_pr.append(shd)


def set_cell_margins(cell, top=80, start=120, bottom=80, end=120) -> None:
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = tc_pr.first_child_found_in("w:tcMar")
    if tc_mar is None:
        tc_mar = OxmlElement("w:tcMar")
        tc_pr.append(tc_mar)
    for edge, value in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        element = tc_mar.find(qn(f"w:{edge}")) or OxmlElement(f"w:{edge}")
        element.set(qn("w:w"), str(value)); element.set(qn("w:type"), "dxa")
        if element.getparent() is None: tc_mar.append(element)


def set_table_geometry(table, widths: list[int], indent: int = 0) -> None:
    table.autofit = False
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    tbl_pr = table._tbl.tblPr
    for tag, attrs in (("tblW", {"w": str(sum(widths)), "type": "dxa"}), ("tblInd", {"w": str(indent), "type": "dxa"})):
        node = tbl_pr.find(qn(f"w:{tag}"))
        if node is None:
            node = OxmlElement(f"w:{tag}")
        for key, value in attrs.items(): node.set(qn(f"w:{key}"), value)
        if node.getparent() is None: tbl_pr.append(node)
    grid = table._tbl.tblGrid
    for child in list(grid): grid.remove(child)
    for width in widths:
        col = OxmlElement("w:gridCol"); col.set(qn("w:w"), str(width)); grid.append(col)
    for row in table.rows:
        for cell, width in zip(row.cells, widths):
            tc_w = cell._tc.get_or_add_tcPr().find(qn("w:tcW"))
            if tc_w is None:
                tc_w = OxmlElement("w:tcW")
            tc_w.set(qn("w:w"), str(width)); tc_w.set(qn("w:type"), "dxa")
            if tc_w.getparent() is None: cell._tc.get_or_add_tcPr().append(tc_w)
            set_cell_margins(cell)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def add_numbering(doc: Document, fmt: str, text: str) -> int:
    numbering = doc.part.numbering_part.element
    abstract_ids = [int(x.get(qn("w:abstractNumId"))) for x in numbering.findall(qn("w:abstractNum"))]
    num_ids = [int(x.get(qn("w:numId"))) for x in numbering.findall(qn("w:num"))]
    abstract_id = max(abstract_ids, default=0) + 1
    num_id = max(num_ids, default=0) + 1
    abstract = OxmlElement("w:abstractNum"); abstract.set(qn("w:abstractNumId"), str(abstract_id))
    lvl = OxmlElement("w:lvl"); lvl.set(qn("w:ilvl"), "0")
    start = OxmlElement("w:start"); start.set(qn("w:val"), "1")
    num_fmt = OxmlElement("w:numFmt"); num_fmt.set(qn("w:val"), fmt)
    lvl_text = OxmlElement("w:lvlText"); lvl_text.set(qn("w:val"), text)
    ppr = OxmlElement("w:pPr"); tabs = OxmlElement("w:tabs"); tab = OxmlElement("w:tab")
    tab.set(qn("w:val"), "num"); tab.set(qn("w:pos"), "720"); tabs.append(tab)
    ind = OxmlElement("w:ind"); ind.set(qn("w:left"), "720"); ind.set(qn("w:hanging"), "360")
    spacing = OxmlElement("w:spacing"); spacing.set(qn("w:after"), "80"); spacing.set(qn("w:line"), "276"); spacing.set(qn("w:lineRule"), "auto")
    ppr.extend([tabs, ind, spacing]); lvl.extend([start, num_fmt, lvl_text, ppr]); abstract.append(lvl); numbering.append(abstract)
    num = OxmlElement("w:num"); num.set(qn("w:numId"), str(num_id)); ref = OxmlElement("w:abstractNumId"); ref.set(qn("w:val"), str(abstract_id)); num.append(ref); numbering.append(num)
    return num_id


def apply_num(paragraph, num_id: int) -> None:
    ppr = paragraph._p.get_or_add_pPr(); num_pr = OxmlElement("w:numPr")
    ilvl = OxmlElement("w:ilvl"); ilvl.set(qn("w:val"), "0")
    num = OxmlElement("w:numId"); num.set(qn("w:val"), str(num_id))
    num_pr.extend([ilvl, num]); ppr.append(num_pr)


def style_document(doc: Document) -> tuple[int, int]:
    section = doc.sections[0]
    section.page_width = Inches(8.5); section.page_height = Inches(11)
    section.top_margin = section.right_margin = section.bottom_margin = section.left_margin = Inches(1)
    section.header_distance = section.footer_distance = Inches(0.492)
    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"; normal.font.size = Pt(11); normal.font.color.rgb = RGBColor(0, 0, 0)
    normal._element.rPr.rFonts.set(qn("w:ascii"), "Arial"); normal._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    normal.paragraph_format.space_before = Pt(0); normal.paragraph_format.space_after = Pt(8); normal.paragraph_format.line_spacing = 1.15
    normal.paragraph_format.keep_together = True
    for name, size, before, after, color in (("Heading 1",20,20,6,"000000"),("Heading 2",16,18,6,"000000"),("Heading 3",14,16,4,"434343")):
        style = styles[name]; style.font.name = "Arial"; style.font.size = Pt(size); style.font.bold = False; style.font.color.rgb = RGBColor.from_string(color)
        style._element.rPr.rFonts.set(qn("w:ascii"), "Arial"); style._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
        style.paragraph_format.space_before = Pt(before); style.paragraph_format.space_after = Pt(after); style.paragraph_format.keep_with_next = True
    doc.core_properties.author = "SookTa handover package"
    doc.core_properties.last_modified_by = "SookTa handover package"
    doc.core_properties.created = datetime(2026, 8, 23, tzinfo=timezone.utc)
    doc.core_properties.modified = datetime(2026, 8, 23, tzinfo=timezone.utc)
    return add_numbering(doc, "bullet", "●"), add_numbering(doc, "decimal", "%1.")


def add_title(doc: Document, title: str, subtitle: str) -> None:
    paragraph = doc.add_paragraph()
    paragraph.paragraph_format.space_before = Pt(0); paragraph.paragraph_format.space_after = Pt(3)
    run = paragraph.add_run(title); run.font.name = "Arial"; run.font.size = Pt(26); run.font.bold = False; run.font.color.rgb = RGBColor(0, 0, 0)
    run._element.rPr.rFonts.set(qn("w:ascii"), "Arial"); run._element.rPr.rFonts.set(qn("w:hAnsi"), "Arial")
    p = doc.add_paragraph(subtitle); p.paragraph_format.space_after = Pt(12)
    for r in p.runs: r.font.color.rgb = RGBColor.from_string("555555")


def add_bullets(doc: Document, items: Iterable[str], num_id: int) -> None:
    for item in items:
        p = doc.add_paragraph(item); apply_num(p, num_id)


def add_table(doc: Document, headers: list[str], rows: list[list[str]], widths: list[int]) -> None:
    table = doc.add_table(rows=1, cols=len(headers))
    table.style = "Table Grid"
    for idx, header in enumerate(headers):
        cell = table.rows[0].cells[idx]; cell.text = header
        for run in cell.paragraphs[0].runs: run.bold = True
        set_cell_shading(cell, "F5F5F5")
    for values in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(values): cells[idx].text = value
    set_table_geometry(table, widths, indent=0)
    doc.add_paragraph()


def build_release_doc(output: Path) -> None:
    doc = Document(); bullets, numbers = style_document(doc)
    add_title(doc, "Final Release and Scope Closure", f"SookTa {VERSION} | Evidence basis {EVIDENCE_DATE}")
    doc.add_heading("Release status", level=1)
    doc.add_paragraph("Evidence-backed freeze: the handover baseline is version 1.3.11+28 at Git commit " + COMMIT + ", tree " + TREE + ". The evidence date is the release-date basis for this package; it is not proof of App Store or Google Play publication.")
    add_table(doc, ["Control", "Status", "Evidence / required action"], [
        ["Version freeze", "Verified", "pubspec.yaml and lib/app/build_info.dart at bf8867a; full hash in Evidence references"],
        ["Static analysis", "Passed", "flutter analyze; exit 0; no issues found"],
        ["Automated tests", "Passed", "135 tests passed; 0 failures"],
        ["Android technical build", "Passed - non-production classification", "AAB build exit 0; upload-signing ownership unverified"],
        ["iOS technical build", "Passed - non-production classification", "IPA build exit 0; distribution signing and App Store ownership unverified"],
        ["Production delivery", "Pending Owner Action", "Provide signing/access evidence and complete store submission/acceptance"],
    ], [2100, 2400, 4860])
    doc.add_heading("Final feature list", level=1)
    add_bullets(doc, [
        "Thai and English onboarding, language selection, terms acknowledgement, and local research-profile setup.",
        "Farmer/participant profile creation, selection, editing, avatar selection, and local management without login.",
        "Coffee-farming activity selection for transplanting, fertilizing, pesticide spraying, pruning, harvesting, and transport.",
        "Camera and gallery image acquisition with image-quality guidance, four-angle slot guidance, and multi-person rejection.",
        "On-device MoveNet Thunder pose estimation and normalized joint-feature extraction.",
        "REBA-based posture scoring for all activities, with ISO 11228 inputs where lifting, carrying, pushing, pulling, or repetition applies.",
        "Advisory-only XGBoost inference that does not replace official REBA/ISO scores.",
        "Before/after assessment flow, risk level, score explanation, body-risk map, and categorized risk-reduction recommendations.",
        "Farmer-friendly result summary plus expandable technical details and reference sources.",
        "Estimated economic-impact layer based on cost inputs; it is not a training label or clinical outcome.",
        "Local assessment history, detail view, per-farmer filtering, trend views, resumable drafts, and durable persistence.",
        "CSV/research export and user-initiated share flow, including assessment references and training-data export tools.",
        "Help, bundled user manual, references, profile editing, and text-to-speech support.",
        "Offline-first primary assessment; Firebase Analytics and Crashlytics are optional and disabled by default unless explicitly enabled at build time.",
        "Portrait-only mobile interface for iOS and Android, with platform-specific native build integration.",
    ], bullets)
    doc.add_heading("Original-scope closure", level=1)
    doc.add_paragraph("Status: Pending Owner Action. The governing requirements define the required closure fields, but the available evidence does not include an authoritative original contract scope mapped item-by-item to the final implementation. A historical contract exists in Drive metadata only and was not treated as final authority.")
    add_bullets(doc, [
        "Owner action: approve the original-scope baseline and identify each item as delivered, changed, deferred, or not developed.",
        "Owner action: provide reasons and approvals for every changed/deferred item.",
        "Researcher action: confirm whether the final feature list above reflects the intended live-system scope.",
    ], bullets)
    doc.add_heading("Later additions and request provenance", level=1)
    add_table(doc, ["Addition evidenced in Git", "Technical source", "Request source / rationale status"], [
        ["Durable history, confirmation, drafts, expanded breakdown/export", "Commits 2571910 through 031191a (2026-07-06)", "Requester and business rationale not evidenced - human confirmation required"],
        ["Image requirements and quality checks", "Commits d219681, 4e424b5, 1e395d0 (2026-07-07)", "Requester not evidenced; technical intent is visible in commit messages"],
        ["Farmer-friendly summaries and recommendation groups", "Commits 041ac06 through 98f60ea (2026-07-09 to 2026-07-12)", "Requester/rationale approval record required"],
        ["Platform parity and production-readiness gates", "Commits df6715a through 8b009df (2026-07-19)", "Technical issue provenance only; owner approval required"],
        ["Approved bilingual recommendation catalog", "Commits 754771d through eb4cc87 (2026-07-30)", "Approval is named in Git; approving person/source document not present in current evidence"],
        ["Official REBA score preservation and iOS simulator-slice rejection", "Commits cd21b6e and bf8867a (2026-08-12 to 2026-08-13)", "Technical defect fixes; request source not evidenced"],
    ], [3100, 2900, 3360])
    doc.add_heading("Known limitations and known-bug status", level=1)
    add_bullets(doc, [
        "Not a medical device or diagnostic tool; results support risk communication, education, and research.",
        "Daily Injury Logistic coefficients remain a template and must be fitted and externally validated before interpreting probability.",
        "XGBoost output is advisory-only; official REBA/ISO scoring remains deterministic/rule-based.",
        "Primary workflows are offline/local. Cloud sync, remote database, web admin, and authenticated user accounts are not evidenced in the baseline.",
        "Telemetry is disabled by default and requires an explicit build flag; any enablement requires updated privacy/store declarations.",
        "Assessment quality depends on suitable images and visible pose landmarks; multiple-person images are rejected.",
        "The tracked README reports an older current version in its status prose while pubspec/build constants are 1.3.11+28; treat pubspec/build constants as authoritative and update README after owner review.",
        "No reproducible application defect was identified by the 2026-08-23 analyze/test run. This is not a claim that zero defects remain; device acceptance, signing, and store review remain open.",
    ], bullets)
    doc.add_heading("Release notes - 1.3.11+28", level=1)
    add_bullets(doc, [
        "Advanced version name/build number from 1.3.10+27 to 1.3.11+28.",
        "Hardened the iOS build process to invalidate staged native assets when platform SDK metadata changes.",
        "Added release-artifact checks that reject simulator Mach-O slices and verify code-signing structure.",
        "Preserved official REBA scoring changes carried forward from the 1.3.10+27 baseline.",
        "Reproduced dependency resolution, static analysis, 135 automated tests, Android AAB, iOS IPA, Android debug fallback, and iOS no-codesign fallback from the authoritative commit.",
    ], bullets)
    doc.add_heading("Acceptance actions", level=1)
    add_bullets(doc, [
        "Developer/owner: provide Android upload-key ownership/transfer evidence without placing secrets in this package.",
        "Developer/owner: provide Apple distribution certificate/provisioning and App Store Connect access evidence through approved secure channels.",
        "Researcher: review original scope, later-addition provenance, known limitations, and final feature list.",
        "Authorized parties: sign the final acceptance and rights documents; this unsigned document does not replace those approvals.",
    ], numbers)
    doc.add_heading("Evidence references", level=1)
    add_bullets(doc, [
        f"Git commit {COMMIT}; tree {TREE}; version {VERSION}.",
        "/private/tmp/fsookta-final-handover/evidence/final_source_metadata.txt and metadata_validation.json.",
        "Task 2 analyze/test/build logs under /private/tmp/fsookta-final-handover/evidence/.",
        "Task 1 requirements.json, evidence_map.json, source_inventory.json, and drive_source_inventory.json.",
    ], bullets)
    output.parent.mkdir(parents=True, exist_ok=True); doc.save(output)


def build_repository_doc(output: Path, repo: Path, archive_manifest: dict) -> None:
    doc = Document(); bullets, numbers = style_document(doc)
    add_title(doc, "Repository Ownership and IP Handover", f"SookTa {VERSION} | Evidence basis {EVIDENCE_DATE}")
    doc.add_heading("Repository authority and source package", level=1)
    add_table(doc, ["Item", "Verified record"], [
        ["Authoritative repository", "https://github.com/odinos/fSookta.git (origin recorded locally)"],
        ["Authoritative baseline", f"Commit {COMMIT}; tree {TREE}; version {VERSION}"],
        ["Full-history branch", "Remote branch origin/codex/fix-reba-score-release-1.3.8 contains the baseline; permanent branch designation requires owner confirmation"],
        ["Clean snapshot branch", "codex/client-handover-source-2026-08-23 at ce9c6dd; supplementary one-commit convenience snapshot only"],
        ["Source archive", f"{Path(archive_manifest['archive']).name}; {archive_manifest['file_count']} files; SHA-256 {archive_manifest['sha256']}"],
        ["Proposed final tag", f"sookta-v{VERSION}; proposal only - tag has not been created"],
    ], [2300, 7060])
    doc.add_paragraph("Authority rule: retain the full Git history and authoritative commit. Do not substitute the clean one-commit snapshot or source archive for the full-history repository.")
    doc.add_heading("Branches and preservation actions", level=1)
    add_bullets(doc, [
        "Preserve the repository default branch and all existing full-history branches/tags before transfer.",
        f"Protect the authoritative commit {COMMIT} from force-push or history rewriting.",
        f"After owner approval, create an annotated tag such as sookta-v{VERSION} that points exactly to the authoritative commit, then verify the tag object and remote ref.",
        "Keep the clean snapshot branch and deterministic archive labeled supplementary to avoid provenance confusion.",
        "Export repository access/transfer evidence separately; no access tokens or credentials belong in this package.",
    ], bullets)
    doc.add_heading("Setup, build, and deploy", level=1)
    for heading, steps in [
        ("1. Prepare", ["Install Flutter 3.41.9 stable/Dart 3.11.5 or validate with the project SDK constraints.", "For iOS install Xcode 26.6 and CocoaPods; for Android install a compatible Android SDK and Java 17.", "Provide Firebase configuration through the approved owner-controlled channel; the source archive intentionally excludes local Firebase configuration."]),
        ("2. Resolve and test", ["Run flutter pub get.", "Run flutter analyze --no-pub.", "Run flutter test --no-pub; the reference run passed 135 tests.", "After native plugin changes, run pod install from ios/."]),
        ("3. Technical builds", ["Android: flutter build appbundle --release.", "iOS: flutter build ipa --release.", "These commands prove technical buildability only until production signing ownership and store access are evidenced."]),
        ("4. Production deployment", ["Android: supply the owner-controlled upload keystore through a secure channel, configure android/key.properties locally, verify signer, then upload to Google Play Console.", "iOS: use the owner-controlled distribution identity/provisioning profile and App Store Connect role, verify archive/export, then upload through the approved release process.", "Complete store metadata, privacy declarations, device acceptance, staged release, and acceptance evidence before calling the build production-delivered."]),
    ]:
        doc.add_heading(heading, level=2); add_bullets(doc, steps, bullets)
        if heading == "4. Production deployment":
            final_deployment_bullet = doc.paragraphs[-1]
            final_deployment_bullet.paragraph_format.page_break_before = True
            final_deployment_bullet.paragraph_format.keep_together = True
            final_deployment_bullet.paragraph_format.space_before = Pt(6)
    doc.add_heading("Component ownership and restrictions", level=1)
    add_table(doc, ["Component", "Technical evidence", "Ownership / restriction status"], [
        ["Flutter application source, tests, tooling, documentation", f"Tracked at {COMMIT}", "Contractual ownership/assignment not verified - authorized parties must confirm"],
        ["MoveNet TFLite artifacts", "Bundled under assets/ml", "Third-party model/license and distribution terms require final license review"],
        ["ONNX Runtime and local wrapper override", "pubspec.lock, Podfile.lock, third_party/onnxruntime_16kb/LICENSE", "Third-party; local wrapper carries MIT text; runtime/native notices require register review"],
        ["XGBoost advisory model and project JSON models", "Bundled under assets/models and assets/ml", "Model/data provenance and permitted research/publication use require owner/researcher confirmation"],
        ["Firebase SDKs/services", "pubspec.lock, Podfile.lock, opt-in code paths", "Third-party service terms plus account/project ownership; access transfer pending"],
        ["Research data/media", "Inventoried by metadata only; excluded from archive", "Data governance, consent, access, and publication rights require researcher approval"],
        ["App-store accounts/signing assets", "No secret-bearing evidence inspected", "Ownership and transfer unverified; secure owner action required"],
    ], [2500, 3000, 3860])
    doc.add_heading("Rights statement - human action required", level=1)
    doc.add_paragraph("This document does not itself grant or confirm legal rights. The available technical evidence cannot establish contractual ownership, assignment, or permission to use, modify, continue development, distribute, conduct research, or publish. An authorized rights holder and the researcher must review the governing agreement, identify third-party restrictions, and sign the statement below or a counsel-approved equivalent.")
    doc.add_paragraph("Proposed confirmation for authorized review: Subject to the governing project agreement and identified third-party licenses, the researcher may use, modify, maintain, and continue development of the delivered project materials and may use authorized project outputs for research and publication. Any participant data/media remains governed by consent, ethics, privacy, and data-management approvals. [NOT EFFECTIVE UNTIL COMPLETED AND SIGNED BY AUTHORIZED PARTIES]")
    doc.add_heading("Access and handover record", level=1)
    add_bullets(doc, [
        "Repository owner/admin access: Pending Owner Action; capture invitation/role evidence and acceptance timestamp.",
        "Firebase, Apple Developer, App Store Connect, and Google Play Console: Pending Owner Action; see the separate access checklist.",
        "Signing materials and recovery controls: transfer only through an approved secure channel; record completion without embedding secret values.",
        "Third-party license review: Pending Owner/Legal Action; resolve every 'Human verification required' row in the license register.",
    ], bullets)
    doc.add_heading("Signature blocks", level=1)
    add_table(doc, ["Role", "Name / organization", "Signature", "Date"], [
        ["Authorized developer / rights holder", "________________", "________________", "__________"],
        ["Researcher / receiving owner", "________________", "________________", "__________"],
        ["Repository transfer administrator", "________________", "________________", "__________"],
    ], [2800, 2600, 2400, 1560])
    doc.add_heading("Evidence references", level=1)
    add_bullets(doc, [
        f"Git origin, refs, commit {COMMIT}, tree {TREE}, pubspec.yaml, pubspec.lock, ios/Podfile.lock, Android Gradle configuration, and tracked third-party LICENSE.",
        "Task 1 source and Drive metadata inventories; governing requirements Drive file 1OVOS13DQCNxK2PjAnmk3CYckcf2KCH0N.",
        "Task 2 evidence logs and safe metadata under /private/tmp/fsookta-final-handover/evidence/.",
    ], bullets)
    output.parent.mkdir(parents=True, exist_ok=True); doc.save(output)


def write_checksums(paths: Iterable[Path], output: Path) -> None:
    rows = []
    for path in sorted(paths, key=lambda p: p.name):
        rows.append(f"{hashlib.sha256(path.read_bytes()).hexdigest()}  {path.name}")
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text("\n".join(rows) + "\n", encoding="ascii")


def main() -> None:
    parser = argparse.ArgumentParser()
    parser.add_argument("--repository", type=Path, required=True)
    parser.add_argument("--staging", type=Path, default=STAGING)
    args = parser.parse_args()
    archive = args.staging / "archives" / f"SookTa-{VERSION}-source-snapshot.tar.gz"
    source_manifest = args.staging / "manifests" / "source_snapshot_manifest.json"
    dependency_json = args.staging / "working" / "task3" / "dependencies.json"
    raw_dir = args.staging / "working" / "task3" / "raw_docx"
    raw_dir.mkdir(parents=True, exist_ok=True)
    archive_manifest = build_archive(args.repository, archive, source_manifest)
    parse_dependencies(args.repository, dependency_json)
    build_release_doc(raw_dir / "01_Final_Release_and_Scope_Closure.docx")
    build_repository_doc(raw_dir / "02_Repository_Ownership_and_IP_Handover.docx", args.repository, archive_manifest)
    print(json.dumps({"archive_manifest": archive_manifest, "dependency_json": str(dependency_json), "raw_docx_dir": str(raw_dir)}, indent=2))


if __name__ == "__main__":
    main()
