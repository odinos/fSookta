#!/usr/bin/env python3
"""Recalculate Task 7 formulas in LibreOffice and persist cached results."""
import os,shutil,subprocess,tempfile
from pathlib import Path
from openpyxl import load_workbook
ROOT=Path(os.environ.get('SOOKTA_HANDOVER_ROOT','/private/tmp/fsookta-final-handover'));path=ROOT/'artifacts/09_Security_and_Access_Control_Matrices.xlsx'
soffice=os.environ.get('SOOKTA_SOFFICE') or shutil.which('soffice')
if not soffice:raise FileNotFoundError('SOOKTA_SOFFICE or soffice is required for formula recalculation')
with tempfile.TemporaryDirectory(prefix='task7-calc-') as out,tempfile.TemporaryDirectory(prefix='task7-lo-') as profile:
 subprocess.run([soffice,f'-env:UserInstallation=file://{profile}','--headless','--convert-to','xlsx','--outdir',out,str(path)],check=True,capture_output=True,text=True)
 converted=Path(out)/path.name;assert converted.is_file();shutil.copy2(converted,path)
assert load_workbook(path,data_only=False)['Control Summary']['B9'].value=='=IF(COUNTIF(\'Risk Register\'!E5:E9,"Open*")>0,"OPEN ACTIONS","NO OPEN ROWS")'
assert load_workbook(path,data_only=True)['Control Summary']['B9'].value=='OPEN ACTIONS'
print('recalculated=1 cached_status=OPEN ACTIONS')
