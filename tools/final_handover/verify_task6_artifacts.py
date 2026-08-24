#!/usr/bin/env python3
"""Canonical independent verifier for all Task 6 data, document and render surfaces."""
from __future__ import annotations

import hashlib
import json
import os
import re
import subprocess
import sys
from collections import Counter
from pathlib import Path
from datetime import date, datetime, time, timezone

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

BASELINE='1.3.11+28'; COMMIT='bf8867a2083357cb9d60915bf6c2233801f923d8'
RX=re.compile(r'^\d\d:\d\d \+(\d+)(?: -\d+)?: (.*)$')
PASS_HEADERS=['Case ID','Requirement ID','Test name','Category','Precondition / input','Expected','Actual','Status','Baseline / version','Tester category','Timestamp / date','Exact method / command','Raw evidence path','SHA-256','Limitations','Source path']
PASS_KEYS=('case_id','requirement_id','name','category','precondition_input','expected','actual','status','baseline','tester_category','timestamp','method','raw_path','sha256','limitations','source_path')
STATUS={'Pending Owner Action':137,'Exception Approval Required':12,'Pending Researcher Evidence':6}
EXPECTED_MODELS=json.loads(Path(__file__).with_name('task6_expected_artifact_models.json').read_text())

def sha(p): return hashlib.sha256(Path(p).read_bytes()).hexdigest()
def normalized(value):
 if isinstance(value,(datetime,date,time)): return value.isoformat()
 if isinstance(value,bytes): return value.hex()
 if value is None or isinstance(value,(str,int,float,bool)): return value
 return str(value)
def sheet_model(sheet):
 cells=[]
 for _,cell in sorted(sheet._cells.items()):
  cells.append({'coordinate':cell.coordinate,'value':normalized(cell.value),'data_type':cell.data_type,'number_format':cell.number_format,'style_id':cell.style_id,'hyperlink':cell.hyperlink.target if cell.hyperlink else None,'comment':cell.comment.text if cell.comment else None})
 validations=[]
 for item in sheet.data_validations.dataValidation:
  validations.append({key:normalized(getattr(item,key)) for key in ('sqref','type','operator','formula1','formula2','allowBlank','errorTitle','error','promptTitle','prompt','showErrorMessage','showInputMessage')})
 rows={str(index):{'height':row.height,'hidden':row.hidden,'outlineLevel':row.outlineLevel} for index,row in sorted(sheet.row_dimensions.items())}
 columns={index:{'width':column.width,'hidden':column.hidden,'outlineLevel':column.outlineLevel} for index,column in sorted(sheet.column_dimensions.items())}
 return {'title':sheet.title,'max_row':sheet.max_row,'max_column':sheet.max_column,'cells':cells,'merged_ranges':sorted(str(item) for item in sheet.merged_cells.ranges),'validations':validations,'row_dimensions':rows,'column_dimensions':columns,'freeze_panes':str(sheet.freeze_panes) if sheet.freeze_panes else None,'auto_filter':str(sheet.auto_filter.ref) if sheet.auto_filter.ref else None,'print_area':str(sheet.print_area) if sheet.print_area else None,'page_setup':{'orientation':sheet.page_setup.orientation,'paperSize':sheet.page_setup.paperSize,'fitToWidth':sheet.page_setup.fitToWidth,'fitToHeight':sheet.page_setup.fitToHeight}}
def model_sha(model): return hashlib.sha256(json.dumps(model,ensure_ascii=False,sort_keys=True,separators=(',',':')).encode()).hexdigest()
def verify_committed_artifact_expectations(root):
 assert EXPECTED_MODELS['schema_version']==1 and len(EXPECTED_MODELS['artifacts'])==10
 for name,expected in EXPECTED_MODELS['artifacts'].items():
  path=root/'artifacts'/name; assert path.is_file() and sha(path)==expected, f'unexpected artifact bytes: {name}'
def category(path,name):
 v=(path+' '+name).lower()
 for label,words in [('Algorithm / reference',('algorithm','reba','iso11228','pose','xgboost','logistic')),('Invalid / boundary',('invalid','required','malformed','missing','zero-height','error')),('UI / regression',('capture','screen','layout','responsive','widget','navigation')),('Data / persistence',('persist','draft','export','history','storage'))]:
  if any(x in v for x in words): return label
 return 'Functional / regression'
def reconstruct(text):
 out=[]; done=0; pending=None
 for line in text.splitlines():
  m=RX.fullmatch(line)
  if not m: continue
  n=int(m.group(1)); label=m.group(2).strip()
  if n!=done:
   assert n==done+1 and pending is not None
   path,name=pending; out.append({'ordinal':n,'case_id':f'AUTO-{n:03d}','source_path':path,'name':name,'category':category(path,name)}); done=n
  if label=='All tests passed!' or label.startswith('loading '): continue
  assert '.dart: ' in label
  prefix,name=label.rsplit('.dart: ',1); pending=(prefix.rsplit('/source/',1)[-1]+'.dart',name)
 assert done==len(out)==135 and out[-1]['name']=='clears the draft after a successful evaluation save'
 return out
def source_root(root): return sorted((root/'authoritative-materializations').glob('source-*/source'))[0]
def table_rows(sheet,width): return [tuple(sheet.cell(r,c).value for c in range(1,width+1)) for r in range(5,sheet.max_row+1)]
def expected_row(r):
 if r.get('status')!='PASS':
  return (r['case_id'],None,r['name'],r['category'],'Human/runtime evidence required','Final-version evidence meeting the plan',r['actual'],r['status'],None,None,None,None,None,None,r.get('limitations'),None)
 out=[]
 for k in PASS_KEYS:
  value=r.get(k,'')
  if k=='timestamp' and value:
   value=(datetime.fromisoformat(value.replace('Z','+00:00')).astimezone(timezone.utc).replace(tzinfo=None)-datetime(1899,12,30)).total_seconds()/86400
  out.append(None if value in (None,'') else value)
 return tuple(out)
def assert_join(sheet, rows):
 assert [sheet.cell(4,c).value for c in range(1,17)]==PASS_HEADERS
 actual=table_rows(sheet,16); expected=[expected_row(r) for r in rows]
 assert actual==expected, f'{sheet.title} normalized join mismatch'
 for row,src in zip(actual,rows,strict=True):
  if row[7]=='PASS':
   assert row[0]==src['case_id'] and row[8]==BASELINE and all(row[i] not in (None,'') for i in (0,2,3,4,5,6,7,8,9,10,11,12,13))
   assert re.fullmatch(r'[0-9a-f]{64}',row[13])

def verify_payload(root):
 p=json.loads((root/'working/task6/task6_corrected_payload.json').read_text()); tests=reconstruct((root/'evidence/flutter_test_1.3.11+28.log').read_text())
 fields=('ordinal','case_id','source_path','name','category'); assert [tuple(x[k] for k in fields) for x in p['tests']]==[tuple(x[k] for k in fields) for x in tests]
 assert Counter(x['category'] for x in tests)=={'Algorithm / reference':55,'Functional / regression':33,'Data / persistence':26,'UI / regression':14,'Invalid / boundary':7}
 src=json.loads((root/'evidence_map.json').read_text())['records']; assert len(src)==len(p['requirements'])==155 and Counter(x['status'] for x in src)==STATUS and p['requirement_status_counts']==STATUS
 preserved=('requirement_id','status','authoritative_sources','human_action','final_version_claim','rationale')
 for a,b in zip(src,p['requirements'],strict=True): assert all(a.get(k)==b.get(k) for k in preserved) and b['task2_evidence']['commit']==COMMIT and b['task6_evidence']
 required={'case_id','requirement_id','precondition_input','expected','actual','status','baseline','tester_category','timestamp','method','raw_path','sha256'}
 for r in p['result_rows']:
  if r['status']!='PASS': continue
  assert required<=r.keys() and all(r[k] not in (None,'') for k in required-{'requirement_id'})
  raw=root/r['raw_path']; assert raw.is_file() and sha(raw)==r['sha256'] and re.fullmatch(r'[0-9a-f]{64}',r['sha256'])
  assert all(r['evidence_tuple'][k]==r[k] for k in ('case_id','baseline','timestamp','method','raw_path','sha256'))
 assert next(x for x in p['result_rows'] if x['case_id']=='FIREBASE-PLAN-001')['status'].startswith('Not Executed')
 h3,h7=p['historical_uat'][2],p['historical_uat'][6]; assert (h3['date'],h3['round_revision'],h3['version'])==('2026-06-06','r2','1.1.2+10'); assert h7['version']=='not stated'
 return p

def actual_git_history():
 repo=Path(os.environ.get('FSOOKTA_AUTHORITATIVE_GIT_REPO',str(Path(__file__).resolve().parents[2])))
 assert repo.is_dir()
 def git(*args,binary=False):
  result=subprocess.run(['git','-C',str(repo),*args],check=True,capture_output=True,text=not binary)
  return result.stdout
 assert git('cat-file','-t',COMMIT).strip()=='commit'
 raw=git('log','--format=ENTRY%x00%H%x00%aI%x00%s%x00','--name-only','-z',COMMIT,binary=True)
 refs={}
 for line in git('show-ref').splitlines():
  oid,ref=line.split(' ',1)
  if ref.startswith('refs/heads/'): short=ref.removeprefix('refs/heads/')
  elif ref.startswith('refs/remotes/'): short=ref.removeprefix('refs/remotes/')
  else: continue
  refs.setdefault(oid,set()).add(short)
 history=[]
 for block in raw.split(b'ENTRY\x00')[1:]:
  fields=block.split(b'\x00'); commit=fields[0].decode(); authored=fields[1].decode(); subject=fields[2].decode(); paths=[value.decode().strip() for value in fields[4:] if value.decode().strip()]
  history.append({'commit':commit,'date':authored,'subject':subject,'refs':refs.get(commit,set()),'changed_paths':'; '.join(paths) if paths else 'Merge/no path list'})
 assert len(history)==137 and history[0]['commit']==COMMIT
 return history
def verify_git_history(root,p):
 history=p['git_history']; actual=actual_git_history(); assert len(history)==137 and len({x['commit'] for x in history})==137 and all(re.fullmatch(r'[0-9a-f]{40}',x['commit']) for x in history)
 assert len(actual)==len(history)
 for staged,authoritative in zip(history,actual,strict=True):
  assert staged['commit']==authoritative['commit'] and staged['date']==authoritative['date'] and staged['subject']==authoritative['subject'] and staged['changed_paths']==authoritative['changed_paths']
  assert {x.strip() for x in staged['refs'].split(',') if x.strip()}==authoritative['refs']
 pos={x['commit']:i for i,x in enumerate(history)}; sr=source_root(root)
 for d in p['defects']:
  assert d['before_commit'] in pos and d['after_commit'] in pos and pos[d['before_commit']]>pos[d['after_commit']]
  assert d['before_git_object']==history[pos[d['before_commit']]] and d['after_git_object']==history[pos[d['after_commit']]]
  assert d['source_path'] in d['after_git_object']['changed_paths'].split('; ')
  assert (sr/d['evidence_path']).is_file() and sha(sr/d['evidence_path'])==d['sha256']

def verify_workbooks(root,p):
 A=root/'artifacts'; audit=load_workbook(A/'06_Development_Audit_Trail.xlsx',data_only=False); master=load_workbook(A/'07_Master_Test_and_Verification_Package.xlsx',data_only=False); uat=load_workbook(A/'08_UAT_Field_Test_and_Usability_Package.xlsx',data_only=False)
 assert (len(audit.sheetnames),len(master.sheetnames),len(uat.sheetnames))==(11,16,14)
 for filename,workbook in [('06_Development_Audit_Trail.xlsx',audit),('07_Master_Test_and_Verification_Package.xlsx',master),('08_UAT_Field_Test_and_Usability_Package.xlsx',uat)]:
  expected=EXPECTED_MODELS['workbooks'][filename]; assert workbook.sheetnames==expected['sheet_order']
  actual_models={sheet.title:model_sha(sheet_model(sheet)) for sheet in workbook.worksheets}; assert actual_models==expected['sheet_model_sha256'], f'complete sheet model mismatch: {filename}'
 auto=[x for x in p['result_rows'] if x['case_id'].startswith('AUTO-')]; alg=[x for x in auto if x['category']=='Algorithm / reference']; boundary=[x for x in auto if x['category']=='Invalid / boundary']; functional=[x for x in auto if x['category']=='Functional / regression']
 assert_join(master['Automated Suite'],auto); assert_join(master['Integration Regression'],auto); assert_join(master['Algorithm Reference'],alg); assert_join(master['Threshold Boundaries'],boundary); assert_join(master['Invalid Missing Inputs'],boundary); assert_join(master['Functional Cases'],functional); assert_join(master['Network Error Applicability'],[next(x for x in p['result_rows'] if x['case_id']=='FIREBASE-PLAN-001')]); assert_join(master['Test Plan'],[x for x in p['result_rows'] if not x['case_id'].startswith('AUTO-')]); assert_join(master['Evidence Index'],[x for x in p['result_rows'] if x['status']=='PASS'])
 req=table_rows(audit['Requirement Traceability'],8); src=json.loads((root/'evidence_map.json').read_text())['records']; assert len(req)==155 and Counter(x[1] for x in req)==STATUS
 for a,b in zip(src,req,strict=True): assert b[0]==a['requirement_id'] and b[1]==a['status'] and json.loads(b[2])==a['authoritative_sources'] and b[3]==a['human_action'] and b[4]==a['rationale']
 rt=table_rows(master['Requirement Test Trace'],6); assert [(x[0],x[1],x[2],x[3]) for x in rt]==[(x['requirement_id'],x['status'],x['human_action'],x['rationale']) for x in p['requirements']]
 hist=table_rows(uat['Historical Evidence'],13); keys=('evidence_id','date','round_revision','version','device','runner_harness','build_workaround','assessment_bypass','blocker','observation','evidence_path','sha256','evidence_layer'); assert hist==[tuple(None if x[k]=='' else x[k] for k in keys) for x in p['historical_uat']]
 s=uat['SUS Response Form']; wording=[s.cell(r,2).value for r in range(5,15)]; assert wording==[x['wording'] for x in p['sus']['items']] and s['D16'].value=='=IF(COUNT(D5:D14)=10,(SUM(D5,D7,D9,D11,D13)-5+25-SUM(D6,D8,D10,D12,D14))*2.5,"")'; assert any(v.type=='whole' and v.formula1=='1' and v.formula2=='5' and 'D5:D14' in str(v.sqref) for v in s.data_validations.dataValidation)
 defects=table_rows(master['Bugs Corrections Retest'],13); dk=('defect_id','version','description','before_commit','after_commit','source_path');
 for row,d in zip(defects,p['defects'],strict=True): assert row[:6]==tuple(d[k] for k in dk) and row[6]==d['before_git_object']['changed_paths'] and row[7]==d['after_git_object']['changed_paths'] and row[10:13]==(d['evidence_path'],d['sha256'],d['status'])
 before=table_rows(audit['Before After Evidence'],10); bk=('change_id','before_commit','after_commit','source_path','before','after','before_changed_paths','after_changed_paths','evidence_path','sha256'); assert before==[tuple(x[k] for k in bk) for x in p['before_after']]
 version_rows=table_rows(audit['Version History'],4); assert len(version_rows)==137
 for row,actual in zip(version_rows,actual_git_history(),strict=True):
  excel_date=(datetime.fromisoformat(actual['date']).astimezone(timezone.utc).replace(tzinfo=None)-datetime(1899,12,30)).total_seconds()/86400
  assert row==(actual['commit'],excel_date,None,actual['subject'])
 return {'workbooks':3,'sheets':41}

def doc_text(path):
 d=Document(path); return '\n'.join(x.text for x in d.paragraphs)+'\n'+'\n'.join(c.text for t in d.tables for row in t.rows for c in row.cells)
def verify_reports_pdfs(root,p):
 A=root/'artifacts'; a=doc_text(A/'07_Final_Algorithm_Verification_Report.docx'); u=doc_text(A/'08_UAT_and_Field_Test_Technical_Report.docx')
 assert 'exactly 135 completed Flutter tests' in a and 'Algorithm / reference\n55' in a and 'Invalid / boundary\n7' in a and 'AUTO-135' in a and p['tests'][-1]['name'] in a and 'The summary line is excluded' in a
 assert 'H-UAT-003 | 2026-06-06 | 1.1.2+10' in u and 'Round / revision\nr2' in u and 'H-UAT-007 | 2026-07-19 | not stated' in u and 'BLOCKED BY DEVICE CONNECTION' in u
 expected=json.loads((root/'manifests/task6_pdf_page_counts.json').read_text()); actual={name:len(PdfReader(A/name).pages) for name in expected}; assert actual==expected and all(v>0 for v in actual.values())
 return {'documents':2,'pdfs':len(actual),'pdf_pages':sum(actual.values())}
def verify_visual(root):
 M=root/'manifests'; expected_path=M/'task6_expected_visual_renders.json'; expected=json.loads(expected_path.read_text()); record=json.loads((M/'task6_visual_qa_manifest.json').read_text()); decision_path=Path(record['external_decision_path']); decision=json.loads(decision_path.read_text())
 assert record['source'].startswith('external manual decision') and record['decision']=='PASS' and record['expected_manifest_sha256']==sha(expected_path) and record['external_decision_sha256']==sha(decision_path) and decision['expected_manifest_sha256']==sha(expected_path)
 assert expected['count']==record['expected_render_count']==len(expected['renders'])
 for x in expected['renders']: assert (root/x['path']).is_file() and sha(root/x['path'])==x['sha256']
 return {'renders':expected['count']}
def verify(root):
 verify_committed_artifact_expectations(root); p=verify_payload(root); verify_git_history(root,p); out={'status':'PASS','tests':135,'algorithm':55,'boundary':7,'requirements':155,'pass_results':136}; out.update(verify_workbooks(root,p)); out.update(verify_reports_pdfs(root,p)); out.update(verify_visual(root)); return out
if __name__=='__main__': print(json.dumps(verify(Path(sys.argv[1] if len(sys.argv)>1 else '/private/tmp/fsookta-final-handover')),indent=2,sort_keys=True))
