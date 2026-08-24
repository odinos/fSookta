#!/usr/bin/env python3
"""Independent, fail-closed verifier for Task 6 fix round 1.

This module deliberately does not import the builder.  It reconstructs the
expanded-reporter state, source requirement records, workbook surfaces, and
render bindings from authoritative inputs.
"""
from __future__ import annotations

import hashlib
import json
import re
import sys
from collections import Counter
from pathlib import Path

from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader

BASELINE = "1.3.11+28"
COMMIT = "bf8867a2083357cb9d60915bf6c2233801f923d8"
REPORTER = re.compile(r"^(?P<clock>\d{2}:\d{2}) \+(?P<count>\d+)(?: -\d+)?: (?P<label>.+)$")
PASS_FIELDS = {"case_id", "baseline", "timestamp", "method", "raw_path", "sha256"}
EXPECTED_STATUS = {"Pending Owner Action": 137, "Exception Approval Required": 12, "Pending Researcher Evidence": 6}
SUS_WORDING = [
    "I think that I would like to use this system frequently.",
    "I found the system unnecessarily complex.",
    "I thought the system was easy to use.",
    "I think that I would need the support of a technical person to use this system.",
    "I found the various functions in this system were well integrated.",
    "I thought there was too much inconsistency in this system.",
    "I would imagine that most people would learn to use this system very quickly.",
    "I found the system very cumbersome to use.",
    "I felt very confident using the system.",
    "I needed to learn a lot of things before I could get going with this system.",
]


def sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def classify(path: str, name: str) -> str:
    value = f"{path} {name}".lower()
    groups = (
        ("Algorithm / reference", ("algorithm", "reba", "iso11228", "pose", "xgboost", "logistic")),
        ("Invalid / boundary", ("invalid", "required", "malformed", "missing", "zero-height", "error")),
        ("UI / regression", ("capture", "screen", "layout", "responsive", "widget", "navigation")),
        ("Data / persistence", ("persist", "draft", "export", "history", "storage")),
    )
    return next((label for label, needles in groups if any(n in value for n in needles)), "Functional / regression")


def reconstruct_tests(log_text: str) -> list[dict]:
    """Attribute each counter increment to the preceding displayed test."""
    completed: list[dict] = []
    pending: tuple[str, str] | None = None
    observed = 0
    for raw in log_text.splitlines():
        match = REPORTER.fullmatch(raw)
        if not match:
            continue
        counter = int(match.group("count"))
        label = match.group("label").strip()
        if counter != observed:
            if counter != observed + 1 or pending is None:
                raise AssertionError(f"invalid expanded-reporter transition {observed}->{counter}")
            path, name = pending
            completed.append({"ordinal": counter, "case_id": f"AUTO-{counter:03d}", "source_path": path, "name": name, "category": classify(path, name)})
            observed = counter
        if label == "All tests passed!" or label.startswith("loading "):
            continue
        if ".dart: " not in label:
            raise AssertionError(f"unrecognized reporter label: {label}")
        prefix, name = label.rsplit(".dart: ", 1)
        source_relative = prefix.rsplit("/source/", 1)[-1] + ".dart"
        pending = (source_relative, name)
    assert observed == 135 and len(completed) == 135
    assert completed[-1]["name"] == "clears the draft after a successful evaluation save"
    assert all(x["name"] != "All tests passed!" for x in completed)
    return completed


def verify_payload(root: Path) -> dict:
    payload = json.loads((root / "working/task6/task6_corrected_payload.json").read_text())
    tests = reconstruct_tests((root / "evidence/flutter_test_1.3.11+28.log").read_text())
    test_fields = ("ordinal", "case_id", "source_path", "name", "category")
    assert [tuple(x[k] for k in test_fields) for x in payload["tests"]] == [tuple(x[k] for k in test_fields) for x in tests]

    source = json.loads((root / "evidence_map.json").read_text())["records"]
    assert len(source) == len(payload["requirements"]) == 155
    assert Counter(x["status"] for x in source) == EXPECTED_STATUS
    assert payload["requirement_status_counts"] == EXPECTED_STATUS
    preserved = ("requirement_id", "status", "authoritative_sources", "human_action", "final_version_claim", "rationale")
    for original, layered in zip(source, payload["requirements"], strict=True):
        assert all(layered.get(key) == original.get(key) for key in preserved)
        assert layered["task2_evidence"]["commit"] == COMMIT
        assert layered["task6_evidence"]["claim"] == "Traceability layer; source status not overwritten"

    for row in payload["result_rows"]:
        if row["status"] != "PASS":
            continue
        evidence = row.get("evidence_tuple", {})
        assert PASS_FIELDS <= evidence.keys()
        assert evidence["case_id"] == row["case_id"] and evidence["baseline"] == BASELINE
        assert evidence["timestamp"] and evidence["method"]
        raw_path = root / evidence["raw_path"]
        assert raw_path.is_file() and re.fullmatch(r"[0-9a-f]{64}", evidence["sha256"])
        assert sha256(raw_path) == evidence["sha256"]
    firebase = next(x for x in payload["result_rows"] if x["case_id"] == "FIREBASE-PLAN-001")
    assert firebase["status"].startswith("Not Executed") and not firebase["evidence_tuple"]
    return payload


def rows(sheet, start: int, end: int, width: int) -> list[tuple]:
    return [tuple(cell.value for cell in row[:width]) for row in sheet.iter_rows(min_row=start, max_row=end)]


def verify_workbooks(root: Path, payload: dict) -> dict:
    artifacts = root / "artifacts"
    audit = load_workbook(artifacts / "06_Development_Audit_Trail.xlsx", data_only=False)
    master = load_workbook(artifacts / "07_Master_Test_and_Verification_Package.xlsx", data_only=False)
    uat = load_workbook(artifacts / "08_UAT_Field_Test_and_Usability_Package.xlsx", data_only=False)

    requirement_rows = rows(audit["Requirement Traceability"], 5, 159, 8)
    assert len(requirement_rows) == 155 and Counter(x[1] for x in requirement_rows) == EXPECTED_STATUS
    source = json.loads((root / "evidence_map.json").read_text())["records"]
    for expected, actual in zip(source, requirement_rows, strict=True):
        assert actual[0] == expected["requirement_id"] and actual[1] == expected["status"]
        assert json.loads(actual[2]) == expected["authoritative_sources"]
        assert actual[3] == expected["human_action"] and actual[4] == expected["rationale"]
        assert actual[5] == str(expected["final_version_claim"]).lower()
        assert json.loads(actual[6])["commit"] == COMMIT and json.loads(actual[7])["claim"].endswith("not overwritten")

    auto = rows(master["Automated Suite"], 5, 139, 12)
    fields = ("case_id", "name", "category")
    assert [x[:3] for x in auto] == [tuple(x[k] for k in fields) for x in payload["tests"]]
    assert auto[-1][1] == "clears the draft after a successful evaluation save"
    assert all(x[3] == "PASS" and x[4] == BASELINE and x[5] and x[6] and x[7] and re.fullmatch(r"[0-9a-f]{64}", x[8]) for x in auto)
    plans = rows(master["Test Plan"], 5, 7, 12)
    assert plans[0][0] == "STATIC-001" and plans[0][3] == "PASS" and all(plans[0][i] for i in range(4, 9))
    assert all(x[3].startswith("Not Executed") for x in plans[1:])

    hist_fields = ("evidence_id", "date", "version", "device", "runner_harness", "build_workaround", "assessment_bypass", "blocker", "observation", "evidence_path", "sha256", "evidence_layer")
    historical = rows(uat["Historical Evidence"], 5, 11, 12)
    assert historical == [tuple(x[k] for k in hist_fields) for x in payload["historical_uat"]]
    assert historical[5][2:9] == ("1.3.6+21", "Physical iPhone SE iOS 26.5", "No physical Android detected; iPad simulator only", "Uploaded-video UAT", "Temporary assessment bypass", "Offline relaunch pending; simulator runner idle", "User-observed iPhone checks; historical only")
    assert historical[6][4] == "iOS simulator + Android emulator" and historical[6][6] == "None - production assessment path" and historical[6][7] == "BLOCKED BY DEVICE CONNECTION"
    assert "assessment bypass" not in historical[2][5].lower()

    sus = uat["SUS Response Form"]
    assert [sus.cell(r, 2).value for r in range(5, 15)] == SUS_WORDING
    assert all(sus.cell(r, 3).value == "1 Strongly disagree / 5 Strongly agree" for r in range(5, 15))
    assert sus["D16"].value == '=IF(COUNT(D5:D14)=10,(SUM(D5,D7,D9,D11,D13)-5+25-SUM(D6,D8,D10,D12,D14))*2.5,"")'
    validations = list(sus.data_validations.dataValidation)
    assert any(v.type == "whole" and v.formula1 == "1" and v.formula2 == "5" and "D5:D14" in str(v.sqref) for v in validations)
    assert "includes row 5" in str(uat["Control Summary"]["C8"].value).lower()

    defect_fields = ("defect_id", "version", "description", "before_commit", "after_commit", "source_path", "fix", "retest", "evidence_path", "sha256", "status")
    defects = rows(master["Bugs Corrections Retest"], 5, 7, 11)
    assert defects == [tuple(x[k] for k in defect_fields) for x in payload["defects"]]
    before_fields = ("change_id", "before_commit", "after_commit", "source_path", "before", "after", "evidence_path", "sha256")
    before = rows(audit["Before After Evidence"], 5, 7, 8)
    assert before == [tuple(x[k] for k in before_fields) for x in payload["before_after"]]
    assert all(re.fullmatch(r"[0-9a-f]{40}", x[1]) and re.fullmatch(r"[0-9a-f]{40}", x[2]) and re.fullmatch(r"[0-9a-f]{64}", x[7]) for x in before)
    return {"workbooks": 3, "sheets": len(audit.sheetnames) + len(master.sheetnames) + len(uat.sheetnames)}


def verify_reports_and_pdfs(root: Path) -> dict:
    artifacts = root / "artifacts"
    docs = [artifacts / "07_Final_Algorithm_Verification_Report.docx", artifacts / "08_UAT_and_Field_Test_Technical_Report.docx"]
    texts = ["\n".join(p.text for p in Document(path).paragraphs) + "\n" + "\n".join(c.text for t in Document(path).tables for row in t.rows for c in row.cells) for path in docs]
    assert "clears the draft after a successful evaluation save" in texts[0]
    assert "+135 All tests passed proves AUTO-135" in texts[0]
    assert "No final participant UAT" in texts[1] and "Temporary assessment bypass" in texts[1] and "BLOCKED BY DEVICE CONNECTION" in texts[1]
    pdfs = {
        "06_Development_Audit_Trail.pdf": 24,
        "07_Master_Test_and_Verification_Package.pdf": 36,
        "08_UAT_Field_Test_and_Usability_Package.pdf": 15,
        "07_Final_Algorithm_Verification_Report.pdf": 12,
        "08_UAT_and_Field_Test_Technical_Report.pdf": 6,
    }
    actual = {name: len(PdfReader(artifacts / name).pages) for name in pdfs}
    assert actual == pdfs
    return {"documents": 2, "pdfs": 5, "pdf_pages": sum(actual.values())}


def verify_visuals(root: Path) -> dict:
    manifests = root / "manifests"
    expected_path = manifests / "task6_fix_round1_expected_visual_renders.json"
    contacts_path = manifests / "task6_fix_round1_contact_sheets.json"
    decision = json.loads((manifests / "task6_fix_round1_visual_decision.json").read_text())
    expected = json.loads(expected_path.read_text())
    contacts = json.loads(contacts_path.read_text())
    assert expected["count"] == len(expected["renders"]) == 134
    assert contacts["count"] == len(contacts["sheets"]) == 12
    assert decision["decision"] == "PASS" and decision["inspector"] and decision["inspected_at"]
    assert decision["expected_manifest_sha256"] == sha256(expected_path)
    assert decision["contact_sheet_manifest_sha256"] == sha256(contacts_path)
    assert decision["expected_render_count"] == 134 and decision["contact_sheet_count"] == 12
    assert "must not create or change" in decision["automation_boundary"]
    for entry in expected["renders"] + contacts["sheets"]:
        path = root / entry["path"]
        assert path.is_file() and sha256(path) == entry["sha256"]
    return {"renders": 134, "contact_sheets": 12}


def verify(root: Path) -> dict:
    payload = verify_payload(root)
    result = {"status": "PASS", "tests": len(payload["tests"]), "requirements": len(payload["requirements"]), "pass_results": sum(x["status"] == "PASS" for x in payload["result_rows"])}
    result.update(verify_workbooks(root, payload))
    result.update(verify_reports_and_pdfs(root))
    result.update(verify_visuals(root))
    return result


if __name__ == "__main__":
    base = Path(sys.argv[1] if len(sys.argv) > 1 else "/private/tmp/fsookta-final-handover")
    print(json.dumps(verify(base), indent=2, sort_keys=True))
