"""Deterministic Markdown-to-DOCX tooling for the Sookta 2.1.0 documents."""

from __future__ import annotations

import re
from dataclasses import dataclass
from pathlib import Path
from typing import Mapping

from docx import Document
from docx.document import Document as DocumentObject
from docx.enum.section import WD_ORIENT
from docx.enum.style import WD_STYLE_TYPE
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Mm, Pt, Twips
from PIL import Image, ImageDraw, ImageFont


A4_CONTENT_WIDTH_DXA = 9026
_A4_CONTENT_WIDTH_INCHES = A4_CONTENT_WIDTH_DXA / 1440
_DIAGRAM_MARKER = re.compile(r"^<!--\s*DOCX_DIAGRAM:([a-z0-9-]+)\s*-->$")
_HEADING = re.compile(r"^(#{1,4})\s+(.+)$")
_UNORDERED_LIST = re.compile(r"^\s*[-*+]\s+(.+)$")
_ORDERED_LIST = re.compile(r"^\s*(\d+)[.)]\s+(.+)$")
_BLOCKQUOTE = re.compile(r"^>\s?(.*)$")
_IMAGE = re.compile(r"^!\[([^\]]*)]\(([^)]+)\)$")
_INLINE_TOKEN = re.compile(
    r"(\*\*.+?\*\*|__.+?__|`.+?`|\*[^*]+?\*|_[^_]+?_|"
    r"\[[^\]]+]\([^)]+\))"
)


@dataclass(frozen=True)
class DocumentProfile:
    name: str
    body_size_pt: float
    title_size_pt: float
    heading_sizes_pt: tuple[float, float, float]
    line_spacing: float
    header_text: str
    footer_status: str


@dataclass(frozen=True)
class DiagramSpec:
    title: str
    nodes: tuple[tuple[str, str], ...]
    edges: tuple[tuple[str, str, str], ...]


BALANCED_PROFILE = DocumentProfile(
    name="balanced",
    body_size_pt=11.0,
    title_size_pt=25.0,
    heading_sizes_pt=(18.0, 15.0, 12.5),
    line_spacing=1.15,
    header_text="Sookta Application 2.1.0 | เอกสารฉบับสมดุล",
    footer_status="Controlled copy — Release 2.1.0",
)

TECHNICAL_PROFILE = DocumentProfile(
    name="technical",
    body_size_pt=10.5,
    title_size_pt=24.0,
    heading_sizes_pt=(17.5, 14.5, 12.0),
    line_spacing=1.1,
    header_text="Sookta Application 2.1.0 | Technical Specification",
    footer_status="Technical controlled copy — Release 2.1.0",
)


def _set_run_font(run, font_name: str = "Sarabun") -> None:
    run.font.name = font_name
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:ascii"), font_name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:hAnsi"), font_name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:eastAsia"), font_name)
    run._element.get_or_add_rPr().get_or_add_rFonts().set(qn("w:cs"), font_name)


def _set_style_font(style, font_name: str = "Sarabun") -> None:
    style.font.name = font_name
    fonts = style.element.get_or_add_rPr().get_or_add_rFonts()
    for attribute in ("ascii", "hAnsi", "eastAsia", "cs"):
        fonts.set(qn(f"w:{attribute}"), font_name)


def _configure_styles(document: DocumentObject, profile: DocumentProfile) -> None:
    styles = document.styles
    for style in styles:
        if style.type in (WD_STYLE_TYPE.PARAGRAPH, WD_STYLE_TYPE.CHARACTER):
            _set_style_font(style)

    normal = styles["Normal"]
    normal.font.size = Pt(profile.body_size_pt)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = profile.line_spacing

    title = styles["Title"]
    title.font.size = Pt(profile.title_size_pt)
    title.font.bold = True
    title.font.color.rgb = None
    title.paragraph_format.space_after = Pt(18)
    title.paragraph_format.alignment = WD_ALIGN_PARAGRAPH.LEFT

    subtitle = styles["Subtitle"]
    subtitle.font.size = Pt(profile.body_size_pt + 1)
    subtitle.font.italic = False

    for index, size in enumerate(profile.heading_sizes_pt, start=1):
        heading = styles[f"Heading {index}"]
        heading.font.size = Pt(size)
        heading.font.bold = True
        heading.paragraph_format.keep_with_next = True
        heading.paragraph_format.space_before = Pt(12 if index == 1 else 9)
        heading.paragraph_format.space_after = Pt(5)

    styles["Caption"].font.size = Pt(9)
    styles["Caption"].font.italic = True

    if "Code Block" not in styles:
        code_style = styles.add_style("Code Block", WD_STYLE_TYPE.PARAGRAPH)
    else:
        code_style = styles["Code Block"]
    _set_style_font(code_style)
    code_style.font.size = Pt(9)
    code_style.paragraph_format.left_indent = Cm(0.5)
    code_style.paragraph_format.space_after = Pt(1)


def _set_cell_shading(cell, fill: str) -> None:
    properties = cell._tc.get_or_add_tcPr()
    shading = properties.find(qn("w:shd"))
    if shading is None:
        shading = OxmlElement("w:shd")
        properties.append(shading)
    shading.set(qn("w:fill"), fill)


def _set_cell_width(cell, width_dxa: int) -> None:
    properties = cell._tc.get_or_add_tcPr()
    width = properties.find(qn("w:tcW"))
    if width is None:
        width = OxmlElement("w:tcW")
        properties.append(width)
    width.set(qn("w:type"), "dxa")
    width.set(qn("w:w"), str(width_dxa))


def _set_repeat_table_header(row) -> None:
    properties = row._tr.get_or_add_trPr()
    repeat = OxmlElement("w:tblHeader")
    repeat.set(qn("w:val"), "true")
    properties.append(repeat)


def _set_table_cant_split(row) -> None:
    properties = row._tr.get_or_add_trPr()
    cant_split = OxmlElement("w:cantSplit")
    properties.append(cant_split)


def _set_paragraph_numbering(paragraph, num_id: int, level: int = 0) -> None:
    properties = paragraph._p.get_or_add_pPr()
    number_properties = properties.get_or_add_numPr()
    level_element = number_properties.get_or_add_ilvl()
    level_element.val = level
    number_id = number_properties.get_or_add_numId()
    number_id.val = num_id


def _new_numbering_instance(
    document: DocumentObject,
    *,
    template_num_id: int,
    start: int,
) -> int:
    numbering = document.part.numbering_part.element
    template = next(
        num
        for num in numbering.findall(qn("w:num"))
        if int(num.get(qn("w:numId"))) == template_num_id
    )
    abstract_id = template.find(qn("w:abstractNumId")).get(qn("w:val"))
    new_num_id = max(
        int(num.get(qn("w:numId")))
        for num in numbering.findall(qn("w:num"))
    ) + 1

    num = OxmlElement("w:num")
    num.set(qn("w:numId"), str(new_num_id))
    abstract = OxmlElement("w:abstractNumId")
    abstract.set(qn("w:val"), abstract_id)
    num.append(abstract)
    override = OxmlElement("w:lvlOverride")
    override.set(qn("w:ilvl"), "0")
    start_override = OxmlElement("w:startOverride")
    start_override.set(qn("w:val"), str(start))
    override.append(start_override)
    num.append(override)
    numbering.append(num)
    return new_num_id


def _append_page_field(paragraph) -> None:
    run = paragraph.add_run()
    begin = OxmlElement("w:fldChar")
    begin.set(qn("w:fldCharType"), "begin")
    instruction = OxmlElement("w:instrText")
    instruction.set(qn("xml:space"), "preserve")
    instruction.text = " PAGE "
    separate = OxmlElement("w:fldChar")
    separate.set(qn("w:fldCharType"), "separate")
    text = OxmlElement("w:t")
    text.text = "1"
    end = OxmlElement("w:fldChar")
    end.set(qn("w:fldCharType"), "end")
    run._r.extend((begin, instruction, separate, text, end))
    _set_run_font(run)


def _configure_page(document: DocumentObject, profile: DocumentProfile) -> None:
    section = document.sections[0]
    section.orientation = WD_ORIENT.PORTRAIT
    section.page_width = Mm(210)
    section.page_height = Mm(297)
    section.top_margin = Cm(2.54)
    section.bottom_margin = Cm(2.54)
    section.left_margin = Cm(2.54)
    section.right_margin = Cm(2.54)
    section.header_distance = Cm(1.25)
    section.footer_distance = Cm(1.25)

    header = section.header
    header_paragraph = header.paragraphs[0]
    header_paragraph.text = profile.header_text
    header_paragraph.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    for run in header_paragraph.runs:
        _set_run_font(run)
        run.font.size = Pt(8.5)

    footer = section.footer
    footer_paragraph = footer.paragraphs[0]
    footer_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = footer_paragraph.add_run(f"{profile.footer_status}  |  หน้า ")
    _set_run_font(run)
    run.font.size = Pt(8.5)
    _append_page_field(footer_paragraph)


def _add_inline_markdown(paragraph, text: str) -> None:
    cursor = 0
    for match in _INLINE_TOKEN.finditer(text):
        if match.start() > cursor:
            _set_run_font(paragraph.add_run(text[cursor : match.start()]))
        token = match.group(0)
        if token.startswith(("**", "__")):
            run = paragraph.add_run(token[2:-2])
            run.bold = True
        elif token.startswith("`"):
            run = paragraph.add_run(token[1:-1])
            run.font.size = Pt(9)
        elif token.startswith(("*", "_")):
            run = paragraph.add_run(token[1:-1])
            run.italic = True
        elif token.startswith("["):
            label, url = re.match(r"^\[([^\]]+)]\(([^)]+)\)$", token).groups()
            run = paragraph.add_run(f"{label} ({url})")
            run.font.underline = True
        else:
            run = paragraph.add_run(token)
        _set_run_font(run)
        cursor = match.end()
    if cursor < len(text):
        _set_run_font(paragraph.add_run(text[cursor:]))


def _is_table_separator(line: str) -> bool:
    cells = [cell.strip() for cell in line.strip().strip("|").split("|")]
    return bool(cells) and all(re.fullmatch(r":?-{3,}:?", cell) for cell in cells)


def _parse_table(lines: list[str], start_index: int) -> tuple[list[list[str]], int]:
    rows: list[list[str]] = []
    index = start_index
    while index < len(lines):
        line = lines[index].strip()
        if not (line.startswith("|") and line.endswith("|")):
            break
        if not _is_table_separator(line):
            rows.append([cell.strip() for cell in line.strip("|").split("|")])
        index += 1
    return rows, index


def _table_widths(column_count: int) -> list[int]:
    """Return fixed A4-width columns, with extra room for traceability evidence."""
    if column_count == 7:
        ratios = (0.13, 0.14, 0.16, 0.13, 0.15, 0.13, 0.16)
        widths = [round(A4_CONTENT_WIDTH_DXA * ratio) for ratio in ratios]
    elif column_count == 6:
        ratios = (0.13, 0.17, 0.15, 0.23, 0.15, 0.17)
        widths = [round(A4_CONTENT_WIDTH_DXA * ratio) for ratio in ratios]
    elif column_count == 5:
        ratios = (0.16, 0.19, 0.20, 0.25, 0.20)
        widths = [round(A4_CONTENT_WIDTH_DXA * ratio) for ratio in ratios]
    else:
        widths = [A4_CONTENT_WIDTH_DXA // column_count] * column_count
    widths[-1] += A4_CONTENT_WIDTH_DXA - sum(widths)
    return widths


def _add_table(document: DocumentObject, rows: list[list[str]]) -> None:
    if not rows:
        return
    column_count = max(len(row) for row in rows)
    widths = _table_widths(column_count)
    compact_font_size = (
        7.5
        if column_count >= 7
        else 8.25
        if column_count >= 6
        else 9.0
        if column_count >= 5
        else None
    )
    table = document.add_table(rows=len(rows), cols=column_count)
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    table.autofit = False

    for index, grid_column in enumerate(table._tbl.tblGrid.gridCol_lst):
        grid_column.set(qn("w:w"), str(widths[index]))

    for row_index, row_values in enumerate(rows):
        row = table.rows[row_index]
        _set_table_cant_split(row)
        if row_index == 0:
            _set_repeat_table_header(row)
        for column_index, cell in enumerate(row.cells):
            _set_cell_width(cell, widths[column_index])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            if row_index == 0:
                _set_cell_shading(cell, "DDEBF7")
            paragraph = cell.paragraphs[0]
            paragraph.paragraph_format.space_after = Pt(1 if compact_font_size else 2)
            if compact_font_size:
                paragraph.paragraph_format.line_spacing = 1.0
            value = row_values[column_index] if column_index < len(row_values) else ""
            _add_inline_markdown(paragraph, value)
            if compact_font_size:
                for run in paragraph.runs:
                    run.font.size = Pt(compact_font_size)
            if row_index == 0:
                for run in paragraph.runs:
                    run.bold = True


def _load_diagram_font(font_path: Path, size: int):
    try:
        return ImageFont.truetype(str(font_path), size)
    except (OSError, ValueError):
        return ImageFont.load_default()


def _wrap_text(draw: ImageDraw.ImageDraw, text: str, font, max_width: int) -> list[str]:
    words = text.split()
    if not words:
        return [""]
    lines: list[str] = []
    current = words[0]
    for word in words[1:]:
        candidate = f"{current} {word}"
        if draw.textbbox((0, 0), candidate, font=font)[2] <= max_width:
            current = candidate
        else:
            lines.append(current)
            current = word
    lines.append(current)
    return lines


def _draw_centered_multiline(
    draw: ImageDraw.ImageDraw,
    bounds: tuple[int, int, int, int],
    text: str,
    font,
    fill: str,
) -> None:
    left, top, right, bottom = bounds
    lines = _wrap_text(draw, text, font, right - left - 36)
    line_height = max(28, draw.textbbox((0, 0), "กA", font=font)[3] + 8)
    total_height = len(lines) * line_height
    y = top + ((bottom - top) - total_height) / 2
    for line in lines:
        bbox = draw.textbbox((0, 0), line, font=font)
        x = left + ((right - left) - (bbox[2] - bbox[0])) / 2
        draw.text((x, y), line, font=font, fill=fill)
        y += line_height


def render_flow_diagram(
    spec: DiagramSpec,
    output_path: Path,
    font_path: Path,
) -> None:
    """Render a compact, deterministic flow diagram as a high-resolution PNG."""

    output_path.parent.mkdir(parents=True, exist_ok=True)
    width = 1800
    columns = min(3, max(1, len(spec.nodes)))
    rows = (len(spec.nodes) + columns - 1) // columns
    height = max(620, 230 + rows * 235)
    image = Image.new("RGB", (width, height), "#F7FAFC")
    draw = ImageDraw.Draw(image)
    title_font = _load_diagram_font(font_path, 46)
    node_font = _load_diagram_font(font_path, 31)
    label_font = _load_diagram_font(font_path, 24)

    draw.rounded_rectangle((25, 25, width - 25, height - 25), radius=26, outline="#CBD5E1", width=4)
    title_bbox = draw.textbbox((0, 0), spec.title, font=title_font)
    draw.text(
        ((width - (title_bbox[2] - title_bbox[0])) / 2, 60),
        spec.title,
        font=title_font,
        fill="#17324D",
    )

    node_width = 470
    node_height = 125
    x_gap = (width - 2 * 110 - columns * node_width) / max(1, columns - 1)
    positions: dict[str, tuple[int, int, int, int]] = {}
    for index, (node_id, label) in enumerate(spec.nodes):
        row = index // columns
        col = index % columns
        x = int(110 + col * (node_width + x_gap))
        y = int(155 + row * 235)
        bounds = (x, y, x + node_width, y + node_height)
        positions[node_id] = bounds
        fill = "#E8F3FA" if index not in (0, len(spec.nodes) - 1) else "#D9EAD3"
        draw.rounded_rectangle(bounds, radius=22, fill=fill, outline="#236A8D", width=4)
        _draw_centered_multiline(draw, bounds, label, node_font, "#102A43")

    for source, target, label in spec.edges:
        if source not in positions or target not in positions:
            continue
        source_box = positions[source]
        target_box = positions[target]
        source_center = ((source_box[0] + source_box[2]) // 2, (source_box[1] + source_box[3]) // 2)
        target_center = ((target_box[0] + target_box[2]) // 2, (target_box[1] + target_box[3]) // 2)
        if abs(target_center[1] - source_center[1]) < 50:
            start = (source_box[2] + 8, source_center[1])
            end = (target_box[0] - 8, target_center[1])
        else:
            start = (source_center[0], source_box[3] + 8)
            end = (target_center[0], target_box[1] - 8)
        draw.line((start, end), fill="#526D82", width=7)
        angle_x = 18 if end[0] >= start[0] else -18
        angle_y = 18 if end[1] >= start[1] else -18
        draw.polygon(
            (end, (end[0] - angle_x, end[1] - angle_y // 2), (end[0] - angle_x // 2, end[1] - angle_y)),
            fill="#526D82",
        )
        if label:
            midpoint = ((start[0] + end[0]) // 2, (start[1] + end[1]) // 2 - 30)
            label_bbox = draw.textbbox((0, 0), label, font=label_font)
            pad = 10
            draw.rounded_rectangle(
                (
                    midpoint[0] - pad,
                    midpoint[1] - pad,
                    midpoint[0] + label_bbox[2] + pad,
                    midpoint[1] + label_bbox[3] + pad,
                ),
                radius=9,
                fill="#FFFFFF",
            )
            draw.text(midpoint, label, font=label_font, fill="#334E68")

    image.save(output_path, "PNG", optimize=True)


def _add_image(document: DocumentObject, image_path: Path, alt_text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = paragraph.add_run()
    shape = run.add_picture(str(image_path), width=Mm(153))
    shape._inline.docPr.set("name", alt_text or image_path.stem)
    shape._inline.docPr.set("title", alt_text or image_path.stem)
    shape._inline.docPr.set("descr", alt_text or image_path.stem)


def _add_diagram(
    document: DocumentObject,
    key: str,
    spec: DiagramSpec,
    output_dir: Path,
    font_path: Path,
) -> None:
    diagram_path = output_dir / f"{key}.png"
    render_flow_diagram(spec, diagram_path, font_path)
    _add_image(document, diagram_path, spec.title)
    caption = document.add_paragraph(style="Caption")
    caption.alignment = WD_ALIGN_PARAGRAPH.CENTER
    _add_inline_markdown(caption, f"แผนภาพ: {spec.title}")


def _add_rule(document: DocumentObject) -> None:
    paragraph = document.add_paragraph()
    properties = paragraph._p.get_or_add_pPr()
    borders = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "6")
    bottom.set(qn("w:space"), "1")
    bottom.set(qn("w:color"), "AAB7C4")
    borders.append(bottom)
    properties.append(borders)


def _add_blockquote(document: DocumentObject, text: str) -> None:
    paragraph = document.add_paragraph()
    paragraph.paragraph_format.left_indent = Cm(0.6)
    paragraph.paragraph_format.right_indent = Cm(0.2)
    properties = paragraph._p.get_or_add_pPr()

    borders = OxmlElement("w:pBdr")
    left = OxmlElement("w:left")
    left.set(qn("w:val"), "single")
    left.set(qn("w:sz"), "16")
    left.set(qn("w:space"), "8")
    left.set(qn("w:color"), "4F81BD")
    borders.append(left)
    properties.append(borders)

    shading = OxmlElement("w:shd")
    shading.set(qn("w:fill"), "EEF5FB")
    properties.append(shading)
    _add_inline_markdown(paragraph, text)


def build_docx(
    markdown_path: Path,
    output_path: Path,
    profile: DocumentProfile,
    diagrams: Mapping[str, DiagramSpec],
    font_path: Path,
) -> None:
    """Build an A4 portrait Word document from the supported Markdown subset."""

    markdown_path = Path(markdown_path)
    output_path = Path(output_path)
    output_path.parent.mkdir(parents=True, exist_ok=True)
    diagram_dir = output_path.parent / f".{output_path.stem}-diagrams"
    document = Document()
    _configure_page(document, profile)
    _configure_styles(document, profile)

    lines = markdown_path.read_text(encoding="utf-8").splitlines()
    index = 0
    in_code_fence = False
    code_language = ""
    ordered_num_id: int | None = None
    while index < len(lines):
        raw_line = lines[index]
        line = raw_line.rstrip()
        stripped = line.strip()

        if stripped.startswith("```"):
            if not in_code_fence:
                in_code_fence = True
                code_language = stripped[3:].strip().lower()
            else:
                in_code_fence = False
                code_language = ""
            index += 1
            continue

        if in_code_fence:
            if code_language != "mermaid":
                paragraph = document.add_paragraph(style="Code Block")
                _set_run_font(paragraph.add_run(line))
            index += 1
            continue

        if not stripped:
            ordered_num_id = None
            index += 1
            continue

        if stripped == "<!-- PAGE_BREAK -->":
            ordered_num_id = None
            document.add_page_break()
            index += 1
            continue

        diagram_match = _DIAGRAM_MARKER.match(stripped)
        if diagram_match:
            ordered_num_id = None
            key = diagram_match.group(1)
            if key not in diagrams:
                raise KeyError(f"Missing DiagramSpec for marker: {key}")
            _add_diagram(document, key, diagrams[key], diagram_dir, font_path)
            index += 1
            continue

        if stripped.startswith("<!--") and stripped.endswith("-->"):
            ordered_num_id = None
            index += 1
            continue

        heading_match = _HEADING.match(stripped)
        if heading_match:
            ordered_num_id = None
            level = len(heading_match.group(1))
            text = heading_match.group(2)
            style = "Title" if level == 1 else f"Heading {min(level - 1, 3)}"
            paragraph = document.add_paragraph(style=style)
            _add_inline_markdown(paragraph, text)
            index += 1
            continue

        if stripped.startswith("|") and stripped.endswith("|"):
            ordered_num_id = None
            rows, index = _parse_table(lines, index)
            _add_table(document, rows)
            continue

        unordered_match = _UNORDERED_LIST.match(line)
        if unordered_match:
            ordered_num_id = None
            item_lines = [unordered_match.group(1)]
            index += 1
            while index < len(lines):
                continuation = lines[index]
                if not continuation.strip() or not continuation[:1].isspace():
                    break
                if (
                    _UNORDERED_LIST.match(continuation)
                    or _ORDERED_LIST.match(continuation)
                    or _BLOCKQUOTE.match(continuation.strip())
                ):
                    break
                item_lines.append(continuation.strip())
                index += 1
            paragraph = document.add_paragraph(style="List Bullet")
            _set_paragraph_numbering(paragraph, num_id=1)
            _add_inline_markdown(paragraph, " ".join(item_lines))
            continue

        ordered_match = _ORDERED_LIST.match(line)
        if ordered_match:
            if ordered_num_id is None:
                ordered_num_id = _new_numbering_instance(
                    document,
                    template_num_id=5,
                    start=int(ordered_match.group(1)),
                )
            item_lines = [ordered_match.group(2)]
            index += 1
            while index < len(lines):
                continuation = lines[index]
                if not continuation.strip() or not continuation[:1].isspace():
                    break
                if (
                    _UNORDERED_LIST.match(continuation)
                    or _ORDERED_LIST.match(continuation)
                    or _BLOCKQUOTE.match(continuation.strip())
                ):
                    break
                item_lines.append(continuation.strip())
                index += 1
            paragraph = document.add_paragraph(style="List Number")
            _set_paragraph_numbering(paragraph, num_id=ordered_num_id)
            _add_inline_markdown(paragraph, f"\u00a0{' '.join(item_lines)}")
            continue

        image_match = _IMAGE.match(stripped)
        if image_match:
            ordered_num_id = None
            alt_text, image_reference = image_match.groups()
            image_path = (markdown_path.parent / image_reference).resolve()
            _add_image(document, image_path, alt_text)
            index += 1
            continue

        if stripped in ("---", "***", "___"):
            ordered_num_id = None
            _add_rule(document)
            index += 1
            continue

        blockquote_match = _BLOCKQUOTE.match(stripped)
        if blockquote_match:
            ordered_num_id = None
            quote_lines = [blockquote_match.group(1)]
            index += 1
            while index < len(lines):
                next_match = _BLOCKQUOTE.match(lines[index].strip())
                if next_match is None:
                    break
                quote_lines.append(next_match.group(1))
                index += 1
            _add_blockquote(document, " ".join(quote_lines))
            continue

        paragraph_lines = [stripped]
        ordered_num_id = None
        index += 1
        while index < len(lines):
            candidate = lines[index].strip()
            if (
                not candidate
                or candidate.startswith(("#", "```", "|", "<!--", "!["))
                or _UNORDERED_LIST.match(lines[index])
                or _ORDERED_LIST.match(lines[index])
                or _BLOCKQUOTE.match(candidate)
                or candidate in ("---", "***", "___")
            ):
                break
            paragraph_lines.append(candidate)
            index += 1
        paragraph = document.add_paragraph()
        _add_inline_markdown(paragraph, " ".join(paragraph_lines))

    document.core_properties.title = markdown_path.stem
    document.core_properties.subject = f"Sookta Application 2.1.0 — {profile.name}"
    document.core_properties.author = "Sookta Project"
    document.core_properties.keywords = "Sookta, ergonomics, REBA, occupational health"
    document.save(output_path)


__all__ = [
    "A4_CONTENT_WIDTH_DXA",
    "BALANCED_PROFILE",
    "TECHNICAL_PROFILE",
    "DiagramSpec",
    "DocumentProfile",
    "build_docx",
    "render_flow_diagram",
]
