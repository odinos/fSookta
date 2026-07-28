from __future__ import annotations

import tempfile
import unittest
from pathlib import Path

from docx import Document
from docx.enum.section import WD_ORIENT
from docx.oxml.ns import qn

from tooling.sookta_document_builder import (
    A4_CONTENT_WIDTH_DXA,
    BALANCED_PROFILE,
    DiagramSpec,
    _table_widths,
    build_docx,
)


class SooktaDocumentBuilderTests(unittest.TestCase):
    def setUp(self) -> None:
        self.temp_dir = tempfile.TemporaryDirectory()
        self.addCleanup(self.temp_dir.cleanup)
        self.root = Path(self.temp_dir.name)
        self.font_path = Path("/System/Library/Fonts/Supplemental/Arial Unicode.ttf")

    def _build_sample(self) -> Document:
        markdown_path = self.root / "sample.md"
        output_path = self.root / "sample.docx"
        markdown_path.write_text(
            """# เอกสารทดสอบ

## หัวข้อหลัก

ย่อหน้าภาษาไทยสำหรับตรวจรูปแบบ

- รายการแบบจุด
  พร้อมข้อความต่อเนื่อง
- รายการแบบจุดลำดับสอง

1. รายการลำดับเลข
   พร้อมข้อความต่อเนื่อง
2. รายการลำดับเลขข้อสอง

| ตัวแปร | ความหมาย |
|---|---|
| BMI | ดัชนีมวลกาย |

> **หมายเหตุ:** ข้อความอธิบายสำคัญ

1. รายการเริ่มใหม่

<!-- DOCX_DIAGRAM:test-flow -->

```text
ตัวอย่างโค้ด
```
""",
            encoding="utf-8",
        )
        diagrams = {
            "test-flow": DiagramSpec(
                title="แผนภาพทดสอบ",
                nodes=(("start", "เริ่มต้น"), ("end", "สิ้นสุด")),
                edges=(("start", "end", "ถัดไป"),),
            )
        }
        build_docx(
            markdown_path=markdown_path,
            output_path=output_path,
            profile=BALANCED_PROFILE,
            diagrams=diagrams,
            font_path=self.font_path,
        )
        return Document(output_path)

    def test_balanced_profile_uses_a4_sarabun_and_expected_margins(self) -> None:
        document = self._build_sample()
        section = document.sections[0]

        self.assertEqual(section.orientation, WD_ORIENT.PORTRAIT)
        self.assertAlmostEqual(section.page_width.mm, 210.0, places=1)
        self.assertAlmostEqual(section.page_height.mm, 297.0, places=1)
        self.assertAlmostEqual(section.top_margin.cm, 2.54, places=2)
        self.assertAlmostEqual(section.bottom_margin.cm, 2.54, places=2)
        self.assertAlmostEqual(section.left_margin.cm, 2.54, places=2)
        self.assertAlmostEqual(section.right_margin.cm, 2.54, places=2)
        self.assertEqual(document.styles["Normal"].font.name, "Sarabun")

    def test_markdown_headings_lists_tables_and_diagram_marker_render(self) -> None:
        document = self._build_sample()
        paragraph_styles = [paragraph.style.name for paragraph in document.paragraphs]

        self.assertIn("Title", paragraph_styles)
        self.assertIn("Heading 1", paragraph_styles)
        self.assertEqual(len(document.tables), 1)
        self.assertGreaterEqual(len(document.inline_shapes), 1)
        self.assertNotIn("<!-- DOCX_DIAGRAM:test-flow -->", "\n".join(
            paragraph.text for paragraph in document.paragraphs
        ))
        self.assertIn(
            "หมายเหตุ: ข้อความอธิบายสำคัญ",
            [paragraph.text for paragraph in document.paragraphs],
        )
        self.assertNotIn(
            ">",
            "\n".join(paragraph.text for paragraph in document.paragraphs),
        )

    def test_table_geometry_matches_a4_content_width(self) -> None:
        document = self._build_sample()
        table = document.tables[0]
        grid_widths = [
            int(grid_column.get(qn("w:w")))
            for grid_column in table._tbl.tblGrid.gridCol_lst
        ]

        self.assertEqual(sum(grid_widths), A4_CONTENT_WIDTH_DXA)
        self.assertEqual(table.autofit, False)

    def test_wide_traceability_table_prioritizes_source_column(self) -> None:
        widths = _table_widths(6)

        self.assertEqual(sum(widths), A4_CONTENT_WIDTH_DXA)
        self.assertEqual(len(widths), 6)
        self.assertEqual(max(widths), widths[3])
        self.assertGreater(widths[3], widths[0])

    def test_numbered_and_bullet_lists_use_word_numbering(self) -> None:
        document = self._build_sample()
        list_paragraphs = [
            paragraph
            for paragraph in document.paragraphs
            if paragraph.text.strip().startswith("รายการ")
        ]

        self.assertEqual(len(list_paragraphs), 5)
        for paragraph in list_paragraphs:
            self.assertIsNotNone(paragraph._p.pPr)
            self.assertIsNotNone(paragraph._p.pPr.numPr)
            self.assertIsNotNone(paragraph._p.pPr.numPr.numId)
        self.assertIn(
            "รายการแบบจุด พร้อมข้อความต่อเนื่อง",
            [paragraph.text.strip() for paragraph in list_paragraphs],
        )
        self.assertIn(
            "รายการลำดับเลข พร้อมข้อความต่อเนื่อง",
            [paragraph.text.strip() for paragraph in list_paragraphs],
        )
        ordered = [
            paragraph
            for paragraph in list_paragraphs
            if paragraph.style.name == "List Number"
        ]
        self.assertEqual(ordered[0]._p.pPr.numPr.numId.val, ordered[1]._p.pPr.numPr.numId.val)
        self.assertNotEqual(ordered[1]._p.pPr.numPr.numId.val, ordered[2]._p.pPr.numPr.numId.val)

    def test_all_styles_set_ascii_hansi_eastasia_and_complex_script_to_sarabun(self) -> None:
        document = self._build_sample()
        required_styles = (
            "Normal",
            "Title",
            "Subtitle",
            "Heading 1",
            "Heading 2",
            "Heading 3",
            "List Bullet",
            "List Number",
            "Caption",
        )

        for style_name in required_styles:
            with self.subTest(style=style_name):
                fonts = document.styles[style_name].element.get_or_add_rPr().get_or_add_rFonts()
                self.assertEqual(fonts.get(qn("w:ascii")), "Sarabun")
                self.assertEqual(fonts.get(qn("w:hAnsi")), "Sarabun")
                self.assertEqual(fonts.get(qn("w:eastAsia")), "Sarabun")
                self.assertEqual(fonts.get(qn("w:cs")), "Sarabun")


if __name__ == "__main__":
    unittest.main()
